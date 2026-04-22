# filepath: app/warehouse_manager_nosqlite_excel_billing.py
# Python 3.10+

try:
    import logging_setup
    logging_setup.install_crash_logging("HosnyPOS-ZAY")
except Exception:
    pass

import json
import os
import hashlib
import sys, subprocess
import sqlite3
import time
import tempfile
import webbrowser
from dataclasses import dataclass
from datetime import datetime, date   # +date for calendar
import calendar                       # ADD
from typing import Any, Dict, List, Optional, Sequence, Tuple

import tkinter as tk
from tkinter import ttk, messagebox, filedialog, simpledialog

DB_PATH = "warehouse_data.sqlite3"
LEGACY_JSON_PATH = "warehouse_data.json"
ADMIN_PASSWORD_PLAIN = "1234"
ADMIN_PASSWORD_HASH_PREFIX = "sha256$"
WAREHOUSE_DEVICE_NAME = "WAREHOUSE"
WAREHOUSE_RETURN_LABEL = "الى المصنع"
# Bill rows for stock sent to warehouse review — not retail sales (no POS income / dashboard sales).
WAREHOUSE_RETURN_BILL_TYPE = "WAREHOUSE_RETURN"
BRANCH_TARGET_PREFIX = "فرع: "
POS_CASHIER_LOCKDOWN = True
POS_DISCOUNT_PRESETS = (5, 10, 15)
POS_MANAGER_FEATURE_DEFAULTS = {
    "allow_inventory_window": not POS_CASHIER_LOCKDOWN,
    "allow_manual_incoming": not POS_CASHIER_LOCKDOWN,
    "allow_bulk_price": not POS_CASHIER_LOCKDOWN,
    "allow_excel_import": not POS_CASHIER_LOCKDOWN,
    "allow_manual_adjustment": not POS_CASHIER_LOCKDOWN,
    "allow_reset_counts": not POS_CASHIER_LOCKDOWN,
    "allow_inventory_delete": not POS_CASHIER_LOCKDOWN,
    "allow_inventory_price_edit": not POS_CASHIER_LOCKDOWN,
    "allow_inventory_specs_edit": not POS_CASHIER_LOCKDOWN,
    "allow_size_profile_edit": not POS_CASHIER_LOCKDOWN,
}
DEFAULT_BRANCH_POS_NAMES = [
    "POS-ZAY",
    "POS-OCT",
    "POS-OBO",
    "POS-GESR",
    "POS-BAH",
    "POS-CEN",
]
BRANCH_UI_NAME_BY_DEVICE = {
    "POS-ZAY": "فرع زايد",
    "POS-OCT": "فرع اكتوبر",
    "POS-BAH": "فرع بهتيم",
    "POS-CEN": "فرع السنتر",
    "POS-OBO": "فرع العبور",
    "POS-GESR": "فرع جسر السويس",
}
BRANCH_DEVICE_BY_UI_NAME = {v: k for k, v in BRANCH_UI_NAME_BY_DEVICE.items()}

ALLOWED_NUMERIC_RANGES = {
    (0, 16): [str(i) for i in range(0, 18, 2)],
    (6, 22): [str(i) for i in range(6, 24, 2)],
    (14, 28): [str(i) for i in range(14, 30, 2)],
    (18, 30): [str(i) for i in range(18, 32, 2)],
    (0, 9):  [str(i) for i in range(0, 10, 1)],
    (32, 62): [str(i) for i in range(32, 63,2)],

}
NUMERIC_RANGE_LABELS = [
    f"{a} → {b}" for (a, b) in ALLOWED_NUMERIC_RANGES.keys()
]

ALPHA_SIZES = ["S", "M", "L", "XL", "2XL", "3XL", "4XL", "5XL"]


def merged_numeric_size_labels_from_profile(
    r1s: Optional[int],
    r1e: Optional[int],
    r2s: Optional[int],
    r2e: Optional[int],
) -> List[str]:
    """Union of predefined numeric columns for range 1 and range 2 (no duplicate sizes)."""
    chunks: List[List[str]] = []
    if r1s is not None and r1e is not None:
        labs = ALLOWED_NUMERIC_RANGES.get((r1s, r1e))
        if labs:
            chunks.append(list(labs))
    if r2s is not None and r2e is not None:
        labs = ALLOWED_NUMERIC_RANGES.get((r2s, r2e))
        if labs:
            chunks.append(list(labs))
    if not chunks:
        return []
    numeric_vals: List[int] = []
    extras: List[str] = []
    for chunk in chunks:
        for s in chunk:
            t = str(s).strip()
            if not t:
                continue
            if t.isdigit():
                numeric_vals.append(int(t))
            elif t not in extras:
                extras.append(t)
    uniq_sorted = sorted(set(numeric_vals))
    out = [str(i) for i in uniq_sorted]
    out.extend(extras)
    return out


def now_iso() -> str:
    return datetime.now().isoformat(timespec="seconds")


def _cashier_lockdown_message(action: str = "") -> str:
    base = "هذا الإجراء متوقف في وضع الكاشير. المخزن فقط هو المخوّل بتغيير المخزون أو الأسعار."
    if action:
        return f"{action}\n\n{base}"
    return base


def _feature_restricted_message(action: str = "") -> str:
    base = "هذه الميزة مقيدة حالياً من تبويب صلاحيات المدير."
    if action:
        return f"{action}\n\n{base}"
    return base


def _extract_warehouse_target(value: Any) -> Optional[str]:
    raw = str(value or "").strip()
    if not raw:
        return None
    if raw == WAREHOUSE_RETURN_LABEL:
        return WAREHOUSE_DEVICE_NAME
    if raw.upper() == WAREHOUSE_DEVICE_NAME:
        return WAREHOUSE_DEVICE_NAME
    if raw.startswith(f"{WAREHOUSE_RETURN_LABEL}:"):
        name = raw.split(":", 1)[1].strip()
        if name.upper() == WAREHOUSE_DEVICE_NAME:
            return WAREHOUSE_DEVICE_NAME
    return None


def _extract_branch_target(value: Any) -> Optional[str]:
    raw = str(value or "").strip()
    if not raw:
        return None
    if raw.startswith(BRANCH_TARGET_PREFIX):
        raw = raw[len(BRANCH_TARGET_PREFIX):].strip()
    if raw in DEFAULT_BRANCH_POS_NAMES:
        return raw
    if raw in BRANCH_DEVICE_BY_UI_NAME:
        return BRANCH_DEVICE_BY_UI_NAME[raw]
    return None


def _branch_display_name(value: Any) -> str:
    raw = str(value or "").strip()
    if not raw:
        return ""
    return BRANCH_UI_NAME_BY_DEVICE.get(raw, raw)


@dataclass
class StockRow:
    id: int
    item_type: str
    school: str
    color: str
    size: str
    unit_price: float
    count: int


# ------------------- Excel exporting -------------------

def export_to_excel(path: str, headers: List[str], rows: Sequence[Sequence[Any]]) -> None:
    ext = os.path.splitext(path)[1].lower()
    try:
        import openpyxl  # type: ignore
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Sheet1"
        ws.append(headers)
        for row in rows:
            ws.append(list(row))
        for idx, _ in enumerate(headers, start=1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(idx)].width = 18
        wb.save(path)
        return
    except Exception:
        xml = []
        xml.append('<?xml version="1.0" encoding="UTF-8"?>')
        xml.append(
            '<Workbook xmlns="urn:schemas-microsoft-com:office:spreadsheet" '
            'xmlns:o="urn:schemas-microsoft-com:office:office" '
            'xmlns:x="urn:schemas-microsoft-com:office:excel" '
            'xmlns:ss="urn:schemas-microsoft-com:office:spreadsheet">'
        )
        xml.append('<Worksheet ss:Name="Sheet1"><Table>')
        xml.append("<Row>")
        for h in headers:
            xml.append(f'<Cell><Data ss:Type="String">{_xml_escape(str(h))}</Data></Cell>')
        xml.append("</Row>")
        for r in rows:
            xml.append("<Row>")
            for v in r:
                if isinstance(v, (int, float)):
                    xml.append(f'<Cell><Data ss:Type="Number">{v}</Data></Cell>')
                else:
                    xml.append(f'<Cell><Data ss:Type="String">{_xml_escape("" if v is None else str(v))}</Data></Cell>')
            xml.append("</Row>")
        xml.append("</Table></Worksheet></Workbook>")
        content = "\n".join(xml)
        if ext not in (".xls", ".xlsx", ".xml"):
            path += ".xls"
        with open(path, "w", encoding="utf-8") as f:
            f.write(content)


def _xml_escape(s: str) -> str:
    return (
        s.replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
        .replace('"', "&quot;")
        .replace("'", "&apos;")
    )

def _file_url(path: str) -> str:
    return f"file:///{path.replace(os.sep, '/')}"

def _print_html_auto(path: str, copies: int = 1, parent: Optional[tk.Widget] = None) -> None:
    """
    Try silent HTML print (Windows COM) or shell 'print'. Else open auto-printing copy.
    """
    copies = max(1, int(copies))
    import time

    if sys.platform.startswith("win"):
        try:
            import win32com.client  # type: ignore
            ie = win32com.client.Dispatch("InternetExplorer.Application")
            ie.Visible = False
            url = _file_url(path)
            for _ in range(copies):
                ie.Navigate(url)
                while True:
                    time.sleep(0.1)
                    try:
                        if not ie.Busy and int(ie.ReadyState) == 4:
                            doc = getattr(ie, "Document", None)
                            if doc is not None and str(getattr(doc, "readyState", "")).lower() == "complete":
                                break
                    except Exception:
                        break
                time.sleep(0.2)
                ie.ExecWB(6, 2)
                time.sleep(0.2)

            ie.Quit()
            return
        except Exception:
            pass
        try:
            for _ in range(copies):
                os.startfile(path, "print")  # type: ignore[attr-defined]
            return
        except Exception:
            pass

    try:
        auto_path = path[:-5] + "_autoprint.html" if path.lower().endswith(".html") else path + "_autoprint.html"
        try:
            with open(path, "r", encoding="utf-8") as f:
                html = f.read()
            edge_meta = '<meta http-equiv="X-UA-Compatible" content="IE=edge" />'
            if "</head>" in html:
                if "X-UA-Compatible" not in html:
                    html = html.replace("</head>", edge_meta + "</head>")
                html = html.replace(
                    "</head>",
                    "<script>window.onload=function(){try{window.print();}catch(e){} setTimeout(()=>window.close(),600);};</script></head>"
                )
            else:
                if "X-UA-Compatible" not in html:
                    html = edge_meta + html
                html = html + "<script>window.onload=function(){try{window.print();}catch(e){} setTimeout(()=>window.close(),600);};</script>"

            with open(auto_path, "w", encoding="utf-8") as f:
                f.write(html)
        except Exception:
            auto_path = path

        for _ in range(copies):
            webbrowser.open_new_tab(_file_url(auto_path))
            time.sleep(0.05)

        if parent is not None:
            messagebox.showinfo("الطباعة", f"تم فتح الفاتورة للطباعة تلقائيًا.\nعدد النسخ: {copies}.", parent=parent)
    except Exception as ex:
        if parent is not None:
            messagebox.showerror("فشل الطباعة", f"{ex}", parent=parent)

def save_bill_as_html(path: str, bill: Dict[str, Any], items: List[Dict[str, Any]]) -> None:
    """Generate receipt-style HTML for thermal printers (80mm width)."""

    def _fmtf(x: Any) -> str:
        try:
            return f"{float(x):.2f}"
        except Exception:
            return "0.00"

    rows_html = ""
    for ln in items:
        line_total = _fmtf(ln.get('line_total') or (float(ln.get('unit_price', 0)) * int(ln.get('qty', 0))))
        rows_html += f"""<tr>
<td>{_html(ln['item_type'])} - {_html(ln['school'])}<br><small>{_html(ln['color'])} / {_html(ln['size'])}</small></td>
<td style="text-align:center">{ln['qty']}</td>
<td style="text-align:left">{_fmtf(ln['unit_price'])}</td>
<td style="text-align:left">{line_total}</td>
</tr>
"""

    customer = _html(bill.get('customer') or '')
    customer_line = f"<div>العميل: {customer}</div>" if customer else ""

    html = f"""<!DOCTYPE html>
<html lang="ar" dir="rtl">
<head>
<meta charset="utf-8">
<title>فاتورة #{bill['id']}</title>
<style>
  @page {{ size: 80mm auto; margin: 2mm; }}
  * {{ box-sizing: border-box; }}
  body {{
    font-family: "Segoe UI", Tahoma, Arial, sans-serif;
    font-size: 11px;
    margin: 0;
    padding: 0;
    width: 76mm;
    direction: rtl;
  }}
  .receipt {{
    padding: 2mm;
  }}
  .center {{ text-align: center; }}
  .sep {{
    border: none;
    border-top: 1px dashed #000;
    margin: 4px 0;
  }}
  h2 {{
    font-size: 14px;
    margin: 4px 0;
    text-align: center;
  }}
  .info {{
    font-size: 11px;
    margin: 2px 0;
  }}
  table {{
    border-collapse: collapse;
    width: 100%;
    font-size: 10px;
  }}
  th {{
    border-bottom: 1px solid #000;
    padding: 3px 2px;
    text-align: right;
  }}
  td {{
    padding: 3px 2px;
    vertical-align: top;
    border-bottom: 1px dotted #ccc;
  }}
  .total-row {{
    font-size: 14px;
    font-weight: bold;
    text-align: center;
    margin: 6px 0;
  }}
  .footer {{
    text-align: center;
    font-size: 10px;
    margin-top: 6px;
  }}
</style>
</head>
<body>
<div class="receipt">
  <h2>فاتورة #{bill['id']}</h2>
  <hr class="sep">
  <div class="info">التاريخ: {bill['created_at']}</div>
  {customer_line}
  <hr class="sep">
  <table>
    <thead>
      <tr>
        <th>الصنف</th>
        <th>الكمية</th>
        <th>السعر</th>
        <th>المجموع</th>
      </tr>
    </thead>
    <tbody>
      {rows_html}
    </tbody>
  </table>
  <hr class="sep">
  <div class="total-row">الإجمالي: {float(bill['total']):.2f}</div>
  <hr class="sep">
  <div class="footer">شكراً لتعاملكم معنا</div>
</div>
</body>
</html>"""

    with open(path, "w", encoding="utf-8") as f:
        f.write(html)

def _html(s: str) -> str:
    return (s or "").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

# Normalize size text for matching (handles Arabic digits & case)
_AR_DIGITS = str.maketrans("٠١٢٣٤٥٦٧٨٩", "0123456789")
def _normalize_size_label(s: str) -> str:
    return (s or "").strip().translate(_AR_DIGITS).upper()


def _parse_money_amount(raw: Any) -> float:
    """Parse user-entered money (Arabic/Latin digits, common decimal separators)."""
    if raw is None:
        return 0.0
    s = str(raw).strip().translate(_AR_DIGITS)
    for ch in ("\u066b", "\u060c"):  # Arabic decimal / thousands comma
        s = s.replace(ch, ".")
    s = s.replace(" ", "").replace("\u00a0", "").replace("\u2009", "")
    if not s:
        return 0.0
    # If only comma as separator, treat as decimal (e.g. 40,50)
    if "," in s and "." not in s:
        s = s.replace(",", ".")
    else:
        s = s.replace(",", "")
    return float(s)


# ------------------- SQLite DB -------------------

def apply_focus_highlight(widget, color="#19E72A"):
    default_bg = widget.cget("background")

    def on_focus_in(e):
        try:
            widget.configure(background=color)
        except:
            pass

    def on_focus_out(e):
        try:
            widget.configure(background=default_bg)
        except:
            pass

    widget.bind("<FocusIn>", on_focus_in)
    widget.bind("<FocusOut>", on_focus_out)

class SqliteDatabase:
    """SQLite persistence layer."""

    def __init__(self, path: str = DB_PATH, legacy_json: str = LEGACY_JSON_PATH) -> None:
        self.path = path
        self.conn = sqlite3.connect(self.path, isolation_level=None, check_same_thread=False)
        self.conn.row_factory = sqlite3.Row
        self.active_shift_id: Optional[int] = None

        self._apply_pragmas()
        self._init_schema()
        self._migrate_from_json_if_empty(legacy_json)


    def _require_shift(self):
        if self.active_shift_id is None:
            raise ValueError("لا يمكن إجراء عمليات بدون وردية مفتوحة.")

    def _apply_pragmas(self) -> None:
        cur = self.conn.cursor()
        cur.execute("PRAGMA journal_mode=WAL;")
        cur.execute("PRAGMA synchronous=NORMAL;")
        cur.execute("PRAGMA foreign_keys=ON;")
        cur.close()

    def _record_sync_event(self, event_type: str, payload: Dict[str, Any]) -> None:
        """Append a business event to sync_outbox. Fail-safe.

        Phase 1: events are written but never pushed. If the sync layer
        is not present or the outbox insert fails, the caller continues
        without error — we must never break a business operation.
        """
        try:
            from sync_core import record_event
            record_event(self.conn, event_type, payload)
        except Exception:
            import traceback
            traceback.print_exc()

    def _record_warehouse_return_event(
        self,
        return_uuid: str,
        note: str,
        lines: List[Dict[str, Any]],
    ) -> None:
        """Append a warehouse-targeted stock return event to sync_outbox."""
        self._record_targeted_inventory_event(
            event_type="STOCK_RETURN_TO_WAREHOUSE",
            target_scope="warehouse",
            payload={
                "return_uuid": return_uuid,
                "from_device": self._current_device_name() or "pos",
                "to_device": WAREHOUSE_DEVICE_NAME,
                "note": note or "",
                "items": lines,
            },
        )

    def _record_transfer_via_warehouse_event(
        self,
        request_uuid: str,
        target_device: str,
        note: str,
        lines: List[Dict[str, Any]],
    ) -> None:
        self._record_targeted_inventory_event(
            event_type="POS_TRANSFER_VIA_WAREHOUSE",
            target_scope="warehouse",
            payload={
                "request_uuid": request_uuid,
                "from_device": self._current_device_name() or "pos",
                "target_device": target_device,
                "note": note or "",
                "items": lines,
            },
        )

    def _current_device_name(self) -> Optional[str]:
        try:
            ident = self.conn.execute(
                "SELECT device_name FROM device_identity WHERE id = 1"
            ).fetchone()
            if ident is not None:
                return ident[0]
        except Exception:
            pass
        return None

    def _record_targeted_inventory_event(
        self,
        event_type: str,
        target_scope: str,
        payload: Dict[str, Any],
    ) -> None:
        import json as _json
        import uuid as _uuid
        try:
            from sync_client import _utc_now_iso as _sync_now  # type: ignore
        except Exception:
            _sync_now = lambda: datetime.utcnow().isoformat(timespec="seconds") + "Z"
        event_uuid = str(_uuid.uuid4())
        now = _sync_now()

        try:
            has_target = False
            try:
                cur = self.conn.execute("PRAGMA table_info(sync_outbox)")
                has_target = any(r[1] == "target_scope" for r in cur.fetchall())
            except Exception:
                has_target = False

            if has_target:
                self.conn.execute(
                    """
                    INSERT INTO sync_outbox
                        (event_uuid, event_type, payload_json,
                         created_at, status, attempts, target_scope)
                    VALUES (?, ?, ?, ?, 'pending', 0, ?)
                    """,
                    (
                        event_uuid,
                        event_type,
                        _json.dumps(payload, ensure_ascii=False, default=str),
                        now,
                        target_scope,
                    ),
                )
            else:
                payload["__target_scope__"] = target_scope
                self.conn.execute(
                    """
                    INSERT INTO sync_outbox
                        (event_uuid, event_type, payload_json,
                         created_at, status, attempts)
                    VALUES (?, ?, ?, ?, 'pending', 0)
                    """,
                    (
                        event_uuid,
                        event_type,
                        _json.dumps(payload, ensure_ascii=False, default=str),
                        now,
                    ),
                )
        except Exception:
            import traceback
            traceback.print_exc()

    def _init_schema(self) -> None:
        cur = self.conn.cursor()
        cur.executescript(
            """
            -- Default price per item (used when income price is empty)
            CREATE TABLE IF NOT EXISTS item_defaults (
                item_type TEXT PRIMARY KEY,
                default_price REAL NOT NULL
            );
            CREATE TABLE IF NOT EXISTS size_profiles (
                id INTEGER PRIMARY KEY AUTOINCREMENT,

                item_type TEXT NOT NULL,
                school TEXT NOT NULL,
                color TEXT NOT NULL,

                -- numeric ranges (nullable)
                num_start_1 INTEGER,
                num_end_1   INTEGER,

                num_start_2 INTEGER,
                num_end_2   INTEGER,

                -- flags
                has_alpha INTEGER NOT NULL DEFAULT 0,

                updated_at TEXT NOT NULL,

                UNIQUE(item_type, school, color)
            );

            CREATE TABLE IF NOT EXISTS stocks(
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                item_type TEXT NOT NULL,
                school TEXT NOT NULL,
                color TEXT NOT NULL,
                size TEXT NOT NULL,
                unit_price REAL NOT NULL,
                count INTEGER NOT NULL
            );
            CREATE INDEX IF NOT EXISTS idx_stocks_specs
            ON stocks(item_type, school, color, size);
            CREATE INDEX IF NOT EXISTS idx_stocks_count
            ON stocks(count);

            CREATE TABLE IF NOT EXISTS movements(
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                ts TEXT NOT NULL,
                direction TEXT NOT NULL, -- IN | OUT | ADJUST_OUT | OUT_FACTORY | PRICE_UPDATE | RESERVE
                stock_id INTEGER,
                qty INTEGER NOT NULL,
                note TEXT,
                bill_id INTEGER,
                item_type TEXT,
                school TEXT,
                color TEXT,
                size TEXT,
                unit_price REAL
            );
            CREATE INDEX IF NOT EXISTS idx_movements_ts ON movements(ts DESC);
            CREATE INDEX IF NOT EXISTS idx_movements_dir ON movements(direction);
            CREATE INDEX IF NOT EXISTS idx_movements_specs ON movements(item_type,school,color,size);

            CREATE TABLE IF NOT EXISTS bills(
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                created_at TEXT NOT NULL,
                customer TEXT,
                total REAL NOT NULL
            );
            CREATE INDEX IF NOT EXISTS idx_bills_created ON bills(created_at DESC);

            CREATE TABLE IF NOT EXISTS bill_items(
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                bill_id INTEGER NOT NULL,
                item_type TEXT NOT NULL,
                school TEXT NOT NULL,
                color TEXT NOT NULL,
                size TEXT NOT NULL,
                unit_price REAL NOT NULL,
                qty INTEGER NOT NULL,
                line_total REAL NOT NULL,
                origin TEXT NOT NULL DEFAULT 'STOCK'
            );
            CREATE INDEX IF NOT EXISTS idx_bill_items_bill ON bill_items(bill_id);

            CREATE TABLE IF NOT EXISTS spec_history(
                field TEXT NOT NULL,
                value TEXT NOT NULL,
                PRIMARY KEY(field, value)
            );

            CREATE TABLE IF NOT EXISTS reservations(
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                created_at TEXT NOT NULL,
                customer TEXT,
                item_type TEXT NOT NULL,
                school TEXT NOT NULL,
                color TEXT NOT NULL,
                size TEXT NOT NULL,
                qty INTEGER NOT NULL,
                unit_price REAL NOT NULL,
                total_amount REAL NOT NULL,
                paid_amount REAL NOT NULL DEFAULT 0,
                status TEXT NOT NULL DEFAULT 'معلق',
                note TEXT
            );
            CREATE INDEX IF NOT EXISTS idx_reservations_status ON reservations(status);
            CREATE INDEX IF NOT EXISTS idx_reservations_created ON reservations(created_at DESC);

            CREATE TABLE IF NOT EXISTS shifts(
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                started_at TEXT NOT NULL,
                ended_at TEXT,
                status TEXT NOT NULL DEFAULT 'OPEN',
                summary_json TEXT,
                note TEXT
            );
            CREATE INDEX IF NOT EXISTS idx_shifts_status ON shifts(status);
            CREATE INDEX IF NOT EXISTS idx_shifts_started ON shifts(started_at DESC);

            CREATE TABLE IF NOT EXISTS income_bills(
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                created_at TEXT NOT NULL,
                supplier TEXT,
                total_qty INTEGER NOT NULL DEFAULT 0,
                total_value REAL NOT NULL DEFAULT 0,
                note TEXT
            );
            CREATE INDEX IF NOT EXISTS idx_income_bills_created ON income_bills(created_at DESC);

            CREATE TABLE IF NOT EXISTS income_bill_items(
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                income_bill_id INTEGER NOT NULL,
                item_type TEXT NOT NULL,
                school TEXT NOT NULL,
                color TEXT NOT NULL,
                size TEXT NOT NULL,
                unit_price REAL NOT NULL,
                qty INTEGER NOT NULL,
                line_total REAL NOT NULL
            );
            CREATE INDEX IF NOT EXISTS idx_income_bill_items_bill ON income_bill_items(income_bill_id);

            CREATE TABLE IF NOT EXISTS app_settings(
                key TEXT PRIMARY KEY,
                value TEXT NOT NULL
            );
            """
        )

        # ensure 'origin' exists
        try:
            cols = {r["name"] for r in self.conn.execute("PRAGMA table_info(bill_items)")}
            if "origin" not in cols:
                self.conn.execute(
                    "ALTER TABLE bill_items ADD COLUMN origin TEXT NOT NULL DEFAULT 'STOCK'"
                )
        except Exception:
            pass

        # Migration: bill_type column on bills table
        try:
            cols = {r["name"] for r in self.conn.execute("PRAGMA table_info(bills)")}
            if "bill_type" not in cols:
                self.conn.execute("ALTER TABLE bills ADD COLUMN bill_type TEXT NOT NULL DEFAULT 'SALE'")
        except Exception:
            pass

        # Migration: immutable/voidable finalized bills
        try:
            cols = {r["name"] for r in self.conn.execute("PRAGMA table_info(bills)")}
            if "status" not in cols:
                self.conn.execute("ALTER TABLE bills ADD COLUMN status TEXT NOT NULL DEFAULT 'CONFIRMED'")
            if "void_reason" not in cols:
                self.conn.execute("ALTER TABLE bills ADD COLUMN void_reason TEXT")
            if "voided_at" not in cols:
                self.conn.execute("ALTER TABLE bills ADD COLUMN voided_at TEXT")
        except Exception:
            pass

        # Migration: remove warehouse_no/package_no from existing tables
        try:
            cols = {r["name"] for r in self.conn.execute("PRAGMA table_info(stocks)")}
            if "warehouse_no" in cols:
                self.conn.executescript("""
                    CREATE TABLE IF NOT EXISTS stocks_new(
                        id INTEGER PRIMARY KEY AUTOINCREMENT,
                        item_type TEXT NOT NULL,
                        school TEXT NOT NULL,
                        color TEXT NOT NULL,
                        size TEXT NOT NULL,
                        unit_price REAL NOT NULL,
                        count INTEGER NOT NULL
                    );
                    INSERT INTO stocks_new(item_type, school, color, size, unit_price, count)
                        SELECT item_type, school, color, size, unit_price, SUM(count)
                        FROM stocks GROUP BY item_type, school, color, size, unit_price;
                    DROP TABLE stocks;
                    ALTER TABLE stocks_new RENAME TO stocks;
                    CREATE INDEX IF NOT EXISTS idx_stocks_specs ON stocks(item_type, school, color, size);
                    CREATE INDEX IF NOT EXISTS idx_stocks_count ON stocks(count);
                """)
        except Exception:
            pass

        try:
            cols = {r["name"] for r in self.conn.execute("PRAGMA table_info(movements)")}
            if "warehouse_no" in cols:
                self.conn.executescript("""
                    CREATE TABLE IF NOT EXISTS movements_new(
                        id INTEGER PRIMARY KEY AUTOINCREMENT,
                        ts TEXT NOT NULL,
                        direction TEXT NOT NULL,
                        stock_id INTEGER,
                        qty INTEGER NOT NULL,
                        note TEXT,
                        bill_id INTEGER,
                        item_type TEXT,
                        school TEXT,
                        color TEXT,
                        size TEXT,
                        unit_price REAL
                    );
                    INSERT INTO movements_new(ts, direction, stock_id, qty, note, bill_id, item_type, school, color, size, unit_price)
                        SELECT ts, direction, stock_id, qty, note, bill_id, item_type, school, color, size, unit_price
                        FROM movements;
                    DROP TABLE movements;
                    ALTER TABLE movements_new RENAME TO movements;
                    CREATE INDEX IF NOT EXISTS idx_movements_ts ON movements(ts DESC);
                    CREATE INDEX IF NOT EXISTS idx_movements_dir ON movements(direction);
                    CREATE INDEX IF NOT EXISTS idx_movements_specs ON movements(item_type,school,color,size);
                """)
        except Exception:
            pass

        try:
            cols = {r["name"] for r in self.conn.execute("PRAGMA table_info(bill_items)")}
            if "warehouse_no" in cols:
                self.conn.executescript("""
                    CREATE TABLE IF NOT EXISTS bill_items_new(
                        id INTEGER PRIMARY KEY AUTOINCREMENT,
                        bill_id INTEGER NOT NULL,
                        item_type TEXT NOT NULL,
                        school TEXT NOT NULL,
                        color TEXT NOT NULL,
                        size TEXT NOT NULL,
                        unit_price REAL NOT NULL,
                        qty INTEGER NOT NULL,
                        line_total REAL NOT NULL,
                        origin TEXT NOT NULL DEFAULT 'STOCK'
                    );
                    INSERT INTO bill_items_new(bill_id, item_type, school, color, size, unit_price, qty, line_total, origin)
                        SELECT bill_id, item_type, school, color, size, unit_price, qty, line_total, COALESCE(origin,'STOCK')
                        FROM bill_items;
                    DROP TABLE bill_items;
                    ALTER TABLE bill_items_new RENAME TO bill_items;
                    CREATE INDEX IF NOT EXISTS idx_bill_items_bill ON bill_items(bill_id);
                """)
        except Exception:
            pass

        # Drop packages table (no longer needed)
        try:
            self.conn.execute("DROP TABLE IF EXISTS packages")
        except Exception:
            pass

        # Migration: remove has_badge
        try:
            cols = {r["name"] for r in self.conn.execute("PRAGMA table_info(stocks)")}
            if "has_badge" in cols:
                self.conn.executescript("""
                    CREATE TABLE IF NOT EXISTS stocks_new2(
                        id INTEGER PRIMARY KEY AUTOINCREMENT,
                        item_type TEXT NOT NULL, school TEXT NOT NULL, color TEXT NOT NULL,
                        size TEXT NOT NULL, unit_price REAL NOT NULL, count INTEGER NOT NULL
                    );
                    INSERT INTO stocks_new2(item_type, school, color, size, unit_price, count)
                        SELECT item_type, school, color, size, unit_price, SUM(count)
                        FROM stocks GROUP BY item_type, school, color, size, unit_price;
                    DROP TABLE stocks;
                    ALTER TABLE stocks_new2 RENAME TO stocks;
                    CREATE INDEX IF NOT EXISTS idx_stocks_specs ON stocks(item_type, school, color, size);
                    CREATE INDEX IF NOT EXISTS idx_stocks_count ON stocks(count);
                """)
        except Exception:
            pass

        try:
            cols = {r["name"] for r in self.conn.execute("PRAGMA table_info(movements)")}
            if "has_badge" in cols:
                self.conn.executescript("""
                    CREATE TABLE IF NOT EXISTS movements_new2(
                        id INTEGER PRIMARY KEY AUTOINCREMENT,
                        ts TEXT NOT NULL, direction TEXT NOT NULL, stock_id INTEGER,
                        qty INTEGER NOT NULL, note TEXT, bill_id INTEGER,
                        item_type TEXT, school TEXT, color TEXT, size TEXT, unit_price REAL
                    );
                    INSERT INTO movements_new2(ts, direction, stock_id, qty, note, bill_id, item_type, school, color, size, unit_price)
                        SELECT ts, direction, stock_id, qty, note, bill_id, item_type, school, color, size, unit_price
                        FROM movements;
                    DROP TABLE movements;
                    ALTER TABLE movements_new2 RENAME TO movements;
                    CREATE INDEX IF NOT EXISTS idx_movements_ts ON movements(ts DESC);
                    CREATE INDEX IF NOT EXISTS idx_movements_dir ON movements(direction);
                    CREATE INDEX IF NOT EXISTS idx_movements_specs ON movements(item_type,school,color,size);
                """)
        except Exception:
            pass

        try:
            cols = {r["name"] for r in self.conn.execute("PRAGMA table_info(bill_items)")}
            if "has_badge" in cols:
                self.conn.executescript("""
                    CREATE TABLE IF NOT EXISTS bill_items_new2(
                        id INTEGER PRIMARY KEY AUTOINCREMENT,
                        bill_id INTEGER NOT NULL, item_type TEXT NOT NULL, school TEXT NOT NULL,
                        color TEXT NOT NULL, size TEXT NOT NULL, unit_price REAL NOT NULL,
                        qty INTEGER NOT NULL, line_total REAL NOT NULL,
                        origin TEXT NOT NULL DEFAULT 'STOCK'
                    );
                    INSERT INTO bill_items_new2(bill_id, item_type, school, color, size, unit_price, qty, line_total, origin)
                        SELECT bill_id, item_type, school, color, size, unit_price, qty, line_total, COALESCE(origin,'STOCK')
                        FROM bill_items;
                    DROP TABLE bill_items;
                    ALTER TABLE bill_items_new2 RENAME TO bill_items;
                    CREATE INDEX IF NOT EXISTS idx_bill_items_bill ON bill_items(bill_id);
                """)
        except Exception:
            pass

        # Migrate reservation statuses from English to Arabic
        try:
            self.conn.execute("UPDATE reservations SET status='معلق' WHERE status='PENDING'")
            self.conn.execute("UPDATE reservations SET status='تم التسليم' WHERE status='COMPLETED'")
            self.conn.commit()
        except Exception:
            pass

        # ------------------- Sync layer (Phase 1) -------------------
        # Additive: creates sync_outbox/inbox/state/device_identity tables
        # and backfills a `uuid` column on every syncable domain table.
        # Failures here must never break the main app.
        try:
            from sync_core import apply_sync_migration, ensure_device_identity
            apply_sync_migration(self.conn)
            ensure_device_identity(
                self.conn,
                default_name="POS-UNCONFIGURED",
                default_role="pos",
            )
            # Performance indexes for common filtered reads (safe/additive).
            self.conn.executescript("""
                CREATE INDEX IF NOT EXISTS idx_stocks_specs_norm
                    ON stocks(LOWER(TRIM(item_type)), LOWER(TRIM(school)), LOWER(TRIM(color)), LOWER(TRIM(size)));
                CREATE INDEX IF NOT EXISTS idx_bills_customer_norm
                    ON bills(LOWER(TRIM(customer)));
                CREATE INDEX IF NOT EXISTS idx_movements_bill_id
                    ON movements(bill_id);
                CREATE INDEX IF NOT EXISTS idx_bill_items_specs
                    ON bill_items(item_type, school, color, size);
                CREATE INDEX IF NOT EXISTS idx_reservations_specs_norm
                    ON reservations(LOWER(TRIM(item_type)), LOWER(TRIM(school)), LOWER(TRIM(color)), LOWER(TRIM(size)), status);
                CREATE INDEX IF NOT EXISTS idx_sync_outbox_status_seq
                    ON sync_outbox(status, local_seq);
                CREATE INDEX IF NOT EXISTS idx_sync_inbox_status_seq
                    ON sync_inbox(apply_status, server_seq);
            """)
        except Exception:
            import traceback
            traceback.print_exc()

        cur.close()

    # ------------------- Size Profiles -------------------

    def get_size_profile(self, item_type: str, school: str, color: str):
        cur = self.conn.execute(
            """
            SELECT
                num_start_1,
                num_end_1,
                num_start_2,
                num_end_2,
                has_alpha
            FROM size_profiles
            WHERE LOWER(TRIM(item_type)) = LOWER(TRIM(?))
              AND LOWER(TRIM(school)) = LOWER(TRIM(?))
              AND LOWER(TRIM(color)) = LOWER(TRIM(?))
            """,
            (item_type, school, color),
        )
        row = cur.fetchone()
        if row is None:
            return None
        return (
            row["num_start_1"],
            row["num_end_1"],
            row["num_start_2"],
            row["num_end_2"],
            row["has_alpha"],
        )

    def _get_size_ranges(self, item_type, school, color):
        profile = self.get_size_profile(item_type, school, color)
        if not profile:
            return []

        r1s, r1e, r2s, r2e, has_alpha = profile
        ranges = []

        if r1s is not None and r1e is not None:
            ranges.append({
                "range_type": "NUMERIC",
                "start": r1s,
                "end": r1e,
            })

        if r2s is not None and r2e is not None:
            ranges.append({
                "range_type": "NUMERIC",
                "start": r2s,
                "end": r2e,
            })

        if has_alpha:
            ranges.append({
                "range_type": "ALPHA",
                "alpha_set": ALPHA_SIZES[:],
            })

        return ranges

    def upsert_size_profile(
        self,
        item_type: str,
        school: str,
        color: str,
        *,
        r1_start: Optional[int],
        r1_end: Optional[int],
        r2_start: Optional[int],
        r2_end: Optional[int],
        has_alpha: bool,
    ) -> None:
        """
        Insert or update a size profile for (item_type, school, color).
        """
        self.conn.execute(
            """
            INSERT INTO size_profiles (
                item_type, school, color,
                num_start_1, num_end_1,
                num_start_2, num_end_2,
                has_alpha, updated_at
            )
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, datetime('now'))
            ON CONFLICT(item_type, school, color)
            DO UPDATE SET
                num_start_1 = excluded.num_start_1,
                num_end_1   = excluded.num_end_1,
                num_start_2 = excluded.num_start_2,
                num_end_2   = excluded.num_end_2,
                has_alpha   = excluded.has_alpha,
                updated_at  = datetime('now')
            """,
            (
                item_type, school, color,
                r1_start, r1_end,
                r2_start, r2_end,
                int(has_alpha),
            )
        )

        self.conn.commit()

    def get_effective_price(self, item_type, school, color, size):
        """
        Priority:
        1) exact last price (history)
        2) default price for item
        3) None if unknown
        """
        try:
            p = self.last_price_for_specs(item_type, school, color, size)
            if p is not None:
                return float(p)
        except Exception:
            pass

        try:
            cur = self.conn.execute(
                "SELECT default_price FROM item_defaults WHERE item_type=?",
                (item_type,)
            )
            row = cur.fetchone()
            if row and row["default_price"] is not None:
                return float(row["default_price"])
        except Exception:
            pass

        return None

    def ensure_default_price(self, item_type, price):
        """
        Save default price once (only if not exists).
        """
        self.conn.execute(
            """
            INSERT INTO item_defaults (item_type, default_price)
            VALUES (?, ?)
            ON CONFLICT(item_type) DO NOTHING
            """,
            (item_type, float(price))
        )
        self.conn.commit()

    def purge_definition(self, field: str, value: str) -> int:
        if field not in ("item_type", "school", "color", "size"):
            raise ValueError("حقل غير مدعوم")

        value = (value or "").strip()
        if not value:
            return 0

        with self.conn:
            cur = self.conn.execute(
                f"SELECT COUNT(*) FROM stocks WHERE {field} = ?",
                (value,),
            )
            count = int(cur.fetchone()[0] or 0)

            self.conn.execute(
                f"DELETE FROM stocks WHERE {field} = ?",
                (value,),
            )

            self.conn.execute(
                "DELETE FROM spec_history WHERE field = ? AND value = ?",
                (field, value),
            )

        return count

    def list_schools_all(self) -> List[str]:
        cur = self.conn.cursor()
        try:
            cur.execute("""
                SELECT DISTINCT TRIM(school) AS s FROM stocks
                UNION
                SELECT DISTINCT TRIM(school) AS s FROM bill_items
                UNION
                SELECT DISTINCT TRIM(value) AS s FROM spec_history WHERE field='school'
            """)
            vals = [r["s"] for r in cur.fetchall() if r["s"]]
            vals.sort(key=lambda v: v.lower())
            return vals
        finally:
            cur.close()

    def list_items_for_school(self, school: str) -> List[Tuple[str, str]]:
        sc = (school or "").strip()
        if not sc:
            return []
        cur = self.conn.cursor()
        try:
            cur.execute("""
                SELECT DISTINCT TRIM(item_type) AS it, TRIM(color) AS cl
                FROM stocks WHERE LOWER(TRIM(school)) = LOWER(TRIM(?))
                UNION
                SELECT DISTINCT TRIM(item_type) AS it, TRIM(color) AS cl
                FROM bill_items WHERE LOWER(TRIM(school)) = LOWER(TRIM(?))
            """, (sc, sc))
            rows = [(r["it"], r["cl"]) for r in cur.fetchall() if r["it"] and r["cl"]]
            rows.sort(key=lambda t: (t[0].lower(), t[1].lower()))
            return rows
        finally:
            cur.close()

    def _size_row(self, school, item_type, color, size):
        """Return dict with size, current count, and last price for a specific size."""
        cur = self.conn.execute(
            "SELECT COALESCE(SUM(count),0) AS c FROM stocks WHERE LOWER(TRIM(item_type))=LOWER(?) AND LOWER(TRIM(school))=LOWER(?) AND LOWER(TRIM(color))=LOWER(?) AND LOWER(TRIM(size))=LOWER(?)",
            (item_type, school, color, size),
        )
        count = int(cur.fetchone()["c"] or 0)
        last_price = self.last_price_for_specs(item_type, school, color, size)
        return {"size": size, "count": count, "last_price": last_price}

    def list_sizes_for_item(self, school, item_type, color):
        ranges = self._get_size_ranges(item_type, school, color)
        out = []

        for r in ranges:
            if r["range_type"] == "NUMERIC":
                for sz in range(r["start"], r["end"] + 1):
                    out.append(self._size_row(school, item_type, color, str(sz)))

            elif r["range_type"] == "ALPHA":
                for sz in r["alpha_set"]:
                    out.append(self._size_row(school, item_type, color, sz))

        return out

    def last_price_for_specs(
        self,
        item_type: str,
        school: str,
        color: str,
        size: str,
    ) -> Optional[float]:
        it = (item_type or "").strip()
        sc = (school or "").strip()
        cl = (color or "").strip()
        sz_raw = (size or "").strip()

        if not (it and sc and cl and sz_raw):
            return None

        sz_norm = _normalize_size_label(sz_raw)
        cur = self.conn.cursor()

        try:
            if sz_norm and sz_norm != sz_raw:
                cur.execute(
                    """
                    SELECT unit_price
                    FROM bill_items
                    WHERE LOWER(TRIM(item_type)) = LOWER(?)
                    AND LOWER(TRIM(school))    = LOWER(?)
                    AND LOWER(TRIM(color))     = LOWER(?)
                    AND (LOWER(TRIM(size))     = LOWER(?) OR LOWER(TRIM(size)) = LOWER(?))
                    ORDER BY id DESC
                    LIMIT 1
                    """,
                    (it, sc, cl, sz_raw, sz_norm),
                )
            else:
                cur.execute(
                    """
                    SELECT unit_price
                    FROM bill_items
                    WHERE LOWER(TRIM(item_type)) = LOWER(?)
                    AND LOWER(TRIM(school))    = LOWER(?)
                    AND LOWER(TRIM(color))     = LOWER(?)
                    AND LOWER(TRIM(size))      = LOWER(?)
                    ORDER BY id DESC
                    LIMIT 1
                    """,
                    (it, sc, cl, sz_raw),
                )

            r = cur.fetchone()
            if r and r[0] is not None:
                return float(r[0])

            if sz_norm and sz_norm != sz_raw:
                cur.execute(
                    """
                    SELECT unit_price
                    FROM stocks
                    WHERE LOWER(TRIM(item_type)) = LOWER(?)
                    AND LOWER(TRIM(school))    = LOWER(?)
                    AND LOWER(TRIM(color))     = LOWER(?)
                    AND (LOWER(TRIM(size))     = LOWER(?) OR LOWER(TRIM(size)) = LOWER(?))
                    ORDER BY id DESC
                    LIMIT 1
                    """,
                    (it, sc, cl, sz_raw, sz_norm),
                )
            else:
                cur.execute(
                    """
                    SELECT unit_price
                    FROM stocks
                    WHERE LOWER(TRIM(item_type)) = LOWER(?)
                    AND LOWER(TRIM(school))    = LOWER(?)
                    AND LOWER(TRIM(color))     = LOWER(?)
                    AND LOWER(TRIM(size))      = LOWER(?)
                    ORDER BY id DESC
                    LIMIT 1
                    """,
                    (it, sc, cl, sz_raw),
                )

            r = cur.fetchone()
            if r and r[0] is not None:
                return float(r[0])

            return None

        finally:
            cur.close()

    def update_specs_by_ids(
        self,
        ids: Sequence[int],
        *,
        item_type: Optional[str] = None,
        school: Optional[str] = None,
        color: Optional[str] = None,
        size: Optional[str] = None,
        note: str = "Edit specs (by ids)",
    ) -> int:
        ids = [int(x) for x in ids if str(x).strip()]
        if not ids:
            return 0

        sets = []
        args: List[Any] = []
        changes = {}

        def _norm(v: Optional[str]) -> Optional[str]:
            v = (v or "").strip()
            return v if v else None

        it = _norm(item_type)
        sc = _norm(school)
        cl = _norm(color)
        sz = _norm(size)

        if it is not None:
            sets.append("item_type = ?"); args.append(it); changes["item_type"] = it
        if sc is not None:
            sets.append("school = ?"); args.append(sc); changes["school"] = sc
        if cl is not None:
            sets.append("color = ?"); args.append(cl); changes["color"] = cl
        if sz is not None:
            sets.append("size = ?"); args.append(sz); changes["size"] = sz

        if not sets:
            return 0

        ph = ",".join("?" for _ in ids)
        with self.conn:
            cur = self.conn.execute(f"SELECT COUNT(*) AS c FROM stocks WHERE id IN ({ph})", ids)
            count = int(cur.fetchone()["c"] or 0)
            if count == 0:
                return 0

            old_specs = self.conn.execute(
                f"SELECT DISTINCT item_type, school, color, size FROM stocks WHERE id IN ({ph})",
                ids,
            ).fetchall()

            self.conn.execute(
                f"UPDATE stocks SET {', '.join(sets)} WHERE id IN ({ph})",
                (*args, *ids),
            )

            self._cascade_spec_rename(old_specs, changes)

            self._upsert_history(changes)
        self.cleanup_unused_specs()
        return count

    def _cascade_spec_rename(
        self,
        old_specs: Sequence[Any],
        changes: Dict[str, Any],
    ) -> None:
        """Propagate a completed spec rename from `stocks` to historical tables.

        For each old (item_type, school, color, size) tuple we just rewrote,
        if no `stocks` row still references that tuple, it was a full rename —
        cascade the same field changes to `movements`, `bill_items`, and
        `reservations` so audit/history views stay consistent.
        """
        if not changes or not old_specs:
            return
        sets_parts: List[str] = []
        new_vals: List[Any] = []
        for fld in ("item_type", "school", "color", "size"):
            if fld in changes:
                sets_parts.append(f"{fld} = ?")
                new_vals.append(changes[fld])
        if not sets_parts:
            return
        set_sql = ", ".join(sets_parts)

        for row in old_specs:
            old_it = row["item_type"] if isinstance(row, sqlite3.Row) else row[0]
            old_sc = row["school"]    if isinstance(row, sqlite3.Row) else row[1]
            old_cl = row["color"]     if isinstance(row, sqlite3.Row) else row[2]
            old_sz = row["size"]      if isinstance(row, sqlite3.Row) else row[3]

            still = self.conn.execute(
                "SELECT 1 FROM stocks WHERE item_type=? AND school=? AND color=? AND size=? LIMIT 1",
                (old_it, old_sc, old_cl, old_sz),
            ).fetchone()
            if still:
                continue

            where_sql = "item_type=? AND school=? AND color=? AND size=?"
            where_args = (old_it, old_sc, old_cl, old_sz)

            for tbl in ("movements", "bill_items", "reservations"):
                try:
                    self.conn.execute(
                        f"UPDATE {tbl} SET {set_sql} WHERE {where_sql}",
                        (*new_vals, *where_args),
                    )
                except sqlite3.OperationalError:
                    pass

    def get_distinct_filtered(self, target: str, constraints: Dict[str, Any]) -> List[str]:
        valid = {"item_type", "school", "color", "size"}
        if target not in valid:
            return []

        where: List[str] = ["1=1"]
        args: List[Any] = []

        def _is_list(x): return isinstance(x, (list, tuple))

        for fld in ("item_type", "school", "color", "size"):
            if fld == target:
                continue
            v = constraints.get(fld)
            if v in (None, ""):
                continue
            if _is_list(v) and v:
                ph = ",".join("?" for _ in v)
                where.append(f"LOWER({fld}) IN ({ph})")
                args += [str(s).strip().lower() for s in v]
            elif isinstance(v, str) and v.strip():
                where.append(f"LOWER({fld}) = LOWER(?)")
                args.append(v.strip())

        cur = self.conn.cursor()
        try:
            cur.execute(
                f"SELECT DISTINCT TRIM({target}) AS v FROM stocks WHERE {' AND '.join(where)}"
                f" ORDER BY LOWER(TRIM({target})) ASC",
                args,
            )

            return [r["v"] for r in cur.fetchall() if r["v"] not in (None, "")]
        finally:
            cur.close()

    def last_color_for_school(self, school: str) -> Optional[str]:
        school = (school or "").strip()
        if not school:
            return None
        cur = self.conn.cursor()
        try:
            cur.execute(
                "SELECT color FROM stocks WHERE LOWER(school)=LOWER(?) AND color<>'' ORDER BY id DESC LIMIT 1",
                (school,)
            )
            r = cur.fetchone()
            if r and r["color"]:
                return str(r["color"])
            cur.execute(
                "SELECT color FROM bill_items WHERE LOWER(school)=LOWER(?) AND color<>'' ORDER BY id DESC LIMIT 1",
                (school,)
            )
            r = cur.fetchone()
            if r and r["color"]:
                return str(r["color"])
            return None
        finally:
            cur.close()

    def update_prices(self, filters: Dict[str, Any], new_price: float, note: str = "Price update") -> int:
        if not self.is_manager_feature_enabled("allow_inventory_price_edit"):
            raise PermissionError(_feature_restricted_message("تعديل الأسعار من نقطة البيع غير مسموح به حالياً."))
        if new_price is None or float(new_price) < 0:
            raise ValueError("New price must be a non-negative number.")
        where: List[str] = []
        args: List[Any] = []

        if "id" in filters and filters["id"]:
            where.append("id = ?")
            args.append(int(filters["id"]))
        else:
            for k in ("item_type", "school", "color", "size"):
                v_raw = (filters.get(k) or "").strip()
                if not v_raw:
                    continue
                if k == "size":
                    v_norm = _normalize_size_label(v_raw)
                    if v_norm and v_norm != v_raw:
                        where.append("(LOWER(TRIM(size)) = LOWER(?) OR LOWER(TRIM(size)) = LOWER(?))")
                        args.extend([v_raw, v_norm])
                    else:
                        where.append("LOWER(TRIM(size)) = LOWER(?)")
                        args.append(v_raw)
                else:
                    where.append(f"LOWER(TRIM({k})) = LOWER(?)")
                    args.append(v_raw)

        if not where:
            raise ValueError("No filter provided for price update.")

        cur = self.conn.cursor()
        cur.execute(f"SELECT * FROM stocks WHERE {' AND '.join(where)}", args)
        rows = cur.fetchall()
        if not rows:
            cur.close()
            return 0

        with self.conn:
            self.conn.execute(
                f"UPDATE stocks SET unit_price = ? WHERE {' AND '.join(where)}",
                (float(new_price), *args),
            )
            ts = now_iso()
            for r in rows:
                self.conn.execute(
                    """INSERT INTO movements
                    (ts,direction,stock_id,qty,note,bill_id,item_type,school,color,size,unit_price)
                    VALUES(?,?,?,?,?,?,?,?,?,?,?)""",
                    (
                        ts, "PRICE_UPDATE", int(r["id"]), 0, note, None,
                        r["item_type"], r["school"], r["color"], r["size"],
                        float(new_price),
                    ),
                )
            self._record_sync_event("PRICE_UPDATE", {
                "mode":    "fixed",
                "new_price": float(new_price),
                "filters": {k: v for k, v in (filters or {}).items() if v not in (None, "", [])},
                "updated_count": len(rows),
                "scope":   "filter",
                "note":    note,
            })
        cur.close()
        return len(rows)

    def _migrate_from_json_if_empty(self, legacy_path: str) -> None:
        cur = self.conn.cursor()
        cur.execute("SELECT COUNT(*) FROM stocks")
        has_data = cur.fetchone()[0] > 0
        cur.execute("SELECT COUNT(*) FROM bills")
        has_data = has_data or cur.fetchone()[0] > 0
        cur.close()
        if has_data or not os.path.exists(legacy_path):
            return
        try:
            with open(legacy_path, "r", encoding="utf-8") as f:
                data = json.load(f)
            with self.conn:
                for s in data.get("stocks", []):
                    self.conn.execute(
                        """INSERT INTO stocks(id,item_type,school,color,size,unit_price,count)
                           VALUES(?,?,?,?,?,?,?)""",
                        (
                            int(s["id"]), s["item_type"], s["school"], s["color"], s["size"],
                            float(s["unit_price"]), int(s["count"]),
                        ),
                    )
                for m in data.get("movements", []):
                    self.conn.execute(
                        """INSERT INTO movements
                           (id,ts,direction,stock_id,qty,note,bill_id,item_type,school,color,size,unit_price)
                           VALUES(?,?,?,?,?,?,?,?,?,?,?,?)""",
                        (
                            int(m["id"]), m["ts"], m["direction"], m.get("stock_id"), int(m["qty"]),
                            m.get("note"), m.get("bill_id"), m.get("item_type"), m.get("school"),
                            m.get("color"), m.get("size"),
                            float(m.get("unit_price")) if m.get("unit_price") is not None else None,
                        ),
                    )
                for b in data.get("bills", []):
                    self.conn.execute(
                        "INSERT INTO bills(id,created_at,customer,total) VALUES(?,?,?,?)",
                        (int(b["id"]), b["created_at"], b.get("customer"), float(b["total"])),
                    )
                for bi in data.get("bill_items", []):
                    self.conn.execute(
                        """INSERT INTO bill_items
                           (id,bill_id,item_type,school,color,size,unit_price,qty,line_total)
                           VALUES(?,?,?,?,?,?,?,?,?)""",
                        (
                            int(bi["id"]), int(bi["bill_id"]), bi["item_type"], bi["school"], bi["color"], bi["size"],
                            float(bi["unit_price"]),
                            int(bi["qty"]), float(bi["line_total"]),
                        ),
                    )
                sh = data.get("spec_history", {})
                for field in ("item_type", "school", "color", "size"):
                    for val in sh.get(field, []):
                        self.conn.execute(
                            "INSERT OR IGNORE INTO spec_history(field,value) VALUES(?,?)", (field, str(val).strip())
                        )
            try:
                os.rename(legacy_path, legacy_path + ".migrated_backup")
            except Exception:
                pass
        except Exception:
            pass

    # -------- Helpers --------
    def list_customers(self) -> List[str]:
        cur = self.conn.cursor()
        cur.execute("""
            SELECT DISTINCT TRIM(customer) AS c
            FROM bills
            WHERE customer IS NOT NULL AND TRIM(customer) <> ''
            ORDER BY LOWER(c)
        """)
        vals = [r["c"] for r in cur.fetchall() if r["c"]]
        cur.close()
        return vals

    def _upsert_history(self, specs: Dict[str, Any]) -> None:
        with self.conn:
            for field in ("item_type", "school", "color", "size"):
                val = str(specs.get(field, "")).strip()
                if val:
                    self.conn.execute(
                        "INSERT OR IGNORE INTO spec_history(field,value) VALUES(?,?)", (field, val)
                    )

    def get_distinct(self, field: str) -> List[str]:
        if field not in ("item_type", "school", "color", "size"):
            return []

        cur = self.conn.cursor()
        try:
            cur.execute(f"""
                SELECT DISTINCT TRIM({field}) AS v
                FROM stocks
                WHERE count > 0
                ORDER BY LOWER(v)
            """)
            return [r["v"] for r in cur.fetchall() if r["v"]]
        finally:
            cur.close()

    # -------- Stocks --------
    def add_stock(
        self,
        item_type: str,
        school: str,
        color: str,
        size: str,
        unit_price: Optional[float],
        count: int,
    ) -> int:
        self._require_shift()

        if count <= 0:
            raise ValueError("Count must be > 0")

        price: Optional[float] = None
        user_provided_price = unit_price is not None and str(unit_price) != ""

        if user_provided_price:
            price = float(unit_price)
            if price < 0:
                raise ValueError("Price must be >= 0")
        else:
            price = self.get_effective_price(item_type, school, color, size)

        if price is None:
            raise ValueError("يجب إدخال السعر مرة واحدة على الأقل لهذا الصنف.")

        if user_provided_price:
            self.ensure_default_price(item_type, price)

        with self.conn:
            cur = self.conn.execute(
                """INSERT INTO stocks
                (item_type,school,color,size,unit_price,count)
                VALUES(?,?,?,?,?,?)""",
                (
                    item_type.strip(),
                    school.strip(),
                    color.strip(),
                    size.strip(),
                    float(price),
                    int(count),
                ),
            )

            stock_id = cur.lastrowid

            self.conn.execute(
                """INSERT INTO movements
                (ts,direction,stock_id,qty,note,bill_id,
                    item_type,school,color,size,unit_price)
                VALUES(?,?,?,?,?,?,?,?,?,?,?)""",
                (
                    now_iso(),
                    "IN",
                    stock_id,
                    int(count),
                    "Income add",
                    None,
                    item_type.strip(),
                    school.strip(),
                    color.strip(),
                    size.strip(),
                    float(price),
                ),
            )

            self._upsert_history(
                {
                    "item_type": item_type,
                    "school": school,
                    "color": color,
                    "size": size,
                }
            )

        return int(stock_id)

    def cleanup_unused_specs(self):
        cur = self.conn.cursor()
        try:
            for field in ("item_type", "school", "color"):
                cur.execute(f"""
                    DELETE FROM spec_history
                    WHERE field = ?
                    AND value NOT IN (
                        SELECT DISTINCT TRIM({field}) FROM stocks
                        UNION
                        SELECT DISTINCT TRIM({field}) FROM bill_items
                    )
                """, (field,))
            self.conn.commit()
        finally:
            cur.close()

    def _filters_where(self, filters: Dict[str, Any], prefix: str = "") -> Tuple[str, List[Any]]:
        """Include count = 0 stock rows so schools/items/colors stay visible when out of stock."""
        where: List[str] = ["1=1"]
        args: List[Any] = []

        candidate_keys_txt = ("item_type", "school", "color", "size")
        all_keys = candidate_keys_txt

        multi_keys = [k for k in all_keys if isinstance(filters.get(k), list) and len(filters.get(k) or []) > 0]
        if len(multi_keys) > 1:
            raise ValueError("يمكن اختيار أكثر من قيمة لحقل واحد فقط، وليس لأكثر من حقل.")

        for k in candidate_keys_txt:
            v = filters.get(k)
            if v and not isinstance(v, list):
                v = str(v).strip()
                if v:
                    where.append(f"LOWER(TRIM({prefix}{k})) = LOWER(?)")
                    args.append(v)

        if multi_keys:
            mk = multi_keys[0]
            vals = [x for x in (filters.get(mk) or []) if x not in (None, "")]
            if vals:
                placeholders = ",".join(["?"] * len(vals))
                where.append(f"LOWER({prefix}{mk}) IN ({placeholders})")
                args.extend([str(x).strip().lower() for x in vals])

        return (" AND ".join(where)) if where else "1=1", args

    def search_stocks(self, filters: Dict[str, Any]) -> List[StockRow]:
        where, args = self._filters_where(filters)
        cur = self.conn.cursor()
        cur.execute(
            f"""SELECT id,item_type,school,color,size,unit_price,count
                FROM stocks
                WHERE {where}
                ORDER BY id ASC""",
            args,
        )
        rows = [
            StockRow(
                id=r["id"], item_type=r["item_type"], school=r["school"], color=r["color"],
                size=r["size"],
                unit_price=r["unit_price"], count=r["count"],
            )
            for r in cur.fetchall()
        ]
        cur.close()
        return rows

    def current_inventory(self, filters: Dict[str, Any]) -> List[Dict[str, Any]]:
        where, args = self._filters_where(filters, prefix="s.")

        cur = self.conn.cursor()
        cur.execute(
            f"""
            SELECT
                MIN(s.id)              AS id,
                s.item_type,
                s.school,
                s.color,
                s.size,
                s.unit_price,
                SUM(s.count)           AS count,
                SUM(s.count * s.unit_price) AS value
            FROM stocks s
            WHERE {where}
            GROUP BY
                s.item_type,
                s.school,
                s.color,
                s.size,
                s.unit_price
            ORDER BY
                s.item_type,
                s.school,
                s.color,

                CASE
                    WHEN TRIM(s.size) NOT GLOB '*[^0-9]*' THEN 0
                    ELSE 1
                END,

                CASE
                    WHEN TRIM(s.size) NOT GLOB '*[^0-9]*'
                    THEN CAST(s.size AS INTEGER)
                    ELSE NULL
                END,

                CASE UPPER(TRIM(s.size))
                    WHEN 'XXS' THEN 1
                    WHEN 'XS'  THEN 2
                    WHEN 'S'   THEN 3
                    WHEN 'M'   THEN 4
                    WHEN 'L'   THEN 5
                    WHEN 'XL'  THEN 6
                    WHEN '2XL' THEN 7
                    WHEN '3XL' THEN 8
                    WHEN '4XL' THEN 9
                    WHEN '5XL' THEN 10
                    ELSE 99
                END
            """,
            args,
        )

        rows = []
        for r in cur.fetchall():
            rows.append(
                dict(
                    id=r["id"],
                    item_type=r["item_type"],
                    school=r["school"],
                    color=r["color"],
                    size=r["size"],
                    unit_price=r["unit_price"],
                    count=int(r["count"]),
                    value=float(r["value"]),
                )
            )

        cur.close()
        return rows

    # -------- Billing --------
    def create_bill(self, customer: str, bill_lines: List[Dict[str, Any]]) -> int:
        self._require_shift()
        if not bill_lines:
            raise ValueError("Bill has no items")

        def _candidates_for_line(line: Dict[str, Any]) -> List[sqlite3.Row]:
            cur = self.conn.cursor()
            try:
                if line.get("stock_id"):
                    cur.execute(
                        "SELECT * FROM stocks WHERE id=? AND count>0 ORDER BY id ASC",
                        (int(line["stock_id"]),),
                    )
                else:
                    where_parts = ["count > 0"]
                    args: List[Any] = []
                    for k in ("item_type", "school", "color", "size"):
                        v = (line.get(k) or "").strip()
                        if v:
                            where_parts.append(f"LOWER(TRIM({k})) = LOWER(?)")
                            args.append(v)
                    cur.execute(f"SELECT * FROM stocks WHERE {' AND '.join(where_parts)} ORDER BY id ASC", args)
                return cur.fetchall()
            finally:
                cur.close()

        with self.conn:
            bill_cur = self.conn.execute(
                "INSERT INTO bills(created_at,customer,total,bill_type,status) VALUES(?,?,?,?,?)",
                (now_iso(), (customer or "").strip() or None, 0.0, "SALE", "CONFIRMED"),
            )
            bill_id = int(bill_cur.lastrowid)
            total = 0.0

            for line in bill_lines:
                try:
                    qty_needed = int(line["qty"])
                except Exception:
                    raise ValueError("Qty must be > 0")
                if qty_needed <= 0:
                    raise ValueError("Qty must be > 0")

                allow_factory = bool(line.get("allow_factory_fill"))

                if allow_factory and not line.get("stock_id"):
                    cands = []
                else:
                    cands = _candidates_for_line(line)

                remaining = qty_needed
                chunks: List[Tuple[sqlite3.Row, int]] = []
                for s in cands:
                    if remaining <= 0:
                        break
                    take = min(int(s["count"]), remaining)
                    if take > 0:
                        chunks.append((s, take))
                        remaining -= take

                if remaining > 0 and not allow_factory:
                    raise ValueError(
                        f"لا توجد كمية كافية للصنف "
                        f"{line.get('item_type','(النوع?)')} / {line.get('school','(المدرسة?)')} / {line.get('size','(المقاس?)')} "
                        f"(المطلوب {qty_needed}، المتاح {qty_needed-remaining})."
                    )

                taken = qty_needed - remaining

                price_base: Optional[float] = None
                if "unit_price" in line and line["unit_price"] is not None and str(line["unit_price"]) != "":
                    price_base = float(line["unit_price"])
                elif taken > 0:
                    value_sum = sum(float(s["unit_price"]) * take for s, take in chunks)
                    qty_sum = sum(take for _, take in chunks)
                    if qty_sum > 0:
                        price_base = round(value_sum / qty_sum, 4)
                if price_base is None:
                    price_base = 0.0

                if taken > 0:
                    for s, take in chunks:
                        stock_price = float(price_base)
                        line_total_chunk = stock_price * int(take)
                        self.conn.execute(
                            """INSERT INTO bill_items
                            (bill_id,item_type,school,color,size,unit_price,qty,line_total,origin)
                            VALUES(?,?,?,?,?,?,?,?,?)""",
                            (
                                bill_id,
                                line["item_type"], line["school"], line["color"], line["size"],
                                stock_price, int(take), line_total_chunk, "STOCK",
                            ),
                        )
                        total += line_total_chunk

                if remaining > 0 and allow_factory:
                    self._upsert_history({
                        "item_type": line.get("item_type"),
                        "school":    line.get("school"),
                        "color":     line.get("color"),
                        "size":      line.get("size"),
                    })

                    factory_price = float(price_base)
                    factory_line_total = factory_price * remaining
                    self.conn.execute(
                        """INSERT INTO bill_items
                        (bill_id,item_type,school,color,size,unit_price,qty,line_total,origin)
                        VALUES(?,?,?,?,?,?,?,?,?)""",
                        (
                            bill_id,
                            line["item_type"], line["school"], line["color"], line["size"],
                            factory_price, int(remaining), factory_line_total, "FACTORY",
                        ),
                    )
                    total += factory_line_total

                for s, take in chunks:
                    self.conn.execute("UPDATE stocks SET count = count - ? WHERE id = ?", (int(take), int(s["id"])))
                    self.conn.execute(
                        """INSERT INTO movements
                        (ts,direction,stock_id,qty,note,bill_id,item_type,school,color,size,unit_price)
                        VALUES(?,?,?,?,?,?,?,?,?,?,?)""",
                        (
                            now_iso(), "OUT", int(s["id"]), int(take), "Bill deduction", bill_id,
                            s["item_type"], s["school"], s["color"], s["size"],
                            float(s["unit_price"]),
                        ),
                    )

                if remaining > 0 and allow_factory:
                    self.conn.execute(
                        """INSERT INTO movements
                        (ts,direction,stock_id,qty,note,bill_id,item_type,school,color,size,unit_price)
                        VALUES(?,?,?,?,?,?,?,?,?,?,?)""",
                        (
                            now_iso(), "OUT_FACTORY", None, int(remaining), "Factory direct", bill_id,
                            line["item_type"], line["school"], line["color"], line["size"],
                            float(price_base),
                        ),
                    )

            self.cleanup_unused_specs()
            self.conn.execute("UPDATE bills SET total=? WHERE id=?", (float(total), bill_id))
            wh_tgt = _extract_warehouse_target(customer)
            if wh_tgt:
                self.conn.execute(
                    "UPDATE bills SET bill_type=?, total=? WHERE id=?",
                    (WAREHOUSE_RETURN_BILL_TYPE, 0.0, bill_id),
                )

            bill_uuid_row = self.conn.execute(
                "SELECT uuid FROM bills WHERE id = ?", (bill_id,)
            ).fetchone()
            bill_uuid = bill_uuid_row[0] if bill_uuid_row else None
            items_payload = [
                {
                    "item_type": r[0], "school": r[1], "color": r[2], "size": r[3],
                    "unit_price": float(r[4]), "qty": int(r[5]),
                    "line_total": float(r[6]), "origin": r[7],
                }
                for r in self.conn.execute(
                    "SELECT item_type,school,color,size,unit_price,qty,line_total,origin "
                    "FROM bill_items WHERE bill_id = ?", (bill_id,)
                ).fetchall()
            ]
            warehouse_target = _extract_warehouse_target(customer)
            branch_target = _extract_branch_target(customer)
            source_device = (self._current_device_name() or "").strip()
            if branch_target and branch_target == source_device:
                branch_target = None
            if warehouse_target:
                self._record_warehouse_return_event(
                    return_uuid=bill_uuid or "",
                    note=f"bill #{bill_id}",
                    lines=[
                        {
                            "item_type": it["item_type"],
                            "school": it["school"],
                            "color": it["color"],
                            "size": it["size"],
                            "unit_price": float(it["unit_price"]),
                            "qty": int(it["qty"]),
                        }
                        for it in items_payload if int(it["qty"]) > 0
                    ],
                )
            elif branch_target:
                self._record_transfer_via_warehouse_event(
                    request_uuid=bill_uuid or "",
                    target_device=branch_target,
                    note=f"bill #{bill_id}",
                    lines=[
                        {
                            "item_type": it["item_type"],
                            "school": it["school"],
                            "color": it["color"],
                            "size": it["size"],
                            "unit_price": float(it["unit_price"]),
                            "qty": int(it["qty"]),
                        }
                        for it in items_payload if int(it["qty"]) > 0
                    ],
                )
            else:
                self._record_sync_event("SALE_CREATED", {
                    "bill_uuid": bill_uuid,
                    "bill_id": bill_id,
                    "customer": (customer or "").strip() or None,
                    "total": float(total),
                    "items": items_payload,
                    "shift_id": self.active_shift_id,
                })

            return bill_id

    # -------- Return Bill Methods --------
    def create_return_bill(self, customer: str, return_lines: List[Dict[str, Any]]) -> int:
        """Create a return bill – adds items back to stock."""
        self._require_shift()
        if not return_lines:
            raise ValueError("لا توجد أصناف في فاتورة المرتجع")

        with self.conn:
            bill_cur = self.conn.execute(
                "INSERT INTO bills(created_at,customer,total,bill_type,status) VALUES(?,?,?,?,?)",
                (now_iso(), (customer or "").strip() or None, 0.0, "RETURN", "CONFIRMED"),
            )
            bill_id = int(bill_cur.lastrowid)
            total = 0.0

            for line in return_lines:
                qty = int(line["qty"])
                if qty <= 0:
                    raise ValueError("الكمية يجب أن تكون أكبر من 0")
                price = float(line.get("unit_price") or 0)
                line_total = price * qty
                total += line_total

                # Add stock back
                cur = self.conn.execute(
                    """INSERT INTO stocks(item_type,school,color,size,unit_price,count)
                    VALUES(?,?,?,?,?,?)""",
                    (line["item_type"].strip(), line["school"].strip(),
                     line["color"].strip(), line["size"].strip(),
                     price, qty),
                )
                stock_id = int(cur.lastrowid)

                # Record bill item
                self.conn.execute(
                    """INSERT INTO bill_items(bill_id,item_type,school,color,size,unit_price,qty,line_total,origin)
                    VALUES(?,?,?,?,?,?,?,?,?)""",
                    (bill_id, line["item_type"], line["school"], line["color"],
                     line["size"], price, qty, line_total, "RETURN"),
                )

                # Record movement
                self.conn.execute(
                    """INSERT INTO movements(ts,direction,stock_id,qty,note,bill_id,item_type,school,color,size,unit_price)
                    VALUES(?,?,?,?,?,?,?,?,?,?,?)""",
                    (now_iso(), "RETURN_IN", stock_id, qty,
                     "مرتجع", bill_id,
                     line["item_type"], line["school"], line["color"],
                     line["size"], price),
                )

                self._upsert_history({
                    "item_type": line.get("item_type"),
                    "school": line.get("school"),
                    "color": line.get("color"),
                    "size": line.get("size"),
                })

            self.conn.execute("UPDATE bills SET total=? WHERE id=?", (float(total), bill_id))

            bill_uuid_row = self.conn.execute(
                "SELECT uuid FROM bills WHERE id = ?", (bill_id,)
            ).fetchone()
            self._record_sync_event("SALE_RETURNED", {
                "bill_uuid": bill_uuid_row[0] if bill_uuid_row else None,
                "bill_id": bill_id,
                "customer": (customer or "").strip() or None,
                "total": float(total),
                "lines": [
                    {
                        "item_type": ln.get("item_type"),
                        "school":    ln.get("school"),
                        "color":     ln.get("color"),
                        "size":      ln.get("size"),
                        "unit_price": float(ln.get("unit_price") or 0),
                        "qty":        int(ln["qty"]),
                    }
                    for ln in return_lines
                ],
                "shift_id": self.active_shift_id,
            })
            return bill_id

    def create_exchange_bill(self, customer: str,
                             return_lines: List[Dict[str, Any]],
                             take_lines: List[Dict[str, Any]]) -> int:
        """Create an exchange bill – returns items to stock and takes new items."""
        self._require_shift()
        if not return_lines and not take_lines:
            raise ValueError("لا توجد أصناف في فاتورة الاستبدال")

        def _candidates_for_line(line: Dict[str, Any]) -> List[Any]:
            cur = self.conn.cursor()
            try:
                where_parts = ["count > 0"]
                args: List[Any] = []
                for k in ("item_type", "school", "color", "size"):
                    v = (line.get(k) or "").strip()
                    if v:
                        where_parts.append(f"LOWER(TRIM({k})) = LOWER(?)")
                        args.append(v)
                cur.execute(f"SELECT * FROM stocks WHERE {' AND '.join(where_parts)} ORDER BY id ASC", args)
                return cur.fetchall()
            finally:
                cur.close()

        with self.conn:
            bill_cur = self.conn.execute(
                "INSERT INTO bills(created_at,customer,total,bill_type,status) VALUES(?,?,?,?,?)",
                (now_iso(), (customer or "").strip() or None, 0.0, "EXCHANGE", "CONFIRMED"),
            )
            bill_id = int(bill_cur.lastrowid)
            return_total = 0.0
            take_total = 0.0

            # --- Process returned items (add to stock) ---
            for line in return_lines:
                qty = int(line["qty"])
                if qty <= 0:
                    raise ValueError("الكمية يجب أن تكون أكبر من 0")
                price = float(line.get("unit_price") or 0)
                line_total = price * qty
                return_total += line_total

                cur = self.conn.execute(
                    """INSERT INTO stocks(item_type,school,color,size,unit_price,count)
                    VALUES(?,?,?,?,?,?)""",
                    (line["item_type"].strip(), line["school"].strip(),
                     line["color"].strip(), line["size"].strip(),
                     price, qty),
                )
                stock_id = int(cur.lastrowid)

                self.conn.execute(
                    """INSERT INTO bill_items(bill_id,item_type,school,color,size,unit_price,qty,line_total,origin)
                    VALUES(?,?,?,?,?,?,?,?,?)""",
                    (bill_id, line["item_type"], line["school"], line["color"],
                     line["size"], price, qty, line_total, "RETURN"),
                )
                self.conn.execute(
                    """INSERT INTO movements(ts,direction,stock_id,qty,note,bill_id,item_type,school,color,size,unit_price)
                    VALUES(?,?,?,?,?,?,?,?,?,?,?)""",
                    (now_iso(), "RETURN_IN", stock_id, qty,
                     "استبدال - مرتجع", bill_id,
                     line["item_type"], line["school"], line["color"],
                     line["size"], price),
                )
                self._upsert_history({
                    "item_type": line.get("item_type"),
                    "school": line.get("school"),
                    "color": line.get("color"),
                    "size": line.get("size"),
                })

            # --- Process taken items (deduct from stock) ---
            for line in take_lines:
                qty_needed = int(line["qty"])
                if qty_needed <= 0:
                    raise ValueError("الكمية يجب أن تكون أكبر من 0")

                cands = _candidates_for_line(line)
                remaining = qty_needed
                chunks: List[Tuple[Any, int]] = []
                for s in cands:
                    if remaining <= 0:
                        break
                    take = min(int(s["count"]), remaining)
                    if take > 0:
                        chunks.append((s, take))
                        remaining -= take

                if remaining > 0:
                    raise ValueError(
                        f"لا توجد كمية كافية للصنف "
                        f"{line.get('item_type','?')} / {line.get('school','?')} / {line.get('size','?')} "
                        f"(المطلوب {qty_needed}، المتاح {qty_needed - remaining})."
                    )

                price = float(line.get("unit_price") or 0)

                for s, take_qty in chunks:
                    lt = price * int(take_qty)
                    take_total += lt
                    self.conn.execute(
                        """INSERT INTO bill_items(bill_id,item_type,school,color,size,unit_price,qty,line_total,origin)
                        VALUES(?,?,?,?,?,?,?,?,?)""",
                        (bill_id, line["item_type"], line["school"], line["color"],
                         line["size"], price, int(take_qty), lt, "STOCK"),
                    )
                    self.conn.execute("UPDATE stocks SET count = count - ? WHERE id = ?",
                                     (int(take_qty), int(s["id"])))
                    self.conn.execute(
                        """INSERT INTO movements(ts,direction,stock_id,qty,note,bill_id,item_type,school,color,size,unit_price)
                        VALUES(?,?,?,?,?,?,?,?,?,?,?)""",
                        (now_iso(), "OUT", int(s["id"]), int(take_qty),
                         "استبدال - مأخوذ", bill_id,
                         s["item_type"], s["school"], s["color"],
                         s["size"], float(s["unit_price"])),
                    )

            self.cleanup_unused_specs()
            diff = take_total - return_total
            self.conn.execute("UPDATE bills SET total=? WHERE id=?", (float(diff), bill_id))

            bill_uuid_row = self.conn.execute(
                "SELECT uuid FROM bills WHERE id = ?", (bill_id,)
            ).fetchone()
            self._record_sync_event("SALE_EXCHANGED", {
                "bill_uuid": bill_uuid_row[0] if bill_uuid_row else None,
                "bill_id": bill_id,
                "customer": (customer or "").strip() or None,
                "return_total": float(return_total),
                "take_total": float(take_total),
                "diff": float(diff),
                "return_lines": [
                    {
                        "item_type": ln.get("item_type"),
                        "school":    ln.get("school"),
                        "color":     ln.get("color"),
                        "size":      ln.get("size"),
                        "unit_price": float(ln.get("unit_price") or 0),
                        "qty":        int(ln["qty"]),
                    }
                    for ln in return_lines
                ],
                "take_lines": [
                    {
                        "item_type": ln.get("item_type"),
                        "school":    ln.get("school"),
                        "color":     ln.get("color"),
                        "size":      ln.get("size"),
                        "unit_price": float(ln.get("unit_price") or 0),
                        "qty":        int(ln["qty"]),
                    }
                    for ln in take_lines
                ],
                "shift_id": self.active_shift_id,
            })
            return bill_id

    # -------- New Reservation Methods --------
    def create_reservation(self, customer, lines, paid_amount=0.0):
        """Create reservation records for items."""
        self._require_shift()
        if not lines:
            raise ValueError("لا توجد أصناف للحجز")
        created = []
        with self.conn:
            for line in lines:
                total = float(line["unit_price"]) * int(line["qty"])
                cur = self.conn.execute(
                    """INSERT INTO reservations(created_at,customer,item_type,school,color,size,qty,unit_price,total_amount,paid_amount,status,note)
                    VALUES(?,?,?,?,?,?,?,?,?,?,?,?)""",
                    (now_iso(), (customer or "").strip(), line["item_type"], line["school"],
                     line["color"], line["size"], int(line["qty"]), float(line["unit_price"]),
                     total, float(paid_amount), "معلق", line.get("note", "")),
                )
                created.append(cur.lastrowid)
                self.conn.execute(
                    """INSERT INTO movements(ts,direction,stock_id,qty,note,bill_id,item_type,school,color,size,unit_price)
                    VALUES(?,?,?,?,?,?,?,?,?,?,?)""",
                    (now_iso(), "RESERVE", None, int(line["qty"]), f"حجز - {customer}",
                     None, line["item_type"], line["school"], line["color"], line["size"],
                     float(line["unit_price"])),
                )
            if created:
                uuid_rows = self.conn.execute(
                    "SELECT id, uuid FROM reservations WHERE id IN (%s)"
                    % ",".join("?" * len(created)),
                    tuple(int(x) for x in created),
                ).fetchall()
                id_to_uuid = {int(r[0]): r[1] for r in uuid_rows}
                self._record_sync_event("RESERVATION_CREATED", {
                    "customer": (customer or "").strip() or None,
                    "paid_amount": float(paid_amount),
                    "reservations": [
                        {
                            "reservation_uuid": id_to_uuid.get(int(rid)),
                            "reservation_id":   int(rid),
                            "item_type": ln.get("item_type"),
                            "school":    ln.get("school"),
                            "color":     ln.get("color"),
                            "size":      ln.get("size"),
                            "qty":        int(ln["qty"]),
                            "unit_price": float(ln["unit_price"]),
                        }
                        for rid, ln in zip(created, lines)
                    ],
                    "shift_id": self.active_shift_id,
                })
        return created

    def list_reservations(self, status=None, date_from=None, date_to=None, school=None, item_type=None, color=None):
        where = ["1=1"]
        args = []
        if status:
            where.append("status = ?")
            args.append(status)
        if date_from:
            where.append("date(created_at) >= date(?)")
            args.append(date_from)
        if date_to:
            where.append("date(created_at) <= date(?)")
            args.append(date_to)
        if school:
            where.append("LOWER(TRIM(school)) = LOWER(?)")
            args.append(school.strip())
        if item_type:
            where.append("LOWER(TRIM(item_type)) = LOWER(?)")
            args.append(item_type.strip())
        if color:
            where.append("LOWER(TRIM(color)) = LOWER(?)")
            args.append(color.strip())
        cur = self.conn.execute(
            f"SELECT * FROM reservations WHERE {' AND '.join(where)} ORDER BY id DESC", args)
        return [dict(r) for r in cur.fetchall()]

    def update_reservation_payment(self, res_id, new_paid):
        self.conn.execute(
            "UPDATE reservations SET paid_amount=? WHERE id=?",
            (float(new_paid), int(res_id)))
        row = self.conn.execute(
            "SELECT uuid FROM reservations WHERE id=?", (int(res_id),)
        ).fetchone()
        self._record_sync_event("RESERVATION_PAYMENT_UPDATED", {
            "reservation_uuid": row[0] if row else None,
            "reservation_id":   int(res_id),
            "paid_amount":      float(new_paid),
        })

    def complete_reservation(self, res_id):
        self.conn.execute(
            "UPDATE reservations SET status='تم التسليم' WHERE id=?", (int(res_id),))
        row = self.conn.execute(
            "SELECT uuid FROM reservations WHERE id=?", (int(res_id),)
        ).fetchone()
        self._record_sync_event("RESERVATION_COMPLETED", {
            "reservation_uuid": row[0] if row else None,
            "reservation_id":   int(res_id),
        })

    def deliver_reservation(self, res_id: int, collected_amount: float = 0.0):
        """Mark reservation as delivered, collect remaining payment, record movement."""
        self._require_shift()
        rid = int(res_id)
        cur = self.conn.execute("SELECT * FROM reservations WHERE id=?", (rid,))
        row = cur.fetchone()
        if not row:
            raise ValueError("الحجز غير موجود")
        if row["status"] == "تم التسليم":
            raise ValueError("تم تسليم هذا الحجز مسبقاً")
        new_paid = float(row["paid_amount"]) + float(collected_amount)
        with self.conn:
            self.conn.execute(
                "UPDATE reservations SET status='تم التسليم', paid_amount=? WHERE id=?",
                (new_paid, rid))
            if float(collected_amount) > 1e-9:
                self.conn.execute(
                    """INSERT INTO movements(ts,direction,stock_id,qty,note,bill_id,item_type,school,color,size,unit_price)
                    VALUES(?,?,?,?,?,?,?,?,?,?,?)""",
                    (now_iso(), "DELIVER_PAY", None, 0,
                     f"تحصيل باقي حجز #{rid}",
                     None, row["item_type"], row["school"], row["color"], row["size"],
                     float(collected_amount)))
            self._record_sync_event("RESERVATION_DELIVERED", {
                "reservation_uuid": row["uuid"] if "uuid" in row.keys() else None,
                "reservation_id":   rid,
                "collected_amount": float(collected_amount),
                "paid_amount_total": float(new_paid),
                "shift_id": self.active_shift_id,
            })

    # -------- Shift Management --------
    def start_shift(self) -> int:
        cur = self.conn.execute("SELECT id FROM shifts WHERE status='OPEN' LIMIT 1")
        if cur.fetchone():
            raise ValueError("يوجد وردية مفتوحة بالفعل.")
        with self.conn:
            c = self.conn.execute(
                "INSERT INTO shifts(started_at, status) VALUES(?, 'OPEN')",
                (now_iso(),)
            )
            shift_id = int(c.lastrowid)
            shift_uuid_row = self.conn.execute(
                "SELECT uuid FROM shifts WHERE id=?", (shift_id,)
            ).fetchone()
            self._record_sync_event("SHIFT_OPENED", {
                "shift_uuid": shift_uuid_row[0] if shift_uuid_row else None,
                "shift_id":   shift_id,
                "started_at": now_iso(),
            })
            return shift_id

    def get_open_shift(self) -> Optional[Dict[str, Any]]:
        cur = self.conn.execute("SELECT * FROM shifts WHERE status='OPEN' LIMIT 1")
        row = cur.fetchone()
        return dict(row) if row else None

    def end_shift(self, shift_id: int, summary_json: str = "") -> None:
        with self.conn:
            self.conn.execute(
                "UPDATE shifts SET ended_at=?, status='CLOSED', summary_json=? WHERE id=? AND status='OPEN'",
                (now_iso(), summary_json, int(shift_id))
            )
            shift_uuid_row = self.conn.execute(
                "SELECT uuid FROM shifts WHERE id=?", (int(shift_id),)
            ).fetchone()
            self._record_sync_event("SHIFT_CLOSED", {
                "shift_uuid":  shift_uuid_row[0] if shift_uuid_row else None,
                "shift_id":    int(shift_id),
                "ended_at":    now_iso(),
                "summary_json": summary_json or "",
            })

    def get_shift_summary(self, shift_id: int) -> Dict[str, Any]:
        cur = self.conn.execute("SELECT * FROM shifts WHERE id=?", (int(shift_id),))
        shift = cur.fetchone()
        if not shift:
            raise ValueError("الوردية غير موجودة")
        started = shift["started_at"]
        ended = shift["ended_at"] or now_iso()

        cur = self.conn.execute(
            "SELECT COUNT(*) as cnt, COALESCE(SUM(qty),0) as total_qty FROM movements WHERE direction='IN' AND ts >= ? AND ts <= ?",
            (started, ended))
        inflow = dict(cur.fetchone())

        cur = self.conn.execute(
            "SELECT item_type, school, color, size, SUM(qty) as qty FROM movements WHERE direction='IN' AND ts >= ? AND ts <= ? GROUP BY item_type, school, color, size",
            (started, ended))
        inflow_items = [dict(r) for r in cur.fetchall()]

        # Retail sales only: SALE/legacy rows, excluding warehouse return (customer الى المصنع)
        cur = self.conn.execute(
            """
            SELECT COUNT(*) as cnt, COALESCE(SUM(total),0) as total FROM bills
             WHERE created_at >= ? AND created_at <= ?
               AND (bill_type='SALE' OR bill_type IS NULL)
               AND TRIM(COALESCE(customer,'')) != ?
            """,
            (started, ended, WAREHOUSE_RETURN_LABEL),
        )
        sales = dict(cur.fetchone())

        cur = self.conn.execute(
            "SELECT COUNT(*) as cnt, COALESCE(SUM(total_amount),0) as total, COALESCE(SUM(paid_amount),0) as paid FROM reservations WHERE created_at >= ? AND created_at <= ?",
            (started, ended))
        res = dict(cur.fetchone())

        # Delivery payments collected during this shift
        cur = self.conn.execute(
            "SELECT COUNT(*) as cnt, COALESCE(SUM(unit_price),0) as total FROM movements WHERE direction='DELIVER_PAY' AND ts >= ? AND ts <= ?",
            (started, ended))
        deliver = dict(cur.fetchone())

        # Returns
        cur = self.conn.execute(
            "SELECT COUNT(*) as cnt, COALESCE(SUM(total),0) as total FROM bills WHERE created_at >= ? AND created_at <= ? AND bill_type='RETURN'",
            (started, ended))
        returns = dict(cur.fetchone())

        # Exchanges (total can be positive=customer paid or negative=refund)
        cur = self.conn.execute(
            "SELECT COUNT(*) as cnt, COALESCE(SUM(total),0) as total FROM bills WHERE created_at >= ? AND created_at <= ? AND bill_type='EXCHANGE'",
            (started, ended))
        exchanges = dict(cur.fetchone())

        return {
            "shift_id": shift["id"], "started_at": started, "ended_at": ended,
            "inflow_count": inflow["cnt"], "inflow_total_qty": inflow["total_qty"],
            "inflow_items": inflow_items,
            "sales_count": sales["cnt"], "sales_total": float(sales["total"]),
            "res_count": res["cnt"], "res_total": float(res["total"]), "res_paid": float(res["paid"]),
            "deliver_count": deliver["cnt"], "deliver_total": float(deliver["total"]),
            "return_count": returns["cnt"], "return_total": float(returns["total"]),
            "exchange_count": exchanges["cnt"], "exchange_total": float(exchanges["total"]),
        }

    def get_all_shifts(self, date_from: Optional[str] = None, date_to: Optional[str] = None) -> List[Dict[str, Any]]:
        """Return all shifts with computed summary columns."""
        wheres: List[str] = []
        params: List[Any] = []
        if date_from:
            wheres.append("s.started_at >= ?")
            params.append(date_from)
        if date_to:
            wheres.append("s.started_at <= ?")
            params.append(date_to + "T23:59:59")
        where_sql = (" WHERE " + " AND ".join(wheres)) if wheres else ""
        rows = self.conn.execute(
            f"SELECT * FROM shifts s{where_sql} ORDER BY s.started_at DESC", params
        ).fetchall()
        results: List[Dict[str, Any]] = []
        for row in rows:
            r = dict(row)
            started = r["started_at"]
            ended = r["ended_at"] or now_iso()
            # Retail sales (exclude warehouse returns)
            cur = self.conn.execute(
                """
                SELECT COUNT(*) as cnt, COALESCE(SUM(total),0) as total FROM bills
                 WHERE created_at >= ? AND created_at <= ?
                   AND (bill_type='SALE' OR bill_type IS NULL)
                   AND TRIM(COALESCE(customer,'')) != ?
                """,
                (started, ended, WAREHOUSE_RETURN_LABEL),
            )
            sales = dict(cur.fetchone())
            # reservations
            cur = self.conn.execute(
                "SELECT COUNT(*) as cnt, COALESCE(SUM(total_amount),0) as total, COALESCE(SUM(paid_amount),0) as paid FROM reservations WHERE created_at >= ? AND created_at <= ?",
                (started, ended))
            res = dict(cur.fetchone())
            # deliveries
            cur = self.conn.execute(
                "SELECT COUNT(*) as cnt, COALESCE(SUM(unit_price),0) as total FROM movements WHERE direction='DELIVER_PAY' AND ts >= ? AND ts <= ?",
                (started, ended))
            deliver = dict(cur.fetchone())
            # inflow
            cur = self.conn.execute(
                "SELECT COUNT(*) as cnt, COALESCE(SUM(qty),0) as total_qty FROM movements WHERE direction='IN' AND ts >= ? AND ts <= ?",
                (started, ended))
            inflow = dict(cur.fetchone())
            # returns
            cur = self.conn.execute(
                "SELECT COUNT(*) as cnt, COALESCE(SUM(total),0) as total FROM bills WHERE created_at >= ? AND created_at <= ? AND bill_type='RETURN'",
                (started, ended))
            returns = dict(cur.fetchone())
            # exchanges
            cur = self.conn.execute(
                "SELECT COUNT(*) as cnt, COALESCE(SUM(total),0) as total FROM bills WHERE created_at >= ? AND created_at <= ? AND bill_type='EXCHANGE'",
                (started, ended))
            exchanges = dict(cur.fetchone())
            r["sales_count"] = sales["cnt"]
            r["sales_total"] = float(sales["total"])
            r["res_count"] = res["cnt"]
            r["res_total"] = float(res["total"])
            r["res_paid"] = float(res["paid"])
            r["deliver_count"] = deliver["cnt"]
            r["deliver_total"] = float(deliver["total"])
            r["inflow_count"] = inflow["cnt"]
            r["inflow_total_qty"] = int(inflow["total_qty"])
            r["return_count"] = returns["cnt"]
            r["return_total"] = float(returns["total"])
            r["exchange_count"] = exchanges["cnt"]
            r["exchange_total"] = float(exchanges["total"])
            # Cash = sales + reservation payments + deliveries - returns + exchange net
            r["cash_collected"] = (r["sales_total"] + r["res_paid"] + r["deliver_total"]
                                   - r["return_total"] + r["exchange_total"])
            results.append(r)
        return results

    # -------- Income Bill Methods --------
    def create_income_bill(self, supplier: str, lines: List[Dict[str, Any]], note: str = "") -> int:
        """Create an income bill that groups multiple stock additions."""
        if not self.is_manager_feature_enabled("allow_manual_incoming"):
            raise PermissionError(_feature_restricted_message("الوارد اليدوي غير مسموح به حالياً في نقطة البيع."))
        self._require_shift()
        if not lines:
            raise ValueError("لا توجد أصناف في فاتورة الوارد")

        total_qty = 0
        total_value = 0.0

        with self.conn:
            cur = self.conn.execute(
                "INSERT INTO income_bills(created_at, supplier, total_qty, total_value, note) VALUES(?,?,0,0,?)",
                (now_iso(), (supplier or "").strip() or None, (note or "").strip() or None)
            )
            bill_id = int(cur.lastrowid)

            for line in lines:
                item_type = str(line["item_type"]).strip()
                school = str(line["school"]).strip()
                color = str(line["color"]).strip()
                size = str(line["size"]).strip()
                qty = int(line["qty"])
                price = float(line["unit_price"])

                if qty <= 0:
                    continue

                line_total = price * qty
                total_qty += qty
                total_value += line_total

                # Add to inventory
                self.add_stock(
                    item_type=item_type, school=school, color=color,
                    size=size, unit_price=price, count=qty,
                )

                # Record in income bill items
                self.conn.execute(
                    """INSERT INTO income_bill_items
                    (income_bill_id, item_type, school, color, size, unit_price, qty, line_total)
                    VALUES(?,?,?,?,?,?,?,?)""",
                    (bill_id, item_type, school, color, size, price, qty, line_total)
                )

            self.conn.execute(
                "UPDATE income_bills SET total_qty=?, total_value=? WHERE id=?",
                (total_qty, total_value, bill_id)
            )
            ib_uuid_row = self.conn.execute(
                "SELECT uuid FROM income_bills WHERE id=?", (bill_id,)
            ).fetchone()
            self._record_sync_event("INCOME_BILL_CREATED", {
                "income_bill_uuid": ib_uuid_row[0] if ib_uuid_row else None,
                "income_bill_id":   bill_id,
                "supplier":    (supplier or "").strip() or None,
                "note":        (note or "").strip() or None,
                "total_qty":   int(total_qty),
                "total_value": float(total_value),
                "lines": [
                    {
                        "item_type": str(ln["item_type"]).strip(),
                        "school":    str(ln["school"]).strip(),
                        "color":     str(ln["color"]).strip(),
                        "size":      str(ln["size"]).strip(),
                        "unit_price": float(ln["unit_price"]),
                        "qty":        int(ln["qty"]),
                    }
                    for ln in lines if int(ln.get("qty") or 0) > 0
                ],
                "shift_id": self.active_shift_id,
            })
            return bill_id

    def list_income_bills(self, date_from: Optional[str] = None, date_to: Optional[str] = None) -> List[Dict[str, Any]]:
        where = ["1=1"]
        args: list = []
        if date_from:
            where.append("date(created_at) >= date(?)")
            args.append(date_from)
        if date_to:
            where.append("date(created_at) <= date(?)")
            args.append(date_to)
        cur = self.conn.execute(
            f"SELECT * FROM income_bills WHERE {' AND '.join(where)} ORDER BY id DESC", args
        )
        return [dict(r) for r in cur.fetchall()]

    def list_income_bill_items(self, income_bill_id: int) -> List[Dict[str, Any]]:
        cur = self.conn.execute(
            "SELECT * FROM income_bill_items WHERE income_bill_id=? ORDER BY id", (int(income_bill_id),)
        )
        return [dict(r) for r in cur.fetchall()]

    def get_income_stats(self, date_from: Optional[str] = None, date_to: Optional[str] = None,
                         school: Optional[str] = None, item_type: Optional[str] = None,
                         color: Optional[str] = None) -> Dict[str, Any]:
        """Get income statistics optionally filtered."""
        where = ["1=1"]
        args: list = []
        if date_from:
            where.append("date(ib.created_at) >= date(?)")
            args.append(date_from)
        if date_to:
            where.append("date(ib.created_at) <= date(?)")
            args.append(date_to)

        item_where = ["1=1"]
        item_args: list = []
        if school:
            item_where.append("LOWER(TRIM(ibi.school)) = LOWER(?)")
            item_args.append(school.strip())
        if item_type:
            item_where.append("LOWER(TRIM(ibi.item_type)) = LOWER(?)")
            item_args.append(item_type.strip())
        if color:
            item_where.append("LOWER(TRIM(ibi.color)) = LOWER(?)")
            item_args.append(color.strip())

        cur = self.conn.execute(
            f"""SELECT COUNT(DISTINCT ib.id) as bill_count,
                       COALESCE(SUM(ibi.qty), 0) as total_qty,
                       COALESCE(SUM(ibi.line_total), 0) as total_value
                FROM income_bills ib
                JOIN income_bill_items ibi ON ibi.income_bill_id = ib.id
                WHERE {' AND '.join(where)} AND {' AND '.join(item_where)}""",
            args + item_args
        )
        row = cur.fetchone()
        return {
            "bill_count": int(row["bill_count"]),
            "total_qty": int(row["total_qty"]),
            "total_value": float(row["total_value"]),
        }

    # -------- Bulk Price Update --------
    def bulk_update_prices(self, mode: str, value: float, constraints: Optional[Dict[str, Any]] = None) -> int:
        """Update prices in bulk. mode='percentage' or 'fixed'. Returns count of updated rows."""
        if not self.is_manager_feature_enabled("allow_bulk_price"):
            raise PermissionError(_feature_restricted_message("تعديل الأسعار من نقطة البيع غير مسموح به حالياً."))
        import math
        where_parts = ["count > 0"]
        args: list = []
        if constraints:
            for k in ("item_type", "school", "color", "size"):
                v = constraints.get(k)
                if v and str(v).strip():
                    where_parts.append(f"LOWER(TRIM({k})) = LOWER(?)")
                    args.append(str(v).strip())

        where = " AND ".join(where_parts)
        cur = self.conn.execute(f"SELECT id, unit_price FROM stocks WHERE {where}", args)
        rows = cur.fetchall()

        if not rows:
            return 0

        def _round_up_5(x: float) -> float:
            return float(math.ceil(x / 5.0) * 5)

        updated = 0
        with self.conn:
            for r in rows:
                old_price = float(r["unit_price"])
                if mode == "percentage":
                    new_price = _round_up_5(old_price * (1 + value / 100.0))
                else:
                    new_price = old_price + value

                if new_price < 0:
                    new_price = 0.0

                if abs(new_price - old_price) > 0.001:
                    self.conn.execute("UPDATE stocks SET unit_price=? WHERE id=?", (new_price, int(r["id"])))
                    self.conn.execute(
                        """INSERT INTO movements(ts,direction,stock_id,qty,note,item_type,school,color,size,unit_price)
                        SELECT ?,'PRICE_UPDATE',id,0,?,item_type,school,color,size,? FROM stocks WHERE id=?""",
                        (now_iso(), f"تعديل سعر جماعي: {old_price:.2f} -> {new_price:.2f}", new_price, int(r["id"]))
                    )
                    updated += 1
            if updated > 0:
                self._record_sync_event("PRICE_UPDATE", {
                    "mode":        mode,
                    "value":       float(value),
                    "constraints": constraints or {},
                    "updated_count": int(updated),
                    "scope":       "bulk",
                })
        return updated

    def get_sales_stats(self, date_from=None, date_to=None, school=None, item_type=None, color=None):
        """Retail sale bill totals (excludes warehouse-return bills) plus reservation delivery cash."""
        # Bill stats - if spec filters exist, join to bill_items
        bill_where = ["1=1"]
        bill_args = []
        if date_from:
            bill_where.append("date(b.created_at) >= date(?)")
            bill_args.append(date_from)
        if date_to:
            bill_where.append("date(b.created_at) <= date(?)")
            bill_args.append(date_to)

        bill_where.append("(COALESCE(b.bill_type,'SALE')='SALE' OR b.bill_type IS NULL)")
        bill_where.append("TRIM(COALESCE(b.customer,'')) != ?")
        bill_args.append(WAREHOUSE_RETURN_LABEL)

        if school or item_type or color:
            # Need to join bill_items for spec filtering
            spec_where = []
            if school:
                spec_where.append("LOWER(TRIM(bi.school)) = LOWER(?)")
                bill_args.append(school.strip())
            if item_type:
                spec_where.append("LOWER(TRIM(bi.item_type)) = LOWER(?)")
                bill_args.append(item_type.strip())
            if color:
                spec_where.append("LOWER(TRIM(bi.color)) = LOWER(?)")
                bill_args.append(color.strip())

            cur = self.conn.execute(
                f"""SELECT COUNT(DISTINCT b.id) as cnt,
                           COALESCE(SUM(bi.line_total), 0) as total
                    FROM bills b
                    JOIN bill_items bi ON bi.bill_id = b.id
                    WHERE {' AND '.join(bill_where)} AND {' AND '.join(spec_where)}""",
                bill_args
            )
        else:
            cur = self.conn.execute(
                f"SELECT COUNT(*) as cnt, COALESCE(SUM(total),0) as total FROM bills b WHERE {' AND '.join(bill_where)}",
                bill_args
            )

        bill_row = dict(cur.fetchone())

        # Reservation stats
        res_where = ["1=1"]
        res_args = []
        if date_from:
            res_where.append("date(created_at) >= date(?)")
            res_args.append(date_from)
        if date_to:
            res_where.append("date(created_at) <= date(?)")
            res_args.append(date_to)
        if school:
            res_where.append("LOWER(TRIM(school)) = LOWER(?)")
            res_args.append(school.strip())
        if item_type:
            res_where.append("LOWER(TRIM(item_type)) = LOWER(?)")
            res_args.append(item_type.strip())
        if color:
            res_where.append("LOWER(TRIM(color)) = LOWER(?)")
            res_args.append(color.strip())

        cur = self.conn.execute(
            f"""SELECT COUNT(*) as cnt, COALESCE(SUM(total_amount),0) as total,
                       COALESCE(SUM(paid_amount),0) as paid
                FROM reservations WHERE {' AND '.join(res_where)}""",
            res_args
        )
        res_row = dict(cur.fetchone())

        dwhere = ["direction='DELIVER_PAY'"]
        dargs: List[Any] = []
        if date_from:
            dwhere.append("date(ts) >= date(?)")
            dargs.append(date_from)
        if date_to:
            dwhere.append("date(ts) <= date(?)")
            dargs.append(date_to)
        if school:
            dwhere.append("LOWER(TRIM(school)) = LOWER(?)")
            dargs.append(school.strip())
        if item_type:
            dwhere.append("LOWER(TRIM(item_type)) = LOWER(?)")
            dargs.append(item_type.strip())
        if color:
            dwhere.append("LOWER(TRIM(color)) = LOWER(?)")
            dargs.append(color.strip())
        cur = self.conn.execute(
            f"SELECT COALESCE(SUM(unit_price), 0) AS dt FROM movements WHERE {' AND '.join(dwhere)}",
            dargs,
        )
        dr = cur.fetchone()
        deliver_cash = float(dr["dt"] if dr and dr["dt"] is not None else 0)

        return {
            "sales_count": bill_row["cnt"],
            "sales_total": float(bill_row["total"]) + deliver_cash,
            "res_count": res_row["cnt"], "res_total": float(res_row["total"]),
            "res_paid": float(res_row["paid"]),
            "deliver_cash": deliver_cash,
        }

    def get_item_movement_stats(self, date_from=None, date_to=None, school=None, item_type=None, color=None):
        """Get received/sold/reserved counts per item."""
        where = ["direction IN ('IN','OUT','OUT_FACTORY','RESERVE','ADJUST_OUT')"]
        args: list = []
        if date_from:
            where.append("date(ts) >= date(?)")
            args.append(date_from)
        if date_to:
            where.append("date(ts) <= date(?)")
            args.append(date_to)
        if school:
            where.append("LOWER(TRIM(school)) = LOWER(?)")
            args.append(school.strip())
        if item_type:
            where.append("LOWER(TRIM(item_type)) = LOWER(?)")
            args.append(item_type.strip())
        if color:
            where.append("LOWER(TRIM(color)) = LOWER(?)")
            args.append(color.strip())
        cur = self.conn.execute(f"""
            SELECT item_type, school, color, size,
                SUM(CASE WHEN direction='IN' THEN qty ELSE 0 END) as received,
                SUM(CASE WHEN direction IN ('OUT','OUT_FACTORY') THEN qty ELSE 0 END) as sold,
                SUM(CASE WHEN direction='RESERVE' THEN qty ELSE 0 END) as reserved,
                SUM(CASE WHEN direction='ADJUST_OUT' THEN qty ELSE 0 END) as adjusted
            FROM movements
            WHERE {' AND '.join(where)}
            GROUP BY item_type, school, color, size
            ORDER BY item_type, school, color, size
        """, args)
        return [dict(r) for r in cur.fetchall()]

    def reset_movement_counts(self):
        """Reset all movement records. Keep item definitions intact."""
        if not self.is_manager_feature_enabled("allow_reset_counts"):
            raise PermissionError(_feature_restricted_message("إعادة التعيين غير مسموح بها حالياً من نقطة البيع."))
        with self.conn:
            self.conn.execute("DELETE FROM movements")
            self.conn.execute("DELETE FROM reservations")
            self.conn.execute("DELETE FROM bills")
            self.conn.execute("DELETE FROM bill_items")

    # -------- Bill history APIs --------
    def list_bills(self) -> List[Dict[str, Any]]:
        cur = self.conn.cursor()
        cur.execute(
            "SELECT id,created_at,customer,total,"
            " COALESCE(status,'CONFIRMED') AS status,"
            " COALESCE(bill_type,'SALE') AS bill_type,"
            " void_reason, voided_at"
            " FROM bills ORDER BY id DESC"
        )
        rows = [dict(r) for r in cur.fetchall()]
        cur.close()
        return rows

    def void_bill(self, bill_id: int, reason: str) -> None:
        reason = (reason or "").strip()
        if not reason:
            raise ValueError("سبب الإلغاء مطلوب.")

        with self.conn:
            bill = self.conn.execute(
                "SELECT id, customer, total, COALESCE(status,'CONFIRMED') AS status, "
                "COALESCE(bill_type,'SALE') AS bill_type, uuid "
                "FROM bills WHERE id = ?",
                (int(bill_id),),
            ).fetchone()
            if bill is None:
                raise ValueError("الفاتورة غير موجودة.")
            if str(bill["status"]).upper() == "VOID":
                raise ValueError("تم إلغاء هذه الفاتورة بالفعل.")
            if str(bill["bill_type"]).upper() != "SALE":
                raise ValueError("الإلغاء المباشر متاح لفواتير البيع فقط. استخدم المرتجع للأنواع الأخرى.")

            items = self.conn.execute(
                "SELECT item_type, school, color, size, unit_price, qty, origin "
                "FROM bill_items WHERE bill_id = ?",
                (int(bill_id),),
            ).fetchall()
            if not items:
                raise ValueError("لا توجد بنود في هذه الفاتورة.")

            for item in items:
                if str(item["origin"] or "STOCK").upper() != "STOCK":
                    continue
                cur = self.conn.execute(
                    """INSERT INTO stocks(item_type,school,color,size,unit_price,count)
                       VALUES(?,?,?,?,?,?)""",
                    (
                        item["item_type"], item["school"], item["color"], item["size"],
                        float(item["unit_price"]), int(item["qty"]),
                    ),
                )
                stock_id = int(cur.lastrowid)
                self.conn.execute(
                    """INSERT INTO movements
                       (ts,direction,stock_id,qty,note,bill_id,item_type,school,color,size,unit_price)
                       VALUES(?,?,?,?,?,?,?,?,?,?,?)""",
                    (
                        now_iso(),
                        "VOID_IN",
                        stock_id,
                        int(item["qty"]),
                        f"Bill void: {reason}",
                        int(bill_id),
                        item["item_type"],
                        item["school"],
                        item["color"],
                        item["size"],
                        float(item["unit_price"]),
                    ),
                )

            self.conn.execute(
                "UPDATE bills SET status='VOID', void_reason=?, voided_at=? WHERE id=?",
                (reason, now_iso(), int(bill_id)),
            )

            self._record_sync_event("SALE_VOIDED", {
                "bill_uuid": bill["uuid"] if "uuid" in bill.keys() else None,
                "bill_id": int(bill_id),
                "customer": bill["customer"],
                "total": float(bill["total"] or 0),
                "reason": reason,
                "shift_id": self.active_shift_id,
            })

    def list_bill_items(self, bill_id: int) -> List[Dict[str, Any]]:
        cur = self.conn.cursor()
        cur.execute(
            """SELECT item_type,school,color,size,unit_price,qty,line_total,origin
            FROM bill_items WHERE bill_id=?""",
            (int(bill_id),),
        )
        rows = [dict(r) for r in cur.fetchall()]
        cur.close()
        return rows

    def export_inventory_excel(self, path: str, rows: Sequence[Dict[str, Any]]) -> None:
        headers = [
            "id","item_type","school","color","size","unit_price","count","value",
        ]
        table = []
        for r in rows:
            value = r.get("value", float(r["unit_price"]) * int(r["count"]))
            table.append(
                [
                    r["id"], r["item_type"], r["school"], r["color"], r["size"],
                    float(r["unit_price"]),
                    int(r["count"]), float(value),
                ]
            )
        export_to_excel(path, headers, table)

    def get_app_setting(self, key: str, default: Optional[str] = None) -> Optional[str]:
        row = self.conn.execute(
            "SELECT value FROM app_settings WHERE key = ?",
            (str(key),),
        ).fetchone()
        if row is None:
            return default
        return str(row["value"])

    def set_app_setting(self, key: str, value: Any) -> None:
        with self.conn:
            self.conn.execute(
                "INSERT INTO app_settings(key, value) VALUES(?, ?) "
                "ON CONFLICT(key) DO UPDATE SET value = excluded.value",
                (str(key), str(value)),
            )

    def is_manager_feature_enabled(self, feature_key: str) -> bool:
        default = "1" if POS_MANAGER_FEATURE_DEFAULTS.get(feature_key, False) else "0"
        raw = (self.get_app_setting(f"feature:{feature_key}", default) or default).strip().lower()
        return raw in ("1", "true", "yes", "on")

    def set_manager_feature_enabled(self, feature_key: str, enabled: bool) -> None:
        self.set_app_setting(f"feature:{feature_key}", "1" if enabled else "0")

    def verify_admin_password(self, plain: str) -> bool:
        stored = self.get_app_setting("admin_password", ADMIN_PASSWORD_PLAIN) or ADMIN_PASSWORD_PLAIN
        if str(stored).startswith(ADMIN_PASSWORD_HASH_PREFIX):
            digest = hashlib.sha256(str(plain).encode("utf-8")).hexdigest()
            return str(stored) == f"{ADMIN_PASSWORD_HASH_PREFIX}{digest}"
        ok = str(plain) == str(stored)
        # Backward-compatible auto-migration from plain-text storage.
        if ok:
            self.set_admin_password(str(plain))
        return ok

    def set_admin_password(self, plain: str) -> None:
        digest = hashlib.sha256(str(plain).encode("utf-8")).hexdigest()
        self.set_app_setting("admin_password", f"{ADMIN_PASSWORD_HASH_PREFIX}{digest}")

    def remove_from_stock(self, stock_id: int, qty: Optional[int], note: str = "Admin remove") -> int:
        if not self.is_manager_feature_enabled("allow_inventory_delete"):
            raise PermissionError(_feature_restricted_message("الحذف اليدوي من مخزون نقطة البيع غير مسموح به حالياً."))
        with self.conn:
            cur = self.conn.execute("SELECT * FROM stocks WHERE id=?", (int(stock_id),))
            s = cur.fetchone()
            if not s:
                raise ValueError("Stock row not found.")
            current = int(s["count"])
            if current <= 0:
                raise ValueError("No quantity available to remove.")
            take = current if qty is None else min(int(qty), current)
            if take <= 0:
                raise ValueError("Quantity must be > 0")
            self.conn.execute("UPDATE stocks SET count = count - ? WHERE id = ?", (int(take), int(stock_id)))
            self.conn.execute(
                """INSERT INTO movements
                   (ts,direction,stock_id,qty,note,bill_id,item_type,school,color,size,unit_price)
                   VALUES(?,?,?,?,?,?,?,?,?,?,?)""",
                (
                    now_iso(), "ADJUST_OUT", int(stock_id), int(take), note, None,
                    s["item_type"], s["school"], s["color"], s["size"],
                    float(s["unit_price"]),
                ),
            )
            self._record_sync_event("STOCK_ADJUST", {
                "stock_uuid": s["uuid"] if "uuid" in s.keys() else None,
                "stock_id":   int(stock_id),
                "direction":  "OUT",
                "qty":        int(take),
                "note":       note,
                "item_type":  s["item_type"],
                "school":     s["school"],
                "color":      s["color"],
                "size":       s["size"],
                "unit_price": float(s["unit_price"]),
            })
            return int(take)

    def list_movements(self, filters: Dict[str, Any]) -> List[Dict[str, Any]]:
        where = ["1=1"]
        args: List[Any] = []
        if filters.get("direction"):
            where.append("direction = ?")
            args.append(filters["direction"])
        for k in ("item_type", "school", "color", "size"):
            v = (filters.get(k) or "").strip()
            if v:
                where.append(f"LOWER(TRIM({k})) = LOWER(TRIM(?))")
                args.append(v)

        cust = (filters.get("customer") or "").strip()
        if cust:
            where.append("bill_id IN (SELECT id FROM bills WHERE LOWER(TRIM(customer)) = LOWER(TRIM(?)))")
            args.append(cust)

        txt = (filters.get("text") or "").strip()
        if txt:
            like = f"%{txt}%"
            where.append("""(
                LOWER(COALESCE(note,''))      LIKE LOWER(?)
            OR LOWER(COALESCE(item_type,'')) LIKE LOWER(?)
            OR LOWER(COALESCE(school,''))    LIKE LOWER(?)
            OR LOWER(COALESCE(color,''))     LIKE LOWER(?)
            OR LOWER(COALESCE(size,''))      LIKE LOWER(?)
            )""")
            args += [like, like, like, like, like]

        df = (filters.get("date_from") or "").strip()
        dt = (filters.get("date_to")   or "").strip()
        if df:
            where.append("date(ts) >= date(?)")
            args.append(df)
        if dt:
            where.append("date(ts) <= date(?)")
            args.append(dt)

        cur = self.conn.cursor()
        cur.execute(f"""
            SELECT id, ts, direction, stock_id, qty, note, bill_id,
                item_type, school, color, size, unit_price
            FROM movements
            WHERE {' AND '.join(where)}
            ORDER BY ts DESC, id DESC
        """, args)
        rows = [dict(r) for r in cur.fetchall()]
        cur.close()
        return rows


# ------------------- UI helpers -------------------
class LabeledStaticCombo(ttk.Frame):
    def __init__(self, master, text: str, values, **kwargs):
        super().__init__(master)
        ttk.Label(self, text=text).pack(side=tk.TOP, anchor="w")
        self.var = tk.StringVar()
        self.cb = ttk.Combobox(self, textvariable=self.var, values=values, state="readonly", **kwargs)
        self.cb.pack(fill=tk.X)

    def get(self) -> str:
        return (self.var.get() or "").strip()

    def set(self, v: str) -> None:
        self.var.set(v)

class LabeledCombobox(ttk.Frame):
    FILTER_IDLE_MS = 60
    MATCH_MODE = "startswith"
    MIN_CHARS_TO_OPEN = 1
    OPEN_ALL_ON_BUTTON = True

    def __init__(self, master, text: str, db: 'SqliteDatabase', field: str, **kwargs):
        super().__init__(master)
        self.db = db
        self.field = field

        ttk.Label(self, text=text).pack(side=tk.TOP, anchor="w")

        row = ttk.Frame(self)
        row.pack(fill=tk.X)

        self.var = tk.StringVar()
        self.cb = ttk.Entry(row, textvariable=self.var, **kwargs)
        self.cb.pack(side=tk.LEFT, fill=tk.X, expand=True)

        self.btn = ttk.Button(row, text="▼", width=2, takefocus=0, command=self._on_caret_click)
        try:
            self.btn.configure(style="Toolbutton")
        except Exception:
            pass
        self.btn.pack(side=tk.RIGHT, padx=(4, 0))

        self._all_values: List[str] = []
        self._supplier = None
        self._debounce_job: Optional[str] = None
        self._suspend_set = False
        self._opened = False
        self._mouse_inside_popup = False

        self._popup: Optional[tk.Toplevel] = None
        self._list: Optional[tk.Listbox] = None
        self._ysb: Optional[ttk.Scrollbar] = None

        self.cb.bind("<FocusIn>", self._on_focus, add="+")
        self.cb.bind("<KeyRelease>", self._on_key_release, add="+")
        self.cb.bind("<Down>", self._on_down, add="+")
        self.cb.bind("<Escape>", self._on_escape, add="+")
        self.cb.bind("<FocusOut>", self._on_focus_out, add="+")

        self.winfo_toplevel().bind("<Configure>", self._reposition_popup_safely, add="+")

    def get(self) -> str:
        return (self.var.get() or "").strip()

    def set(self, v: str) -> None:
        self._suspend_set = True
        try:
            self.var.set(v or "")
        finally:
            self._suspend_set = False
        self._close_popup()

    def set_supplier(self, fn):
        self._supplier = fn
        self.refresh_values()

    def refresh_values(self):
        self._ensure_data(force=True)
        if self._opened:
            self._fill_list(self._filtered(self.get()))

    def _ensure_data(self, force: bool = False) -> None:
        if self._supplier:
            try:
                vals = self._supplier() or []
            except Exception:
                vals = []
            if force or (vals != self._all_values):
                self._all_values = vals[:]
            return

        if not self._all_values or force:
            try:
                self._all_values = self.db.get_distinct_filtered(self.field, {}) or []
            except Exception:
                self._all_values = []

    def _create_popup(self) -> None:
        if self._popup and self._popup.winfo_exists():
            return
        self._popup = tk.Toplevel(self)
        self._popup.overrideredirect(True)
        self._popup.attributes("-topmost", True)
        self._popup.withdraw()

        frm = ttk.Frame(self._popup, borderwidth=1, relief="solid")
        frm.pack(fill=tk.BOTH, expand=True)

        self._list = tk.Listbox(frm, activestyle="none")
        self._list.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)

        self._ysb = ttk.Scrollbar(frm, orient="vertical", command=self._list.yview)
        self._ysb.pack(side=tk.RIGHT, fill=tk.Y)
        self._list.configure(yscrollcommand=self._ysb.set)
        _bind_mousewheel(self._list)

        self._list.bind("<ButtonRelease-1>", self._on_list_click, add="+")
        self._list.bind("<Motion>", lambda e: self._set_active_under_mouse(e), add="+")
        self._popup.bind("<Enter>", lambda *_: self._mark_mouse(True), add="+")
        self._popup.bind("<Leave>", lambda *_: self._mark_mouse(False), add="+")

    def _open_popup(self, *, items: Optional[List[str]] = None) -> None:
        if self._opened and items is None:
            return
        self._create_popup()
        if items is None:
            items = self._filtered(self.get())
        self._fill_list(items or [])
        if not self._list or self._list.size() == 0:
            self._close_popup()
            return

        self._position_popup()
        self._popup.deiconify()
        self._opened = True
        self.after_idle(self._restore_entry_focus)

    def _close_popup(self) -> None:
        if self._popup and self._popup.winfo_exists():
            try:
                self._popup.withdraw()
            except Exception:
                pass
        self._opened = False

    def _position_popup(self) -> None:
        if not (self._popup and self._popup.winfo_exists()):
            return
        try:
            x = self.cb.winfo_rootx()
            y = self.cb.winfo_rooty() + self.cb.winfo_height()
            w = self.cb.winfo_width() + self.btn.winfo_width() + 4
            rows = min(8, max(1, self._list.size() if self._list else 0))
            row_h = max(18, int(self.cb.winfo_fpixels("1.2m")))
            h = rows * row_h + 2
            self._popup.geometry(f"{w}x{h}+{x}+{y}")
        except Exception:
            pass

    def _reposition_popup_safely(self, *_):
        if self._opened:
            self.after_idle(self._position_popup)

    def _filtered(self, text: str) -> List[str]:
        self._ensure_data()
        tl = (text or "").lower()
        if not tl:
            return self._all_values[:]
        if self.MATCH_MODE == "contains":
            return [v for v in self._all_values if tl in str(v).lower()]
        return [v for v in self._all_values if str(v).lower().startswith(tl)]

    def _fill_list(self, items: List[str]) -> None:
        if not self._list:
            return
        self._list.delete(0, tk.END)
        for v in items:
            self._list.insert(tk.END, v)
        if items:
            self._list.activate(0)
            self._list.selection_clear(0, tk.END)
        self._position_popup()

    def _on_focus(self, *_):
        self._ensure_data()

    def _on_key_release(self, event):
        if self._suspend_set:
            return
        if event.keysym in ("Shift_L", "Shift_R", "Control_L", "Control_R", "Alt_L", "Alt_R",
                            "Left", "Right", "Home", "End", "Prior", "Next"):
            return
        if event.keysym == "Up":
            if self._opened and self._list and self._list.size() > 0:
                idx = self._list.index("active")
                idx = max(0, idx - 1)
                self._list.activate(idx); self._ensure_visible(idx)
            return "break"
        if event.keysym == "Down":
            return self._on_down(event)
        if event.keysym in ("Return", "KP_Enter"):
            if self._opened and self._list and self._list.size() > 0:
                self._commit_active_selection()
                return "break"
            return

        if self._debounce_job:
            try: self.after_cancel(self._debounce_job)
            except Exception: pass
        self._debounce_job = self.after(self.FILTER_IDLE_MS, self._apply_filter)

    def _apply_filter(self):
        text = self.get()
        items = self._filtered(text)
        if items:
            if self._opened:
                self._fill_list(items)
            else:
                if len(text) >= self.MIN_CHARS_TO_OPEN:
                    self._open_popup(items=items)
        else:
            self._close_popup()

    def _on_down(self, _event):
        if not self._opened:
            items = self._filtered(self.get())
            if not items and self.OPEN_ALL_ON_BUTTON:
                self._ensure_data()
                items = self._all_values[:]
            self._open_popup(items=items)
            return "break"
        if self._list and self._list.size() > 0:
            idx = self._list.index("active")
            idx = min(self._list.size()-1, max(0, idx + 1))
            self._list.activate(idx); self._ensure_visible(idx)
        return "break"

    def _on_escape(self, *_):
        self._close_popup()

    def _on_focus_out(self, *_):
        self.after(20, self._maybe_close_after_blur)

    def _maybe_close_after_blur(self):
        if self._opened and not self._mouse_inside_popup and self.focus_get() is not self.cb:
            self._close_popup()

    def _restore_entry_focus(self):
        try:
            self.cb.focus_set()
            insert_pos = self.cb.index("insert")
            self.cb.icursor(insert_pos)
            try:
                self.cb.selection_clear()
            except TypeError:
                self.cb.selection_clear(0, tk.END)
        except Exception:
            pass

    def _on_caret_click(self):
        if self._opened:
            self._close_popup()
        else:
            if self.OPEN_ALL_ON_BUTTON and not self.get():
                self._ensure_data()
                self._open_popup(items=self._all_values[:])
            else:
                self._open_popup()
        self.after_idle(self._restore_entry_focus)

    def _on_list_click(self, *_):
        self._commit_active_selection()

    def _commit_active_selection(self):
        if not (self._list and self._list.size() > 0):
            return
        idx = self._list.index("active")
        if 0 <= idx < self._list.size():
            val = self._list.get(idx)
            self._suspend_set = True
            try:
                self.var.set(val)
            finally:
                self._suspend_set = False
            self._close_popup()
            try:
                self.cb.event_generate("<<ComboboxSelected>>")
            except Exception:
                pass
            self.after_idle(self._restore_entry_focus)

    def _set_active_under_mouse(self, event):
        if not self._list:
            return
        try:
            idx = self._list.nearest(event.y)
            if 0 <= idx < self._list.size():
                self._list.activate(idx)
        except Exception:
            pass

    def _ensure_visible(self, idx: int):
        if not self._list:
            return
        top = self._list.nearest(0)
        bot = self._list.nearest(self._list.winfo_height())
        if idx <= top:
            self._list.see(max(0, idx-1))
        elif idx >= bot:
            self._list.see(min(self._list.size()-1, idx+1))

    def _mark_mouse(self, inside: bool):
        self._mouse_inside_popup = inside

class LabeledEntry(ttk.Frame):
    def __init__(self, master, text: str, **kwargs):
        super().__init__(master)
        ttk.Label(self, text=text).pack(side=tk.TOP, anchor="w")
        self.var = tk.StringVar()
        self.entry = ttk.Entry(self, textvariable=self.var, **kwargs)
        self.entry.pack(fill=tk.X)

    def get(self) -> str:
        return self.var.get().strip()

    def set(self, v: str) -> None:
        self.var.set(v)

class DateField(ttk.Frame):
    # Entry + calendar button; popup calendar; returns YYYY-MM-DD; empty allowed.
    def __init__(self, master, label: str):
        super().__init__(master)
        ttk.Label(self, text=label).pack(side=tk.TOP, anchor="w")
        row = ttk.Frame(self); row.pack(fill=tk.X)
        self.var = tk.StringVar()
        self.entry = ttk.Entry(row, textvariable=self.var, width=14)
        self.entry.pack(side=tk.LEFT, fill=tk.X, expand=True)
        self.btn = ttk.Button(row, text="📅", width=3, command=self._open)
        self.btn.pack(side=tk.LEFT, padx=(4,0))
        self._popup: Optional[tk.Toplevel] = None
        self.entry.bind("<Return>", lambda e: self._validate())
        self.entry.bind("<Escape>", lambda e: self._close())
        self.entry.bind("<FocusOut>", lambda e: self._close())

    def _open(self):
        self._close()
        tp = tk.Toplevel(self)
        self._popup = tp
        tp.wm_overrideredirect(True)
        tp.attributes("-topmost", True)
        x = self.entry.winfo_rootx(); y = self.entry.winfo_rooty() + self.entry.winfo_height()
        tp.geometry(f"+{x}+{y}")
        tp.bind("<FocusOut>", lambda e: self._close())
        tp.bind("<Escape>", lambda e: self._close())
        try:
            tp.focus_force()
        except Exception:
            pass

        frm = ttk.Frame(tp, padding=6, borderwidth=1, relief="solid"); frm.pack()
        today = date.today()
        try:
            base = datetime.strptime(self.var.get(), "%Y-%m-%d").date()
        except Exception:
            base = today
        self._year = tk.IntVar(value=base.year)
        self._month = tk.IntVar(value=base.month)

        hdr = ttk.Frame(frm); hdr.pack(fill=tk.X)
        ttk.Button(hdr, text="◄", width=2, command=self._prev_month).pack(side=tk.LEFT)
        self._title = ttk.Label(hdr, font=("", 10, "bold")); self._title.pack(side=tk.LEFT, padx=6)
        ttk.Button(hdr, text="►", width=2, command=self._next_month).pack(side=tk.RIGHT)

        self._grid = ttk.Frame(frm); self._grid.pack()
        self._render()

    def _render(self):
        for w in self._grid.winfo_children(): w.destroy()
        y, m = self._year.get(), self._month.get()
        self._title.config(text=f"{calendar.month_name[m]} {y}")
        hdr = ttk.Frame(self._grid); hdr.grid(row=0, column=0, columnspan=7, pady=(0,4))
        for i, wk in enumerate(["Mo","Tu","We","Th","Fr","Sa","Su"]):
            ttk.Label(hdr, text=wk, width=3, anchor="center").grid(row=0, column=i)
        for r, week in enumerate(calendar.Calendar(firstweekday=0).monthdayscalendar(y, m), start=1):
            for c, d_ in enumerate(week):
                if d_ == 0:
                    ttk.Label(self._grid, text="   ", width=3).grid(row=r, column=c)
                else:
                    ttk.Button(self._grid, text=f"{d_:02d}", width=3,
                               command=lambda dd=d_: self._pick(y, m, dd)).grid(row=r, column=c, padx=1, pady=1)

    def _prev_month(self):
        m, y = self._month.get() - 1, self._year.get()
        if m < 1: m, y = 12, y - 1
        self._month.set(m); self._year.set(y); self._render()

    def _next_month(self):
        m, y = self._month.get() + 1, self._year.get()
        if m > 12: m, y = 1, y + 1
        self._month.set(m); self._year.set(y); self._render()

    def _pick(self, y: int, m: int, d_: int):
        self.var.set(f"{y:04d}-{m:02d}-{d_:02d}")
        self._close()

    def _validate(self):
        txt = self.var.get().strip()
        if not txt: return
        try:
            datetime.strptime(txt, "%Y-%m-%d")
        except Exception:
            messagebox.showerror("صيغة التاريخ", "الرجاء إدخال التاريخ بصيغة YYYY-MM-DD أو استخدم التقويم.", parent=self)
            self.entry.focus_set()

    def _close(self):
        if self._popup and self._popup.winfo_exists():
            self._popup.destroy()
        self._popup = None

    def get(self) -> str:
        return self.var.get().strip()

    def set(self, v: str) -> None:
        self.var.set(v or "")

class SizesGrid(ttk.Frame):
    COLS = 6

    def __init__(self, master, sizes: List[str], **kwargs):
        super().__init__(master, **kwargs)

        self.fixed_sizes = sizes[:]
        self.extra_sizes = ["__custom1__", "__custom2__"]
        self.sizes = self.fixed_sizes + self.extra_sizes

        self.idx_4xl = self.fixed_sizes.index("4XL") if "4XL" in self.fixed_sizes else None
        self.idx_5xl = self.fixed_sizes.index("5XL") if "5XL" in self.fixed_sizes else None

        self.editable_indices = set(filter(
            lambda x: x is not None,
            {
                self.idx_4xl,
                self.idx_5xl,
                len(self.fixed_sizes),
                len(self.fixed_sizes) + 1,
            }
        ))

        self._vars = {}

        for i, sz in enumerate(self.sizes):
            r = i // self.COLS
            c = i % self.COLS

            cell = ttk.Frame(self, padding=(4, 4))
            cell.grid(row=r, column=c, sticky="nsew", padx=2, pady=2)

            is_custom = i in self.editable_indices

            if is_custom:
                v_label = tk.StringVar(value="" if sz.startswith("__") else sz)
                ttk.Entry(cell, textvariable=v_label, width=7, justify="center").pack()
            else:
                v_label = tk.StringVar(value=sz)
                ttk.Label(cell, text=sz, anchor="center").pack()

            ttk.Label(cell, text="الكمية", font=("", 8)).pack()
            v_qty = tk.StringVar()
            ttk.Entry(cell, textvariable=v_qty, width=6).pack()

            ttk.Label(cell, text="السعر", font=("", 8)).pack()
            v_price = tk.StringVar()
            ttk.Entry(cell, textvariable=v_price, width=8).pack()

            self._vars[sz] = (v_qty, v_price, v_label, is_custom)

        for c in range(self.COLS):
            self.columnconfigure(c, weight=1)

    def get_rows(self):
        out = []
        for sz in self.sizes:
            v_qty, v_price, v_label, is_custom = self._vars[sz]

            size_label = v_label.get().strip() if is_custom else sz
            if not size_label:
                continue

            qty_txt = v_qty.get().strip()
            if not qty_txt:
                continue

            try:
                qty = int(qty_txt)
            except Exception:
                continue

            if qty <= 0:
                continue

            price_txt = v_price.get().strip()
            try:
                price = float(price_txt) if price_txt else None
            except Exception:
                price = None

            out.append({
                "size": size_label,
                "qty": qty,
                "price": price
            })

        return out

    def set_price_for_size(self, size_label: str, price):
        """Set the price for a specific size in the grid."""
        for sz in self.sizes:
            v_qty, v_price, v_label, is_custom = self._vars[sz]
            label = v_label.get().strip() if is_custom else sz
            if label == size_label:
                v_price.set(str(price))
                break

    def set_price_for_all(self, price):
        """Set the same price for all sizes in the grid."""
        for sz in self.sizes:
            v_qty, v_price, v_label, is_custom = self._vars[sz]
            v_price.set(str(price))

    def clear_quantities(self):
        """Clear quantity fields in the grid, keeping price fields intact."""
        for sz in self.sizes:
            v_qty, v_price, v_label, is_custom = self._vars[sz]
            v_qty.set("")


# ------------------- Toast Notification -------------------

class ToastNotification:
    """Non-blocking floating notification that auto-dismisses."""
    _active: list = []

    @classmethod
    def show(cls, parent, message, duration=2500, toast_type="success"):
        return cls(parent, message, duration, toast_type)

    def __init__(self, parent, message, duration, toast_type):
        colors = {
            "success": ("#166534", "#dcfce7", "#16a34a"),
            "info":    ("#1e40af", "#dbeafe", "#2563eb"),
            "warning": ("#854d0e", "#fef9c3", "#f59e0b"),
            "error":   ("#991b1b", "#fee2e2", "#dc2626"),
        }
        fg, bg, border = colors.get(toast_type, colors["info"])
        icons = {"success": "\u2714", "info": "\u2139", "warning": "\u26A0", "error": "\u2716"}
        icon = icons.get(toast_type, "")

        y_off = 8
        for t in ToastNotification._active:
            try:
                if t._frame.winfo_exists():
                    y_off += t._frame.winfo_reqheight() + 6
            except Exception:
                pass

        self._frame = tk.Frame(parent, bg=bg, highlightbackground=border,
                               highlightthickness=2, padx=16, pady=10, cursor="hand2")
        tk.Label(self._frame, text=f"{icon}  {message}", bg=bg, fg=fg,
                 font=("Segoe UI", 10, "bold"), wraplength=450, justify="center").pack()
        self._frame.place(relx=0.5, y=y_off, anchor="n")
        self._frame.lift()
        self._frame.bind("<Button-1>", lambda e: self._dismiss())
        ToastNotification._active.append(self)
        parent.after(duration, self._dismiss)

    def _dismiss(self):
        if self in ToastNotification._active:
            ToastNotification._active.remove(self)
        try:
            self._frame.destroy()
        except Exception:
            pass


# ------------------- Tooltip -------------------

class ToolTip:
    """Hover tooltip for any widget."""
    def __init__(self, widget, text_func=None, text="", delay=400):
        self.widget = widget
        self.text_func = text_func
        self.text = text
        self.delay = delay
        self._tip: Optional[tk.Toplevel] = None
        self._job: Optional[str] = None
        widget.bind("<Enter>", self._schedule, add="+")
        widget.bind("<Leave>", self._hide, add="+")
        widget.bind("<ButtonPress>", self._hide, add="+")

    def _schedule(self, event=None):
        self._hide()
        self._job = self.widget.after(self.delay, self._show)

    def _show(self):
        txt = self.text_func() if self.text_func else self.text
        if not txt:
            return
        x = self.widget.winfo_rootx() + 20
        y = self.widget.winfo_rooty() + self.widget.winfo_height() + 4
        self._tip = tk.Toplevel(self.widget)
        self._tip.wm_overrideredirect(True)
        self._tip.attributes("-topmost", True)
        self._tip.geometry(f"+{x}+{y}")
        frm = tk.Frame(self._tip, bg="#1e293b", padx=10, pady=8,
                        highlightbackground="#475569", highlightthickness=1)
        frm.pack()
        tk.Label(frm, text=txt, bg="#1e293b", fg="#f1f5f9",
                 font=("Segoe UI", 9), justify="right", wraplength=350).pack()

    def _hide(self, event=None):
        if self._job:
            try:
                self.widget.after_cancel(self._job)
            except Exception:
                pass
            self._job = None
        if self._tip:
            try:
                self._tip.destroy()
            except Exception:
                pass
            self._tip = None


# ------------------- Treeview Helpers -------------------

def _bind_mousewheel(widget, scrollable=None):
    """Bind mouse-wheel to *widget* so it scrolls on hover (no click needed).

    Works on Windows (<MouseWheel>) and Linux (<Button-4/5>).
    *scrollable* defaults to *widget* if omitted.
    """
    target = scrollable or widget

    def _on_wheel(event):
        try:
            if hasattr(event, "delta") and event.delta:
                target.yview_scroll(int(-1 * (event.delta / 120)), "units")
            elif getattr(event, "num", 0) == 4:
                target.yview_scroll(-3, "units")
            elif getattr(event, "num", 0) == 5:
                target.yview_scroll(3, "units")
        except Exception:
            pass
        return "break"          # prevent double-scroll from native handler

    def _enter(_e=None):
        widget.bind_all("<MouseWheel>", _on_wheel)
        widget.bind_all("<Button-4>", _on_wheel)
        widget.bind_all("<Button-5>", _on_wheel)

    def _leave(_e=None):
        try:
            widget.unbind_all("<MouseWheel>")
            widget.unbind_all("<Button-4>")
            widget.unbind_all("<Button-5>")
        except Exception:
            pass

    widget.bind("<Enter>", _enter, add="+")
    widget.bind("<Leave>", _leave, add="+")


def _apply_zebra_tags(tree):
    """Apply alternating row colors to a treeview after insert."""
    for i, iid in enumerate(tree.get_children()):
        existing = list(tree.item(iid, "tags") or ())
        existing = [t for t in existing if t not in ("even", "odd")]
        existing.append("even" if i % 2 == 0 else "odd")
        tree.item(iid, tags=tuple(existing))


def _add_context_menu(tree, parent_widget, extra_items=None):
    """Add a right-click context menu to a treeview."""
    menu = tk.Menu(parent_widget, tearoff=0)

    def _copy_cell():
        sel = tree.selection()
        if not sel:
            return
        vals = tree.item(sel[0], "values")
        col = tree.identify_column(menu._click_x)
        try:
            idx = int(col.replace("#", "")) - 1
            text = str(vals[idx]) if 0 <= idx < len(vals) else ""
        except Exception:
            text = " | ".join(str(v) for v in vals)
        try:
            parent_widget.clipboard_clear()
            parent_widget.clipboard_append(text)
        except Exception:
            pass

    def _copy_row():
        sel = tree.selection()
        if not sel:
            return
        vals = tree.item(sel[0], "values")
        text = " | ".join(str(v) for v in vals)
        try:
            parent_widget.clipboard_clear()
            parent_widget.clipboard_append(text)
        except Exception:
            pass

    menu.add_command(label="نسخ الخلية", command=_copy_cell)
    menu.add_command(label="نسخ الصف", command=_copy_row)

    if extra_items:
        menu.add_separator()
        for label, cmd in extra_items:
            menu.add_command(label=label, command=cmd)

    def _on_right_click(event):
        iid = tree.identify_row(event.y)
        if iid:
            tree.selection_set(iid)
        menu._click_x = event.x
        try:
            menu.tk_popup(event.x_root, event.y_root)
        finally:
            menu.grab_release()

    tree.bind("<Button-3>", _on_right_click)
    return menu


# ------------------- Dashboard Frame -------------------

class DashboardFrame(ttk.Frame):
    """Home dashboard with today's summary, recent activity, and alerts."""

    def __init__(self, master, db: SqliteDatabase):
        super().__init__(master, padding=10)
        self.db = db
        self._build()

    def _build(self):
        ttk.Label(self, text="\u2302  \u0644\u0648\u062D\u0629 \u0627\u0644\u062A\u062D\u0643\u0645",
                  font=("Segoe UI", 16, "bold")).pack(anchor="w", pady=(0, 12))

        # Stats cards row
        cards = ttk.Frame(self)
        cards.pack(fill=tk.X, pady=(0, 12))

        self._today_sales = tk.StringVar(value="0.00")
        self._today_bills = tk.StringVar(value="0")
        self._today_reservations = tk.StringVar(value="0")
        self._pending_reservations = tk.StringVar(value="0")
        self._low_stock_count = tk.StringVar(value="0")

        card_defs = [
            ("\u0645\u0628\u064A\u0639\u0627\u062A \u0627\u0644\u064A\u0648\u0645", self._today_sales, "#2563eb"),
            ("\u0641\u0648\u0627\u062A\u064A\u0631 \u0627\u0644\u064A\u0648\u0645", self._today_bills, "#16a34a"),
            ("\u062D\u062C\u0648\u0632\u0627\u062A \u0627\u0644\u064A\u0648\u0645", self._today_reservations, "#f59e0b"),
            ("\u062D\u062C\u0648\u0632\u0627\u062A \u0645\u0639\u0644\u0642\u0629", self._pending_reservations, "#dc2626"),
            ("\u0645\u062E\u0632\u0648\u0646 \u0645\u0646\u062E\u0641\u0636", self._low_stock_count, "#7c3aed"),
        ]
        for col, (title, var, color) in enumerate(card_defs):
            card = tk.Frame(cards, bg="white", highlightbackground=color,
                            highlightthickness=2, padx=18, pady=14, cursor="hand2")
            card.grid(row=0, column=col, padx=6, pady=4, sticky="nsew")
            tk.Label(card, text=title, bg="white", fg="#64748b",
                     font=("Segoe UI", 9)).pack(anchor="center")
            tk.Label(card, textvariable=var, bg="white", fg=color,
                     font=("Segoe UI", 20, "bold")).pack(anchor="center", pady=(4, 0))
        for c in range(5):
            cards.columnconfigure(c, weight=1)

        # Bottom: recent bills + low stock alerts
        bottom = ttk.PanedWindow(self, orient=tk.HORIZONTAL)
        bottom.pack(fill=tk.BOTH, expand=True)

        recent = ttk.LabelFrame(bottom, text="\u0622\u062E\u0631 \u0627\u0644\u0641\u0648\u0627\u062A\u064A\u0631")
        bottom.add(recent, weight=1)
        self._recent_tbl = ttk.Treeview(
            recent, columns=("id", "date", "customer", "total"),
            show="headings", height=8)
        for col, txt, w in [("id", "#", 50), ("date", "\u0627\u0644\u062A\u0627\u0631\u064A\u062E", 140),
                             ("customer", "\u0627\u0644\u0639\u0645\u064A\u0644", 150), ("total", "\u0627\u0644\u0625\u062C\u0645\u0627\u0644\u064A", 90)]:
            self._recent_tbl.heading(col, text=txt)
            self._recent_tbl.column(col, width=w, anchor="center")
        self._recent_tbl.pack(fill=tk.BOTH, expand=True, padx=4, pady=4)
        _bind_mousewheel(self._recent_tbl)

        alerts = ttk.LabelFrame(bottom, text="\u062A\u0646\u0628\u064A\u0647\u0627\u062A \u0627\u0644\u0645\u062E\u0632\u0648\u0646 \u0627\u0644\u0645\u0646\u062E\u0641\u0636")
        bottom.add(alerts, weight=1)
        self._alerts_tbl = ttk.Treeview(
            alerts, columns=("item", "school", "color", "size", "count"),
            show="headings", height=8)
        for col, txt, w in [("item", "\u0627\u0644\u0646\u0648\u0639", 100), ("school", "\u0627\u0644\u0645\u062F\u0631\u0633\u0629", 100),
                             ("color", "\u0627\u0644\u0644\u0648\u0646", 70), ("size", "\u0627\u0644\u0645\u0642\u0627\u0633", 50),
                             ("count", "\u0627\u0644\u0643\u0645\u064A\u0629", 50)]:
            self._alerts_tbl.heading(col, text=txt)
            self._alerts_tbl.column(col, width=w, anchor="center")
        self._alerts_tbl.pack(fill=tk.BOTH, expand=True, padx=4, pady=4)
        _bind_mousewheel(self._alerts_tbl)
        self._alerts_tbl.tag_configure("low", foreground="#b45309")
        self._alerts_tbl.tag_configure("critical", foreground="#ffffff", background="#dc2626")

    def refresh(self):
        today = date.today().isoformat()
        try:
            stats = self.db.get_sales_stats(date_from=today, date_to=today)
            self._today_sales.set(f"{stats['sales_total']:.2f}")
            self._today_bills.set(str(stats["sales_count"]))
            self._today_reservations.set(str(stats["res_count"]))
        except Exception:
            pass
        try:
            pending = self.db.list_reservations(status="\u0645\u0639\u0644\u0642")
            self._pending_reservations.set(str(len(pending)))
        except Exception:
            pass
        try:
            self._recent_tbl.delete(*self._recent_tbl.get_children())
            for b in self.db.list_bills()[:10]:
                self._recent_tbl.insert("", tk.END, values=(
                    b["id"], b["created_at"][:16].replace("T", " "),
                    b.get("customer", ""), f"{float(b['total']):.2f}"))
            _apply_zebra_tags(self._recent_tbl)
        except Exception:
            pass
        try:
            self._alerts_tbl.delete(*self._alerts_tbl.get_children())
            low_count = 0
            for r in self.db.current_inventory({}):
                cnt = int(r.get("count", 0))
                if 0 < cnt <= 5:
                    tag = "critical" if cnt <= 2 else "low"
                    self._alerts_tbl.insert("", tk.END, values=(
                        r["item_type"], r["school"], r["color"], r["size"], cnt),
                        tags=(tag,))
                    low_count += 1
            self._low_stock_count.set(str(low_count))
        except Exception:
            pass


# ------------------- Income Frame -------------------
class IncomeFrame(ttk.Frame):

    def __init__(self, master, db: SqliteDatabase):
        super().__init__(master, padding=10)
        self.db = db
        self._income_r1 = tk.StringVar()
        self._income_r2 = tk.StringVar()
        self._income_has_alpha = tk.BooleanVar(value=False)
        self._staging_lines: List[Dict[str, Any]] = []

        self._build()

    def _build(self):
        ttk.Label(self, text="وارد (إضافة أصناف جديدة)", font=("Segoe UI", 12, "bold")).pack(anchor="w", pady=(0, 8))

        self._income_r1.trace_add("write", lambda *_: self._rebuild_sizes_grid())
        self._income_r2.trace_add("write", lambda *_: self._rebuild_sizes_grid())
        self._income_has_alpha.trace_add("write", lambda *_: self._rebuild_sizes_grid())

        # Main horizontal split
        vsplit = ttk.PanedWindow(self, orient=tk.HORIZONTAL)
        vsplit.pack(fill=tk.BOTH, expand=True)

        # LEFT side: existing form
        left = ttk.Frame(vsplit)
        vsplit.add(left, weight=6)

        grid = ttk.LabelFrame(left, text="بيانات الصنف")
        grid.pack(fill=tk.BOTH, expand=True)

        self.item_type = LabeledCombobox(grid, "النوع", self.db, "item_type")
        self.school    = LabeledCombobox(grid, "المدرسة", self.db, "school")
        self.color     = LabeledCombobox(grid, "اللون", self.db, "color")

        self.item_type.set_supplier(lambda: self.db.get_distinct_filtered("item_type", self._build_income_filter_constraints("item_type")))
        self.school.set_supplier(lambda: self.db.get_distinct_filtered("school", self._build_income_filter_constraints("school")))
        self.color.set_supplier(lambda: self.db.get_distinct_filtered("color", self._build_income_filter_constraints("color")))

        ranges_box = ttk.LabelFrame(grid, text="نطاقات المقاسات (للوارد فقط)")
        ranges_box.grid(row=2, column=0, columnspan=2, sticky="ew", padx=6, pady=6)

        ttk.Label(ranges_box, text="النطاق الأول").grid(row=0, column=0, sticky="w")
        ttk.Combobox(
            ranges_box,
            textvariable=self._income_r1,
            values=[""] + NUMERIC_RANGE_LABELS,
            state="readonly",
            width=18
        ).grid(row=0, column=1, sticky="w", padx=4)

        ttk.Label(ranges_box, text="النطاق الثاني").grid(row=1, column=0, sticky="w")
        ttk.Combobox(
            ranges_box,
            textvariable=self._income_r2,
            values=[""] + NUMERIC_RANGE_LABELS,
            state="readonly",
            width=18
        ).grid(row=1, column=1, sticky="w", padx=4)

        ttk.Checkbutton(
            ranges_box,
            text="مقاسات حرفية (S إلى 5XL)",
            variable=self._income_has_alpha
        ).grid(row=2, column=0, columnspan=2, sticky="w")

        for w in (self.item_type.cb, self.school.cb, self.color.cb):
            w.bind("<<ComboboxSelected>>", lambda e: (self._on_income_filter_changed(), self._auto_load_size_profile(), self._auto_fill_price_for_grid()), add="+")
            w.bind("<FocusOut>",           lambda e: (self._auto_load_size_profile(), self._auto_fill_price_for_grid()), add="+")

        sizes_container = ttk.Frame(grid)
        sizes_container.grid(row=3, column=0, columnspan=2, sticky="nsew", padx=6, pady=6)

        sizes_canvas = tk.Canvas(sizes_container, highlightthickness=0)
        sizes_scroll = ttk.Scrollbar(sizes_container, orient="vertical", command=sizes_canvas.yview)
        sizes_canvas.configure(yscrollcommand=sizes_scroll.set)

        sizes_scroll.pack(side="right", fill="y")
        sizes_canvas.pack(side="left", fill="both", expand=True)

        self._sizes_inner = ttk.Frame(sizes_canvas)
        sizes_window = sizes_canvas.create_window((0, 0), window=self._sizes_inner, anchor="nw")

        self.sizes_grid = None

        def _on_sizes_config(event=None):
            try:
                sizes_canvas.configure(scrollregion=sizes_canvas.bbox("all"))
                sizes_canvas.itemconfigure(sizes_window, width=sizes_canvas.winfo_width())
            except Exception:
                pass

        self._sizes_inner.bind("<Configure>", _on_sizes_config)
        sizes_canvas.bind("<Configure>", _on_sizes_config)

        _bind_mousewheel(sizes_canvas)

        self.item_type.grid(row=0, column=0, padx=6, pady=6, sticky="ew")
        self.school.grid(   row=0, column=1, padx=6, pady=6, sticky="ew")
        self.color.grid(    row=1, column=0, padx=6, pady=6, sticky="ew")

        for c in range(2):
            grid.columnconfigure(c, weight=1)

        try:
            grid.rowconfigure(3, weight=1)
            grid.rowconfigure(4, weight=0)
            grid.rowconfigure(5, weight=0)
        except Exception:
            pass

        add_btns = ttk.Frame(grid)
        add_btns.grid(row=5, column=0, columnspan=2, sticky="ew", padx=6, pady=(6, 6))

        ttk.Button(add_btns, text="إضافة للقائمة", command=self._add_to_staging).pack(side=tk.LEFT)
        ttk.Button(add_btns, text="تفريغ", command=self._on_reset_keep_pkg).pack(side=tk.LEFT, padx=8)

        # RIGHT side: staging area
        right = ttk.LabelFrame(vsplit, text="فاتورة الوارد")
        vsplit.add(right, weight=4)

        # Supplier field
        sup_row = ttk.Frame(right)
        sup_row.pack(fill=tk.X, padx=4, pady=(4, 2))
        ttk.Label(sup_row, text="المورد:").pack(side=tk.LEFT)
        self._supplier_var = tk.StringVar()
        ttk.Entry(sup_row, textvariable=self._supplier_var).pack(side=tk.LEFT, fill=tk.X, expand=True, padx=(4, 0))

        # Staging treeview
        tree_wrap = ttk.Frame(right)
        tree_wrap.pack(fill=tk.BOTH, expand=True, padx=4, pady=2)
        self._staging_table = ttk.Treeview(
            tree_wrap,
            columns=("type", "school", "color", "size", "price", "qty", "total"),
            show="headings", height=10,
        )
        for col, txt, w in [
            ("type", "النوع", 100), ("school", "المدرسة", 100), ("color", "اللون", 70),
            ("size", "المقاس", 60), ("price", "السعر", 70), ("qty", "الكمية", 60), ("total", "الإجمالي", 80),
        ]:
            self._staging_table.heading(col, text=txt)
            self._staging_table.column(col, width=w, anchor="center")
        ysb = ttk.Scrollbar(tree_wrap, orient="vertical", command=self._staging_table.yview)
        self._staging_table.configure(yscrollcommand=ysb.set)
        self._staging_table.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        ysb.pack(side=tk.RIGHT, fill=tk.Y)
        _bind_mousewheel(self._staging_table)

        # Context menu on staging table
        _add_context_menu(self._staging_table, self, extra_items=[
            ("حذف السطر", self._remove_staging_line),
        ])

        # Controls
        ctrl_row = ttk.Frame(right)
        ctrl_row.pack(fill=tk.X, padx=4, pady=2)
        ttk.Button(ctrl_row, text="حذف السطر", command=self._remove_staging_line).pack(side=tk.LEFT, padx=2)
        ttk.Button(ctrl_row, text="تفريغ القائمة", command=self._clear_staging).pack(side=tk.LEFT, padx=2)

        # Total
        tot_row = ttk.Frame(right)
        tot_row.pack(fill=tk.X, padx=4, pady=2)
        ttk.Label(tot_row, text="الإجمالي:").pack(side=tk.RIGHT)
        self._staging_total_var = tk.StringVar(value="0.00")
        ttk.Label(tot_row, textvariable=self._staging_total_var, font=("Segoe UI", 12, "bold")).pack(side=tk.RIGHT, padx=(0, 6))

        self._staging_qty_var = tk.StringVar(value="0")
        ttk.Label(tot_row, text="القطع:").pack(side=tk.RIGHT, padx=(12, 0))
        ttk.Label(tot_row, textvariable=self._staging_qty_var, font=("Segoe UI", 10)).pack(side=tk.RIGHT, padx=(0, 6))

        # Finalize button
        ttk.Button(right, text="تأكيد فاتورة الوارد", command=self._finalize_income).pack(fill=tk.X, padx=4, pady=(4, 4))

    def _build_income_filter_constraints(self, exclude_field: str) -> Dict[str, Any]:
        c: Dict[str, Any] = {}
        if exclude_field != "school" and hasattr(self, 'school'):
            v = self.school.get()
            if v:
                c["school"] = v
        if exclude_field != "item_type" and hasattr(self, 'item_type'):
            v = self.item_type.get()
            if v:
                c["item_type"] = v
        if exclude_field != "color" and hasattr(self, 'color'):
            v = self.color.get()
            if v:
                c["color"] = v
        return c

    def _on_income_filter_changed(self, *_):
        self.item_type.refresh_values()
        self.school.refresh_values()
        self.color.refresh_values()

    def _auto_load_size_profile(self):
        """Auto-load size ranges from saved profile when all 3 specs are selected."""
        it = (self.item_type.get() or "").strip()
        sc = (self.school.get() or "").strip()
        cl = (self.color.get() or "").strip()
        if not (it and sc and cl):
            return
        profile = self.db.get_size_profile(it, sc, cl)
        if profile:
            r1s, r1e, r2s, r2e, has_alpha = profile
            r1_label = ""
            r2_label = ""
            for (a, b) in ALLOWED_NUMERIC_RANGES:
                if a == r1s and b == r1e:
                    r1_label = f"{a} \u2192 {b}"
                if a == r2s and b == r2e:
                    r2_label = f"{a} \u2192 {b}"
            self._income_r1.set(r1_label)
            self._income_r2.set(r2_label)
            self._income_has_alpha.set(bool(has_alpha))
        else:
            self._rebuild_sizes_grid()

    def _rebuild_sizes_grid(self):
        try:
            self.sizes_grid.destroy()
        except Exception:
            pass

        numeric_ranges = []

        def _parse_range(label: str):
            if not label:
                return None
            a, b = label.split("\u2192")
            return int(a.strip()), int(b.strip())

        for lbl in (self._income_r1.get(), self._income_r2.get()):
            pair = _parse_range(lbl)
            if pair:
                numeric_ranges.append(pair)

        numeric_ranges.sort()
        merged = []

        for start, end in numeric_ranges:
            if not merged:
                merged.append([start, end])
            else:
                last_start, last_end = merged[-1]
                if start <= last_end:
                    merged[-1][1] = max(last_end, end)
                else:
                    merged.append([start, end])

        sizes = []
        for start, end in merged:
            sizes.extend(str(x) for x in range(start, end + 1, 2))

        if self._income_has_alpha.get():
            sizes.extend(ALPHA_SIZES)

        self.sizes_grid = SizesGrid(self._sizes_inner, sizes)
        self.sizes_grid.pack(fill="both", expand=False)

        self._sizes_inner.update_idletasks()

    def _auto_fill_price_for_grid(self):
        it = (self.item_type.get() or "").strip()
        sc = (self.school.get() or "").strip()
        cl = (self.color.get() or "").strip()
        if not (it and sc and cl):
            return
        try:
            p = self.db.last_price_for_specs(it, sc, cl, "")
            if p is None:
                try:
                    sizes = self.db.list_sizes_for_item(sc, it, cl)
                except Exception:
                    sizes = []
                if sizes:
                    for row in sizes:
                        size_label = str(row.get("size") or "")
                        last_p = row.get("last_price")
                        if last_p is not None:
                            self.sizes_grid.set_price_for_size(size_label, last_p)
                    return
            if p is not None:
                self.sizes_grid.set_price_for_all(p)
        except Exception:
            pass

    def _on_reset_keep_pkg(self):
        for w in (self.item_type, self.school, self.color):
            w.set("")
        self._income_r1.set("")
        self._income_r2.set("")
        self._income_has_alpha.set(False)

        if self.sizes_grid:
            self.sizes_grid.destroy()
            self.sizes_grid = None

        self._sizes_inner.update_idletasks()

    def _add_to_staging(self):
        """Add current form items to the staging list (not yet saved to DB)."""
        item_type = (self.item_type.get() or "").strip()
        school = (self.school.get() or "").strip()
        color = (self.color.get() or "").strip()
        if not item_type or not school or not color:
            messagebox.showwarning("بيانات ناقصة", "يجب اختيار النوع والمدرسة واللون.", parent=self)
            return

        if self.sizes_grid is None:
            messagebox.showwarning("فارغ", "لم تُدخل أي كميات في شبكة المقاسات.", parent=self)
            return

        rows = self.sizes_grid.get_rows()
        added = 0
        for r in rows:
            sz = str(r["size"]).strip()
            qty = int(r["qty"])
            if qty <= 0:
                continue

            price = r["price"]
            if price is None or float(price) <= 0:
                try:
                    price = float(self.db.get_effective_price(item_type, school, color, sz) or 0)
                except Exception:
                    price = 0.0
            else:
                price = float(price)

            # Merge with existing staging line if same specs
            merged = False
            for line in self._staging_lines:
                if (line["item_type"] == item_type and line["school"] == school
                        and line["color"] == color and line["size"] == sz
                        and abs(line["unit_price"] - price) < 0.001):
                    line["qty"] += qty
                    merged = True
                    break

            if not merged:
                self._staging_lines.append({
                    "item_type": item_type,
                    "school": school,
                    "color": color,
                    "size": sz,
                    "unit_price": price,
                    "qty": qty,
                })
            added += 1

        if added == 0:
            messagebox.showwarning("لا توجد كميات", "أدخل كمية واحدة على الأقل.", parent=self)
            return

        self._sync_staging_table()
        # Clear the size grid quantities but keep the prices
        self.sizes_grid.clear_quantities()

    def _sync_staging_table(self):
        """Refresh the staging treeview from in-memory list."""
        self._staging_table.delete(*self._staging_table.get_children())
        total_value = 0.0
        total_qty = 0
        for idx, ln in enumerate(self._staging_lines):
            line_total = float(ln["unit_price"]) * int(ln["qty"])
            total_value += line_total
            total_qty += int(ln["qty"])
            self._staging_table.insert(
                "", tk.END, iid=str(idx),
                values=(ln["item_type"], ln["school"], ln["color"], ln["size"],
                        f"{float(ln['unit_price']):.2f}", ln["qty"], f"{line_total:.2f}")
            )
        self._staging_total_var.set(f"{total_value:.2f}")
        self._staging_qty_var.set(str(total_qty))
        _apply_zebra_tags(self._staging_table)

    def _remove_staging_line(self):
        sel = self._staging_table.selection()
        if not sel:
            return
        idx = int(sel[0])
        if 0 <= idx < len(self._staging_lines):
            self._staging_lines.pop(idx)
            self._sync_staging_table()

    def _clear_staging(self):
        self._staging_lines.clear()
        self._sync_staging_table()

    def _finalize_income(self):
        if not self.db.is_manager_feature_enabled("allow_manual_incoming"):
            messagebox.showwarning("مقيد", _feature_restricted_message("فاتورة الوارد اليدوية متوقفة حالياً في نقطة البيع."), parent=self)
            return
        if not self._staging_lines:
            messagebox.showwarning("فارغ", "لا توجد أصناف في فاتورة الوارد.", parent=self)
            return

        supplier = self._supplier_var.get().strip()

        try:
            bill_id = self.db.create_income_bill(supplier, self._staging_lines)
            total = self._staging_total_var.get()
            qty = self._staging_qty_var.get()
            ToastNotification.show(self.winfo_toplevel(),
                f"تم إنشاء فاتورة وارد #{bill_id} — القطع: {qty} — الإجمالي: {total}",
                toast_type="success")
            self._staging_lines.clear()
            self._sync_staging_table()
            self._supplier_var.set("")
        except Exception as ex:
            messagebox.showerror("خطأ", str(ex), parent=self)

    @staticmethod
    def _err(msg: str) -> None:
        raise RuntimeError(msg)


# ------------------- POS Frame (replaces OutcomeFrame) -------------------

class POSFrame(ttk.Frame):
    """
    Multi-bill POS with step-by-step item navigation.
    Bills 1-3 are regular sales; bill 4 is returns; bill 5 is reservations; bill 6 is exchanges.
    """
    def __init__(self, master, db: SqliteDatabase):
        super().__init__(master, padding=6)
        self.db = db

        # In-memory bill state
        self.bills = {1: [], 2: [], 3: [], 4: [], 5: [], 6: []}
        self.customers = {1: "", 2: "", 3: "", 4: "", 5: "", 6: ""}
        self.active_bill = 1
        self._exchange_mode = "return"  # "return" or "take" — for bill 6

        # Navigation state
        self._sel_school: Optional[str] = None
        self._sel_item: Optional[str] = None
        self._sel_color: Optional[str] = None
        self._sel_size: Optional[str] = None
        self._sizes_cache: List[Dict[str, Any]] = []
        self._size_btns: Dict[str, Any] = {}
        self._selected_size_btn = None
        self._price_user_edited = False

        self._build()

    # ------------------------------------------------------------------ build
    def _build(self):
        # Define styles once
        try:
            s = ttk.Style(self)
            s.configure("Size.TButton", padding=6)
            s.configure("SizeZero.TButton",
                        background="#f8d7da", foreground="#000000",
                        bordercolor="#b91c1c", focusthickness=1, focuscolor="#b91c1c", padding=6)
            s.map("SizeZero.TButton",
                background=[("active", "#f1b0b7"), ("pressed", "#f1b0b7")],
                bordercolor=[("focus", "#b91c1c")])
            s.configure("SizeSelected.TButton",
                        background="#d1fae5", foreground="#000000",
                        bordercolor="#059669", focusthickness=1, focuscolor="#059669", padding=6)
            s.map("SizeSelected.TButton",
                background=[("active", "#a7f3d0"), ("pressed", "#a7f3d0")],
                bordercolor=[("focus", "#059669")])
            s.configure("SizeSelectedZero.TButton",
                        background="#fbd4a8", foreground="#000000",
                        bordercolor="#b45309", focusthickness=1, focuscolor="#b45309", padding=6)
            s.map("SizeSelectedZero.TButton",
                background=[("active", "#f8c58a"), ("pressed", "#f8c58a")],
                bordercolor=[("focus", "#b45309")])
            # Low stock style (count 1-5) — amber/yellow warning
            s.configure("SizeLow.TButton",
                        background="#fef3c7", foreground="#000000",
                        bordercolor="#d97706", focusthickness=1, focuscolor="#d97706", padding=6)
            s.map("SizeLow.TButton",
                background=[("active", "#fde68a"), ("pressed", "#fde68a")],
                bordercolor=[("focus", "#d97706")])
            s.configure("ActiveBill.TButton",
                        background="#059669", foreground="#ffffff",
                        bordercolor="#047857", padding=6)
            s.map("ActiveBill.TButton",
                background=[("active", "#047857"), ("pressed", "#047857")])
        except Exception:
            pass

        # Main horizontal split: left = selector, right = bill panel
        vsplit = ttk.PanedWindow(self, orient=tk.HORIZONTAL)
        vsplit.pack(fill=tk.BOTH, expand=True)

        # ---- LEFT: item selector ----
        left = ttk.Frame(vsplit)
        vsplit.add(left, weight=6)

        # Universal search bar
        search_row = ttk.Frame(left)
        search_row.pack(fill=tk.X, padx=4, pady=(4, 2))
        ttk.Label(search_row, text="\u0628\u062D\u062B \u0633\u0631\u064A\u0639:", font=("Segoe UI", 9)).pack(side=tk.LEFT, padx=(0, 4))
        self._search_var = tk.StringVar()
        self._search_entry = ttk.Entry(search_row, textvariable=self._search_var, width=30)
        self._search_entry.pack(side=tk.LEFT, fill=tk.X, expand=True)
        self._search_entry.bind("<KeyRelease>", lambda e: self._on_search_changed())
        ttk.Button(search_row, text="\u2716", width=3, command=self._clear_search,
                   style="Secondary.TButton").pack(side=tk.LEFT, padx=(4, 0))

        # Quick filter bar
        fbar = ttk.LabelFrame(left, text="\u062A\u0635\u0641\u064A\u0629 \u0633\u0631\u064A\u0639\u0629")
        fbar.pack(fill=tk.X, padx=4, pady=(2, 2))

        self._flt_school = LabeledCombobox(fbar, "\u0627\u0644\u0645\u062F\u0631\u0633\u0629", self.db, "school")
        self._flt_school.set_supplier(lambda: self.db.get_distinct_filtered("school", self._build_filter_constraints("school")))
        self._flt_school.grid(row=0, column=0, padx=4, pady=4, sticky="ew")

        self._flt_item = LabeledCombobox(fbar, "\u0627\u0644\u0646\u0648\u0639", self.db, "item_type")
        self._flt_item.set_supplier(lambda: self.db.get_distinct_filtered("item_type", self._build_filter_constraints("item_type")))
        self._flt_item.grid(row=0, column=1, padx=4, pady=4, sticky="ew")

        self._flt_color = LabeledCombobox(fbar, "\u0627\u0644\u0644\u0648\u0646", self.db, "color")
        self._flt_color.set_supplier(lambda: self.db.get_distinct_filtered("color", self._build_filter_constraints("color")))
        self._flt_color.grid(row=0, column=2, padx=4, pady=4, sticky="ew")

        ttk.Button(fbar, text="\u0645\u0633\u062D", command=self._clear_quick_filter,
                   style="Secondary.TButton").grid(row=0, column=3, padx=4, pady=4)

        for c in range(4):
            fbar.columnconfigure(c, weight=1 if c < 3 else 0)

        # Auto-apply filters when a value is selected from the dropdown
        self._flt_school.cb.bind("<<ComboboxSelected>>", lambda e: self._on_filter_changed(), add="+")
        self._flt_item.cb.bind("<<ComboboxSelected>>", lambda e: self._on_filter_changed(), add="+")
        self._flt_color.cb.bind("<<ComboboxSelected>>", lambda e: self._on_filter_changed(), add="+")

        # Breadcrumb
        crumb = ttk.Frame(left)
        crumb.pack(fill=tk.X, padx=4, pady=(2, 0))
        self._crumb_var = tk.StringVar(value="اختر المدرسة")
        ttk.Label(crumb, textvariable=self._crumb_var, font=("", 10, "bold")).pack(side=tk.LEFT, padx=(2, 6))

        # Scrollable button grid
        sc = ttk.Frame(left)
        sc.pack(fill=tk.BOTH, expand=True, pady=(2, 0))

        self._canvas = tk.Canvas(sc, highlightthickness=0, bd=0)
        self._scroll_y = ttk.Scrollbar(sc, orient="vertical", command=self._canvas.yview)
        self._canvas.configure(yscrollcommand=self._scroll_y.set)

        try:
            bg = ttk.Style(self).lookup("TFrame", "background") or "#ffffff"
            self._canvas.configure(background=bg)
        except Exception:
            pass

        self._grid_host = ttk.Frame(self._canvas)
        self._grid_window = self._canvas.create_window((0, 0), window=self._grid_host, anchor="nw")

        self._canvas.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        self._scroll_y.pack(side=tk.RIGHT, fill=tk.Y)

        def _on_configure(_e=None):
            self._canvas.configure(scrollregion=self._canvas.bbox("all"))
            try:
                self._canvas.itemconfigure(self._grid_window, width=self._canvas.winfo_width())
            except Exception:
                pass

        self._grid_host.bind("<Configure>", _on_configure)
        self._canvas.bind("<Configure>", _on_configure)

        # Scroll wheel support for the button grid
        _bind_mousewheel(self._canvas)

        # Favorites section (most-sold items)
        self._fav_frame = ttk.LabelFrame(left, text="\u2605 \u0627\u0644\u0623\u0643\u062B\u0631 \u0645\u0628\u064A\u0639\u0627\u064B")
        self._fav_frame.pack(fill=tk.X, padx=4, pady=(2, 0))
        self._fav_inner = ttk.Frame(self._fav_frame)
        self._fav_inner.pack(fill=tk.X, padx=4, pady=4)
        self._refresh_favorites()

        # ---- RIGHT: bill management ----
        right = ttk.LabelFrame(vsplit, text="\u0627\u0644\u0641\u0648\u0627\u062A\u064A\u0631")
        vsplit.add(right, weight=4)

        # Bill switcher buttons (1-6)
        switcher = ttk.Frame(right)
        switcher.pack(fill=tk.X, padx=4, pady=(4, 2))
        self._bill_btns: Dict[int, ttk.Button] = {}
        _bill_labels = {
            1: "1", 2: "2", 3: "3",
            4: "4 (مرتجع)", 5: "5 (حجز)", 6: "6 (استبدال)",
        }
        for n in range(1, 7):
            btn = ttk.Button(switcher, text=_bill_labels[n], width=9,
                             command=lambda b=n: self._switch_bill(b))
            btn.pack(side=tk.LEFT, padx=2)
            self._bill_btns[n] = btn
        self._update_bill_btn_styles()

        # Customer field
        self._customer_var = tk.StringVar()
        cust_row = ttk.Frame(right)
        cust_row.pack(fill=tk.X, padx=4, pady=(2, 2))
        ttk.Label(cust_row, text="العميل:").pack(side=tk.LEFT)
        self._cust_entry = ttk.Entry(cust_row, textvariable=self._customer_var)
        self._cust_entry.pack(side=tk.LEFT, fill=tk.X, expand=True, padx=(4, 0))
        _wh_btn = ttk.Button(
            cust_row,
            text=WAREHOUSE_RETURN_LABEL,
            command=lambda: self._customer_var.set(WAREHOUSE_RETURN_LABEL),
            style="Secondary.TButton",
        )
        _wh_btn.pack(side=tk.LEFT, padx=(4, 0))
        ToolTip(_wh_btn, "استخدم الفاتورة الحالية كإرجاع مخزون من الفرع إلى المصنع / المخزن الرئيسي")
        _branch_btn = ttk.Button(
            cust_row,
            text="طلب تحويل",
            command=self._choose_branch_target,
            style="Secondary.TButton",
        )
        _branch_btn.pack(side=tk.LEFT, padx=(4, 0))
        ToolTip(_branch_btn, "سجّل طلب تحويل إلى فرع آخر. التنفيذ الفعلي يتم من خلال المخزن فقط")
        self._cust_entry.bind("<FocusOut>", lambda e: self._save_customer(), add="+")
        self._cust_entry.bind("<KeyRelease>", lambda e: self._autocomplete_customer(), add="+")
        self._cust_listbox = None

        # Bill items treeview
        tree_wrap = ttk.Frame(right)
        tree_wrap.pack(fill=tk.BOTH, expand=True, padx=4, pady=(2, 2))
        self.bill_table = ttk.Treeview(
            tree_wrap,
            columns=("type", "school", "color", "size", "price", "qty", "total"),
            show="headings",
            height=8,
        )
        for col, txt, w in [
            ("type", "النوع", 100), ("school", "المدرسة", 120), ("color", "اللون", 70),
            ("size", "المقاس", 60), ("price", "السعر", 70), ("qty", "الكمية", 60), ("total", "الإجمالي", 80),
        ]:
            self.bill_table.heading(col, text=txt)
            self.bill_table.column(col, width=w, anchor="center")
        ysb = ttk.Scrollbar(tree_wrap, orient="vertical", command=self.bill_table.yview)
        self.bill_table.configure(yscrollcommand=ysb.set)
        self.bill_table.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        ysb.pack(side=tk.RIGHT, fill=tk.Y)
        _bind_mousewheel(self.bill_table)

        # Context menu on bill table
        _add_context_menu(self.bill_table, self, extra_items=[
            ("حذف السطر", self._remove_line),
        ])

        # Qty controls
        qty_row = ttk.Frame(right)
        qty_row.pack(fill=tk.X, padx=4, pady=(2, 2))
        ttk.Button(qty_row, text="-", width=3, command=self._dec_qty).pack(side=tk.LEFT)
        ttk.Button(qty_row, text="+", width=3, command=self._inc_qty).pack(side=tk.LEFT, padx=2)
        ttk.Button(qty_row, text="حذف السطر", command=self._remove_line).pack(side=tk.LEFT, padx=4)
        ttk.Button(qty_row, text="تفريغ", command=self._clear_bill).pack(side=tk.LEFT, padx=4)

        # Total display
        tot_row = ttk.Frame(right)
        tot_row.pack(fill=tk.X, padx=4, pady=(2, 2))
        ttk.Label(tot_row, text="الإجمالي:").pack(side=tk.RIGHT)
        self.total_var = tk.StringVar(value="0.00")
        ttk.Label(tot_row, textvariable=self.total_var, font=("Segoe UI", 12, "bold")).pack(side=tk.RIGHT, padx=(0, 6))

        # Reservation extras (only visible for bill 5)
        self._res_frame = ttk.LabelFrame(right, text="بيانات الحجز")
        self._res_frame.pack(fill=tk.X, padx=4, pady=(2, 2))
        ttk.Label(self._res_frame, text="المبلغ المدفوع:").grid(row=0, column=0, padx=4, pady=4, sticky="e")
        self._paid_var = tk.StringVar(value="0")
        ttk.Entry(self._res_frame, textvariable=self._paid_var, width=12).grid(row=0, column=1, padx=4, pady=4, sticky="w")
        ttk.Label(self._res_frame, text="ملاحظة:").grid(row=1, column=0, padx=4, pady=4, sticky="e")
        self._res_note_var = tk.StringVar()
        ttk.Entry(self._res_frame, textvariable=self._res_note_var, width=24).grid(row=1, column=1, padx=4, pady=4, sticky="ew")
        self._res_frame.columnconfigure(1, weight=1)

        # Return extras (only visible for bill 4)
        self._return_frame = ttk.LabelFrame(right, text="بيانات المرتجع")
        ttk.Label(self._return_frame, text="أضف الأصناف المراد إرجاعها",
                  font=("Segoe UI", 9, "italic")).pack(padx=8, pady=4, anchor="w")

        # Exchange extras (only visible for bill 6)
        self._exchange_frame = ttk.LabelFrame(right, text="بيانات الاستبدال")
        ex_mode_row = ttk.Frame(self._exchange_frame)
        ex_mode_row.pack(fill=tk.X, padx=8, pady=4)
        self._exchange_mode_var = tk.StringVar(value="return")
        ttk.Radiobutton(ex_mode_row, text="↩ إضافة مرتجع", variable=self._exchange_mode_var,
                        value="return", command=self._on_exchange_mode_changed).pack(side=tk.LEFT, padx=(0, 12))
        ttk.Radiobutton(ex_mode_row, text="↪ إضافة مأخوذ", variable=self._exchange_mode_var,
                        value="take", command=self._on_exchange_mode_changed).pack(side=tk.LEFT)
        self._exchange_diff_var = tk.StringVar(value="فرق السعر: 0.00")
        ttk.Label(self._exchange_frame, textvariable=self._exchange_diff_var,
                  font=("Segoe UI", 10, "bold")).pack(padx=8, pady=(0, 4), anchor="w")

        # Finalize button with shortcut hint
        ttk.Button(right, text="تأكيد الفاتورة / الحجز (F8)", command=self._finalize,
                   style="Success.TButton").pack(fill=tk.X, padx=4, pady=(4, 4))

        self._update_res_frame_visibility()

        # Initial render - start with schools
        self._render_schools()

    # ------------------------------------------------------------------ helpers
    def _clear_grid(self):
        for w in self._grid_host.winfo_children():
            w.destroy()
        self._size_btns.clear()
        self._selected_size_btn = None

    def _mk_grid_buttons(self, labels: List[str], on_click, *, cols: int = 5):
        if not labels:
            ttk.Label(self._grid_host, text="لا توجد بيانات.").pack(pady=8)
            return
        row = None
        for i, label in enumerate(labels):
            if i % cols == 0:
                row = ttk.Frame(self._grid_host)
                row.pack(fill=tk.X, padx=4, pady=2)
            btn = ttk.Button(row, text=label, command=lambda v=label: on_click(v))
            btn.pack(side=tk.LEFT, padx=3, pady=3)

    # ------------------------------------------------------------------ navigation
    # Flow: Schools → Items → Colors → Sizes

    def _render_schools(self):
        """Entry point: show school buttons."""
        self._sel_school = None
        self._sel_item = None
        self._sel_color = None
        self._sel_size = None
        self._crumb_var.set("اختر المدرسة")
        self._clear_grid()

        school_f = self._flt_school.get() or None
        item_f = self._flt_item.get() or None
        color_f = self._flt_color.get() or None

        constraints: Dict[str, Any] = {}
        if item_f:
            constraints["item_type"] = item_f
        if color_f:
            constraints["color"] = color_f

        try:
            schools = sorted(self.db.get_distinct_filtered("school", constraints))
        except Exception:
            schools = []

        if school_f:
            schools = [s for s in schools if s == school_f]

        self._mk_grid_buttons(schools, self._select_school, cols=4)

    def _select_school(self, school: str):
        self._sel_school = school
        self._price_user_edited = False
        self._render_items()

    def _render_items(self):
        """Show item type buttons for the selected school."""
        self._sel_item = None
        self._sel_color = None
        self._sel_size = None
        self._crumb_var.set(f"المدرسة: {self._sel_school}  \u27f6  اختر النوع")
        self._clear_grid()

        item_f = self._flt_item.get() or None
        color_f = self._flt_color.get() or None

        constraints: Dict[str, Any] = {"school": self._sel_school}
        if color_f:
            constraints["color"] = color_f

        try:
            items = sorted(self.db.get_distinct_filtered("item_type", constraints))
        except Exception:
            items = []

        if item_f:
            items = [i for i in items if i == item_f]

        self._mk_grid_buttons(items, self._select_item, cols=4)
        ttk.Button(self._grid_host, text="\u25c4 رجوع إلى المدارس", command=self._render_schools)\
            .pack(anchor="w", padx=4, pady=4)

    def _select_item(self, item_type: str):
        self._sel_item = item_type
        self._price_user_edited = False
        self._render_colors()

    def _render_colors(self):
        """Show color buttons for selected school + item."""
        self._sel_color = None
        self._sel_size = None
        self._crumb_var.set(f"المدرسة: {self._sel_school}  \u27f6  النوع: {self._sel_item}  \u27f6  اختر اللون")
        self._clear_grid()

        color_f = self._flt_color.get() or None
        try:
            pairs = self.db.list_items_for_school(self._sel_school or "")
            colors = sorted({cl for (it, cl) in pairs if it == self._sel_item})
        except Exception:
            colors = []

        if color_f:
            colors = [c for c in colors if c == color_f]

        self._mk_grid_buttons(colors, self._select_color, cols=4)
        ttk.Button(self._grid_host, text="\u25c4 رجوع إلى الأنواع", command=self._render_items)\
            .pack(anchor="w", padx=4, pady=4)

    def _select_color(self, color: str):
        self._sel_color = color
        self._price_user_edited = False
        self._render_sizes()

    @staticmethod
    def _spec_match(a: str, b: str) -> bool:
        return (a or "").strip().casefold() == (b or "").strip().casefold()

    def _pending_out_qty_for_specs(self, school: str, item: str, color: str, size: str) -> int:
        """Qty on the active bill that will deduct shop stock when finalized (not yet in DB)."""
        b = self.active_bill
        lines = self.bills.get(b) or []
        total = 0
        if b in (1, 2, 3):
            for ln in lines:
                if not self._spec_match(ln.get("school"), school):
                    continue
                if not self._spec_match(ln.get("item_type"), item):
                    continue
                if not self._spec_match(ln.get("color"), color):
                    continue
                if not self._spec_match(ln.get("size"), size):
                    continue
                try:
                    total += int(ln.get("qty") or 0)
                except Exception:
                    pass
        elif b == 6:
            for ln in lines:
                if ln.get("direction") != "take":
                    continue
                if not self._spec_match(ln.get("school"), school):
                    continue
                if not self._spec_match(ln.get("item_type"), item):
                    continue
                if not self._spec_match(ln.get("color"), color):
                    continue
                if not self._spec_match(ln.get("size"), size):
                    continue
                try:
                    total += int(ln.get("qty") or 0)
                except Exception:
                    pass
        return total

    def _refresh_size_grid_if_current(self, *, preserve_size: Optional[str] = None) -> None:
        """Re-query stock and rebuild size buttons when the grid is visible."""
        if self._sel_school and self._sel_item and self._sel_color:
            try:
                self._render_sizes(preserve_size=preserve_size)
            except Exception:
                pass

    def _render_sizes(self, preserve_size: Optional[str] = None):
        self._sel_size = None
        self._crumb_var.set(
            f"المدرسة: {self._sel_school}  \u27f6  النوع: {self._sel_item}  \u27f6  اللون: {self._sel_color}  \u27f6  اختر المقاس"
        )
        self._clear_grid()

        raw_sizes = self._get_sizes_for_bill(
            self._sel_school or "", self._sel_item or "", self._sel_color or "")

        stock_rows: Dict[str, Any] = {}
        try:
            for r in self.db.current_inventory({
                "school": self._sel_school,
                "item_type": self._sel_item,
                "color": self._sel_color,
            }):
                sz = str(r.get("size") or "").strip()
                if not sz:
                    continue
                stock_rows.setdefault(sz, {"count": 0, "last_price": r.get("unit_price")})
                stock_rows[sz]["count"] += int(r.get("count") or 0)
        except Exception:
            pass

        self._sizes_cache = []
        for sz in raw_sizes:
            r = stock_rows.get(sz, {})
            base = int(r.get("count", 0) or 0)
            pending = self._pending_out_qty_for_specs(
                self._sel_school or "", self._sel_item or "", self._sel_color or "", str(sz))
            eff = max(0, base - pending)
            self._sizes_cache.append({
                "size": sz,
                "count": eff,
                "last_price": r.get("last_price"),
            })

        cols = 4
        row = None
        for i, r in enumerate(self._sizes_cache):
            if i % cols == 0:
                row = ttk.Frame(self._grid_host)
                row.pack(fill=tk.X, padx=4, pady=2)

            sz = str(r.get("size") or "")
            cnt = int(r.get("count") or 0)
            label = f"{sz} ({cnt})"
            style = "SizeZero.TButton" if cnt == 0 else ("SizeLow.TButton" if cnt <= 5 else "Size.TButton")

            btn = ttk.Button(row, text=label, style=style)
            # Double-click adds to active bill directly
            btn.configure(command=lambda v=label, b=btn: self._on_size_click(v, b))
            btn.bind("<Double-Button-1>", lambda e, v=label, b=btn: self._on_size_double_click(v, b))
            btn.pack(side=tk.LEFT, padx=3, pady=3)

            self._size_btns[label] = (btn, cnt == 0)

        ttk.Button(self._grid_host, text="\u25c4 رجوع إلى الألوان", command=self._render_colors)            .pack(anchor="w", padx=4, pady=4)

        if preserve_size:
            psz = str(preserve_size).strip()
            for lbl in list(self._size_btns.keys()):
                if lbl.startswith(psz + " ("):
                    btn, _ = self._size_btns[lbl]
                    self._on_size_click(lbl, btn)
                    break

    def _on_size_click(self, label_with_count: str, btn):
        size = label_with_count.split(" (", 1)[0].strip()
        self._sel_size = size
        self._price_user_edited = False

        if self._selected_size_btn is not None:
            old_btn, was_zero = self._selected_size_btn
            if old_btn.winfo_exists():
                old_btn.configure(style="SizeZero.TButton" if was_zero else "Size.TButton")

        btn.configure(style="SizeSelected.TButton")
        was_zero = "(0)" in label_with_count
        self._selected_size_btn = (btn, was_zero)

    def _on_size_double_click(self, label_with_count: str, btn):
        size = label_with_count.split(" (", 1)[0].strip()
        self._sel_size = size
        self._on_size_click(label_with_count, btn)
        self._add_to_active_bill(size, qty=1)

    def _add_to_active_bill(self, size: str, qty: int = 1):
        if not all([self._sel_school, self._sel_item, self._sel_color]):
            messagebox.showwarning("اختر أولاً", "اختر المدرسة والنوع واللون أولاً.")
            return

        # Find price
        price = self._compute_price_for_size(size)
        if price is None:
            price = 0.0

        # Determine direction tag for exchange bill
        direction = None
        if self.active_bill == 6:
            direction = self._exchange_mode  # "return" or "take"

        bill_lines = self.bills[self.active_bill]
        # Merge if same specs, price, and direction
        for existing in bill_lines:
            if (existing.get("item_type") == self._sel_item
                    and existing.get("school") == self._sel_school
                    and existing.get("color") == self._sel_color
                    and existing.get("size") == size
                    and abs(float(existing.get("unit_price", 0)) - price) < 0.001
                    and existing.get("direction") == direction):
                existing["qty"] = int(existing["qty"]) + qty
                self._sync_bill_table()
                self._refresh_size_grid_if_current(preserve_size=size)
                return

        line = {
            "item_type": self._sel_item,
            "school": self._sel_school,
            "color": self._sel_color,
            "size": size,
            "unit_price": price,
            "qty": qty,
        }
        if direction is not None:
            line["direction"] = direction
        bill_lines.append(line)
        self._sync_bill_table()
        self._refresh_size_grid_if_current(preserve_size=size)

    def _compute_price_for_size(self, target_size: str) -> Optional[float]:
        import math
        try:
            def _ceil10(x):
                return float(int(math.ceil(x / 10.0) * 10))

            known: Dict[str, float] = {}
            for r in (self._sizes_cache or []):
                s = str(r.get("size") or "")
                lp = r.get("last_price")
                if lp is not None:
                    try:
                        known[s] = float(lp)
                    except Exception:
                        pass

            if target_size in known:
                return float(known[target_size])

            try:
                index_of = {lbl: i for i, lbl in enumerate([r.get("size") for r in (self._sizes_cache or [])])}
            except Exception:
                index_of = {}

            points = [(index_of[s], float(p)) for s, p in known.items() if s in index_of]

            if len(points) >= 2:
                xs = [float(x) for x, _ in points]
                ys = [float(y) for _, y in points]
                mean_x = sum(xs) / len(xs)
                mean_y = sum(ys) / len(ys)
                cov = sum((x - mean_x) * (y - mean_y) for x, y in zip(xs, ys))
                varx = sum((x - mean_x) ** 2 for x in xs)
                slope = 0.0 if varx == 0 else cov / varx
                intercept = mean_y - slope * mean_x
                if target_size in index_of:
                    tx = float(index_of[target_size])
                    pred = slope * tx + intercept
                    if pred is not None and pred == pred and pred >= 0:
                        return _ceil10(pred)

            if points and target_size in index_of:
                tgt_idx = index_of[target_size]
                nearest = min(points, key=lambda ip: abs(ip[0] - tgt_idx))
                return _ceil10(nearest[1])

            if points:
                prices = sorted([p for _, p in points])
                return _ceil10(prices[len(prices) // 2])

            try:
                p = self.db.last_price_for_specs(
                    self._sel_item or "", self._sel_school or "", self._sel_color or "", "")
                if p is not None:
                    return float(round(float(p), 2))
            except Exception:
                pass
        except Exception:
            pass
        return None

    def _get_sizes_for_bill(self, school: str, item: str, color: str) -> List[str]:
        profile = self.db.get_size_profile(item, school, color)
        sizes: List[str] = []

        if profile:
            r1s, r1e, r2s, r2e, has_alpha = profile
            sizes.extend(
                merged_numeric_size_labels_from_profile(r1s, r1e, r2s, r2e),
            )
            if has_alpha:
                sizes.extend(ALPHA_SIZES)
            return sizes

        try:
            rows = self.db.current_inventory({"school": school, "item_type": item, "color": color})
            seen: set = set()
            for r in rows:
                sz = str(r.get("size") or "").strip()
                if sz and sz not in seen:
                    seen.add(sz)
                    sizes.append(sz)
            return sizes
        except Exception:
            return []

    # ------------------------------------------------------------------ quick filter
    def _build_filter_constraints(self, exclude_field: str) -> Dict[str, Any]:
        c: Dict[str, Any] = {}
        if exclude_field != "school":
            v = self._flt_school.get()
            if v:
                c["school"] = v
        if exclude_field != "item_type":
            v = self._flt_item.get()
            if v:
                c["item_type"] = v
        if exclude_field != "color":
            v = self._flt_color.get()
            if v:
                c["color"] = v
        return c

    def _on_filter_changed(self):
        """Refresh all filter dropdowns (cascading) then apply."""
        self._flt_school.refresh_values()
        self._flt_item.refresh_values()
        self._flt_color.refresh_values()
        self._auto_apply_filter()

    def _auto_apply_filter(self):
        """Called automatically when any filter combobox value changes."""
        sc = self._flt_school.get()
        it = self._flt_item.get()
        cl = self._flt_color.get()

        if sc and it and cl:
            self._sel_school = sc
            self._sel_item = it
            self._sel_color = cl
            self._render_sizes()
        elif sc and it:
            self._sel_school = sc
            self._sel_item = it
            self._render_colors()
        elif sc:
            self._sel_school = sc
            self._render_items()
        else:
            self._render_schools()

    def _clear_quick_filter(self):
        self._flt_school.set("")
        self._flt_item.set("")
        self._flt_color.set("")
        self._flt_school.refresh_values()
        self._flt_item.refresh_values()
        self._flt_color.refresh_values()
        self._render_schools()

    # ------------------------------------------------------------------ search & favorites
    def _on_search_changed(self):
        """Filter the button grid based on search text."""
        query = self._search_var.get().strip()
        if not query:
            if self._sel_color:
                self._render_sizes()
            elif self._sel_item:
                self._render_colors()
            elif self._sel_school:
                self._render_items()
            else:
                self._render_schools()
            return

        self._clear_grid()
        self._crumb_var.set(f"نتائج البحث: {query}")
        try:
            results = self.db.current_inventory({})
            seen = set()
            matches = []
            for r in results:
                key = (r.get("school", ""), r.get("item_type", ""), r.get("color", ""))
                if key in seen:
                    continue
                text = f"{r.get('school', '')} {r.get('item_type', '')} {r.get('color', '')}"
                if query in text:
                    seen.add(key)
                    matches.append(r)
            cols = 3
            row_frame = None
            for i, r in enumerate(matches[:24]):
                if i % cols == 0:
                    row_frame = ttk.Frame(self._grid_host)
                    row_frame.pack(fill=tk.X, padx=4, pady=2)
                label = f"{r['item_type']} - {r['school']} ({r['color']})"
                btn = ttk.Button(row_frame, text=label,
                                 command=lambda s=r.get("school"), it=r.get("item_type"), c=r.get("color"): self._jump_to(s, it, c))
                btn.pack(side=tk.LEFT, padx=3, pady=3)
        except Exception:
            pass

    def _jump_to(self, school, item_type, color):
        """Navigate directly to sizes view for a specific school/item/color combo."""
        self._sel_school = school
        self._sel_item = item_type
        self._sel_color = color
        self._price_user_edited = False
        self._search_var.set("")
        self._render_sizes()

    def _clear_search(self):
        """Clear the search bar and reset to schools view."""
        self._search_var.set("")
        self._sel_school = None
        self._sel_item = None
        self._sel_color = None
        self._sel_size = None
        self._render_schools()

    def sync_refresh(self):
        """Refresh POS selector state after inbound sync changes stock."""
        try:
            self._flt_school.refresh_values()
            self._flt_item.refresh_values()
            self._flt_color.refresh_values()
        except Exception:
            pass

        try:
            self._refresh_favorites()
        except Exception:
            pass

        query = (self._search_var.get() or "").strip()
        if query:
            self._on_search_changed()
            return

        if self._sel_color:
            self._render_sizes()
        elif self._sel_item:
            self._render_colors()
        elif self._sel_school:
            self._render_items()
        else:
            self._render_schools()

    def _refresh_favorites(self):
        """Show top 5 most-sold item combos as quick-access buttons."""
        for w in self._fav_inner.winfo_children():
            w.destroy()
        try:
            rows = self.db.conn.execute("""
                SELECT item_type, school, color, SUM(qty) as total_qty
                FROM bill_items
                GROUP BY item_type, school, color
                ORDER BY total_qty DESC
                LIMIT 5
            """).fetchall()
            if not rows:
                ttk.Label(self._fav_inner, text="لا توجد مبيعات بعد",
                          font=("Segoe UI", 8)).pack(side=tk.LEFT)
                return
            for r in rows:
                label = f"{r['item_type']}-{r['school']} ({r['color']})"
                btn = ttk.Button(self._fav_inner, text=label, style="Secondary.TButton",
                                 command=lambda s=r["school"], it=r["item_type"], c=r["color"]: self._jump_to(s, it, c))
                btn.pack(side=tk.LEFT, padx=2, pady=2)
        except Exception:
            pass

    # ------------------------------------------------------------------ bill management
    def _switch_bill(self, n: int):
        # Save current customer
        self.customers[self.active_bill] = self._customer_var.get()
        self.active_bill = n
        # Restore customer for this bill
        self._customer_var.set(self.customers[n])
        self._update_bill_btn_styles()
        self._sync_bill_table()
        self._update_res_frame_visibility()
        self._refresh_size_grid_if_current(preserve_size=self._sel_size)

    def _update_bill_btn_styles(self):
        for n, btn in self._bill_btns.items():
            if n == self.active_bill:
                try:
                    btn.configure(style="ActiveBill.TButton")
                except Exception:
                    pass
            else:
                try:
                    btn.configure(style="TButton")
                except Exception:
                    pass

    def _save_customer(self):
        self.customers[self.active_bill] = self._customer_var.get()

    def _choose_branch_target(self):
        current_device = (self.db.conn.execute(
            "SELECT device_name FROM device_identity WHERE id = 1"
        ).fetchone() or [None])[0]
        choices = [n for n in DEFAULT_BRANCH_POS_NAMES if n != current_device]
        ui_choices = [_branch_display_name(n) for n in choices]
        ui_to_dev = {_branch_display_name(n): n for n in choices}
        dlg = tk.Toplevel(self.winfo_toplevel())
        dlg.title("طلب تحويل إلى فرع")
        dlg.geometry("320x140")
        dlg.transient(self.winfo_toplevel())
        dlg.grab_set()
        frm = ttk.Frame(dlg, padding=12)
        frm.pack(fill=tk.BOTH, expand=True)
        ttk.Label(frm, text="اختر الفرع المطلوب لإرسال الطلب إلى المخزن:").pack(anchor="w", pady=(0, 8))
        var = tk.StringVar(value=ui_choices[0] if ui_choices else "")
        cb = ttk.Combobox(frm, textvariable=var, values=ui_choices, state="readonly")
        cb.pack(fill=tk.X)

        def _apply():
            picked = (var.get() or "").strip()
            dev = ui_to_dev.get(picked, picked)
            if dev:
                self._customer_var.set(f"{BRANCH_TARGET_PREFIX}{_branch_display_name(dev)}")
            dlg.destroy()

        btns = ttk.Frame(frm)
        btns.pack(fill=tk.X, pady=(12, 0))
        ttk.Button(btns, text="اختيار", command=_apply).pack(side=tk.LEFT, padx=(0, 8))
        ttk.Button(btns, text="إلغاء", command=dlg.destroy).pack(side=tk.LEFT)

    def _autocomplete_customer(self):
        """Show auto-complete suggestions for customer name."""
        query = self._customer_var.get().strip()
        # Destroy old listbox if exists
        if self._cust_listbox:
            try:
                self._cust_listbox.destroy()
            except Exception:
                pass
            self._cust_listbox = None

        if len(query) < 2:
            return
        try:
            rows = self.db.conn.execute(
                "SELECT DISTINCT customer FROM bills WHERE customer LIKE ? AND customer != '' LIMIT 8",
                (f"%{query}%",)
            ).fetchall()
            names = [r["customer"] for r in rows if r["customer"]]
            special = [WAREHOUSE_RETURN_LABEL] + [
                f"{BRANCH_TARGET_PREFIX}{_branch_display_name(n)}" for n in DEFAULT_BRANCH_POS_NAMES
            ]
            for label in reversed(special):
                if query.lower() in label.lower() and label not in names:
                    names.insert(0, label)
            if not names:
                return

            lb = tk.Listbox(self.winfo_toplevel(), height=min(len(names), 6),
                            font=("Segoe UI", 9), bg="#ffffff", fg="#1e293b",
                            selectbackground="#bfdbfe", bd=1, relief="solid")
            for n in names:
                lb.insert(tk.END, n)

            # Position below the entry widget
            x = self._cust_entry.winfo_rootx()
            y = self._cust_entry.winfo_rooty() + self._cust_entry.winfo_height()
            lb.place(x=x - self.winfo_toplevel().winfo_rootx(),
                     y=y - self.winfo_toplevel().winfo_rooty(),
                     width=self._cust_entry.winfo_width())
            lb.lift()

            def _select(event=None):
                sel = lb.curselection()
                if sel:
                    self._customer_var.set(lb.get(sel[0]))
                lb.destroy()
                self._cust_listbox = None

            lb.bind("<ButtonRelease-1>", _select)
            lb.bind("<Return>", _select)
            self._cust_listbox = lb

            # Auto-dismiss when clicking elsewhere
            def _dismiss(event=None):
                if self._cust_listbox:
                    try:
                        self._cust_listbox.destroy()
                    except Exception:
                        pass
                    self._cust_listbox = None
            self._cust_entry.bind("<Escape>", _dismiss, add="+")
        except Exception:
            pass

    def _update_res_frame_visibility(self):
        self._res_frame.pack_forget()
        self._return_frame.pack_forget()
        self._exchange_frame.pack_forget()
        if self.active_bill == 5:
            self._res_frame.pack(fill=tk.X, padx=4, pady=(2, 2))
        elif self.active_bill == 4:
            self._return_frame.pack(fill=tk.X, padx=4, pady=(2, 2))
        elif self.active_bill == 6:
            self._exchange_frame.pack(fill=tk.X, padx=4, pady=(2, 2))

    def _on_exchange_mode_changed(self):
        self._exchange_mode = self._exchange_mode_var.get()

    def _sync_bill_table(self):
        self.bill_table.delete(*self.bill_table.get_children())
        total = 0.0
        return_total = 0.0
        take_total = 0.0
        is_return = self.active_bill == 4
        is_exchange = self.active_bill == 6

        for idx, ln in enumerate(self.bills[self.active_bill]):
            line_total = float(ln["unit_price"]) * int(ln["qty"])
            display_type = ln["item_type"]

            if is_return:
                display_type = f"\u21a9 {display_type}"
                total += line_total
            elif is_exchange:
                direction = ln.get("direction", "return")
                if direction == "return":
                    display_type = f"\u21a9 {display_type}"
                    return_total += line_total
                else:
                    display_type = f"\u21aa {display_type}"
                    take_total += line_total
            else:
                total += line_total

            self.bill_table.insert(
                "", tk.END, iid=str(idx),
                values=(display_type, ln["school"], ln["color"], ln["size"],
                        f"{float(ln['unit_price']):.2f}", ln["qty"], f"{line_total:.2f}")
            )

        if is_exchange:
            diff = take_total - return_total
            if diff > 0:
                diff_text = f"فرق السعر: {diff:.2f} (يدفع العميل)"
            elif diff < 0:
                diff_text = f"فرق السعر: {abs(diff):.2f} (يسترد العميل)"
            else:
                diff_text = "فرق السعر: 0.00 (متساوي)"
            self._exchange_diff_var.set(diff_text)
            self.total_var.set(f"{diff:.2f}")
        else:
            self.total_var.set(f"{total:.2f}")

        _apply_zebra_tags(self.bill_table)

    def _inc_qty(self):
        sel = self.bill_table.selection()
        if not sel:
            return
        idx = int(sel[0])
        lines = self.bills[self.active_bill]
        if 0 <= idx < len(lines):
            lines[idx]["qty"] = int(lines[idx]["qty"]) + 1
            self._sync_bill_table()
            self.bill_table.selection_set(str(idx))
            self._refresh_size_grid_if_current(preserve_size=self._sel_size)

    def _dec_qty(self):
        sel = self.bill_table.selection()
        if not sel:
            return
        idx = int(sel[0])
        lines = self.bills[self.active_bill]
        if 0 <= idx < len(lines):
            if int(lines[idx]["qty"]) > 1:
                lines[idx]["qty"] = int(lines[idx]["qty"]) - 1
            else:
                lines.pop(idx)
            self._sync_bill_table()
            self._refresh_size_grid_if_current(preserve_size=self._sel_size)

    def _remove_line(self):
        sel = self.bill_table.selection()
        if not sel:
            return
        idx = int(sel[0])
        lines = self.bills[self.active_bill]
        if 0 <= idx < len(lines):
            lines.pop(idx)
            self._sync_bill_table()
            self._refresh_size_grid_if_current(preserve_size=self._sel_size)

    def _clear_bill(self):
        self.bills[self.active_bill].clear()
        self._sync_bill_table()
        self._refresh_size_grid_if_current(preserve_size=self._sel_size)

    # ------------------------------------------------------------------ finalize
    def _finalize(self):
        lines = self.bills[self.active_bill]
        if not lines:
            messagebox.showwarning("فارغ", "لا توجد أصناف في الفاتورة.")
            return

        self._save_customer()
        customer = (self.customers[self.active_bill] or "").strip()
        total_qty = sum(int(ln["qty"]) for ln in lines)
        is_reservation = self.active_bill == 5
        is_return = self.active_bill == 4
        is_exchange = self.active_bill == 6
        warehouse_target = _extract_warehouse_target(customer)

        # Determine title and total display
        if is_return:
            overlay_title = "تأكيد المرتجع"
            review_title = "مراجعة المرتجع"
            total = sum(float(ln["unit_price"]) * int(ln["qty"]) for ln in lines)
            total_label = f"مبلغ الاسترداد: {total:.2f}"
        elif is_exchange:
            overlay_title = "تأكيد الاستبدال"
            review_title = "مراجعة الاستبدال"
            ret_lines = [ln for ln in lines if ln.get("direction") == "return"]
            take_lines = [ln for ln in lines if ln.get("direction") == "take"]
            if not ret_lines and not take_lines:
                messagebox.showwarning("فارغ", "أضف أصناف مرتجعة ومأخوذة.")
                return
            ret_total = sum(float(ln["unit_price"]) * int(ln["qty"]) for ln in ret_lines)
            take_total = sum(float(ln["unit_price"]) * int(ln["qty"]) for ln in take_lines)
            total = take_total - ret_total
            if total > 0:
                total_label = f"يدفع العميل: {total:.2f}"
            elif total < 0:
                total_label = f"يسترد العميل: {abs(total):.2f}"
            else:
                total_label = "متساوي - لا فرق في السعر"
        elif is_reservation:
            overlay_title = "تأكيد الحجز"
            review_title = "مراجعة الحجز"
            total = sum(float(ln["unit_price"]) * int(ln["qty"]) for ln in lines)
            total_label = f"الإجمالي: {total:.2f}"
        elif warehouse_target:
            overlay_title = f"تأكيد {WAREHOUSE_RETURN_LABEL}"
            review_title = f"مراجعة {WAREHOUSE_RETURN_LABEL}"
            total = sum(float(ln["unit_price"]) * int(ln["qty"]) for ln in lines)
            total_label = f"قيمة المرجع: {total:.2f}"
        elif branch_target := _extract_branch_target(customer):
            overlay_title = "تأكيد تحويل إلى فرع"
            review_title = f"مراجعة تحويل إلى {branch_target}"
            total = sum(float(ln["unit_price"]) * int(ln["qty"]) for ln in lines)
            total_label = f"قيمة التحويل: {total:.2f}"
        else:
            overlay_title = "تأكيد الفاتورة"
            review_title = "مراجعة الفاتورة"
            total = sum(float(ln["unit_price"]) * int(ln["qty"]) for ln in lines)
            total_label = f"الإجمالي: {total:.2f}"

        # Show confirmation overlay
        overlay = tk.Toplevel(self)
        overlay.title(overlay_title)
        overlay.geometry("450x420" if is_exchange else "420x360")
        overlay.transient(self.winfo_toplevel())
        overlay.grab_set()

        frm = ttk.Frame(overlay, padding=16)
        frm.pack(fill=tk.BOTH, expand=True)

        ttk.Label(frm, text=review_title, font=("Segoe UI", 13, "bold")).pack(anchor="w", pady=(0, 8))

        if customer:
            ttk.Label(frm, text=f"العميل: {customer}").pack(anchor="w")

        ttk.Label(frm, text=f"عدد الأصناف: {len(lines)}   |   عدد القطع: {total_qty}",
                  font=("Segoe UI", 9)).pack(anchor="w", pady=(4, 4))

        # Items summary table
        cols_frm = ttk.Frame(frm)
        cols_frm.pack(fill=tk.BOTH, expand=True, pady=(4, 8))

        if is_exchange:
            summary_tree = ttk.Treeview(cols_frm, columns=("dir", "spec", "qty", "total"),
                                         show="headings", height=min(len(lines), 6))
            summary_tree.heading("dir", text="نوع")
            summary_tree.heading("spec", text="الصنف")
            summary_tree.heading("qty", text="الكمية")
            summary_tree.heading("total", text="المبلغ")
            summary_tree.column("dir", width=50, anchor="center")
            summary_tree.column("spec", width=200)
            summary_tree.column("qty", width=60, anchor="center")
            summary_tree.column("total", width=80, anchor="center")
            for ln in lines:
                spec = f"{ln['item_type']} - {ln['school']} ({ln['color']}/{ln['size']})"
                lt = float(ln['unit_price']) * int(ln['qty'])
                d = "\u21a9" if ln.get("direction") == "return" else "\u21aa"
                summary_tree.insert("", tk.END, values=(d, spec, ln['qty'], f"{lt:.2f}"))
        else:
            summary_tree = ttk.Treeview(cols_frm, columns=("spec", "qty", "total"),
                                         show="headings", height=min(len(lines), 6))
            summary_tree.heading("spec", text="الصنف")
            summary_tree.heading("qty", text="الكمية")
            summary_tree.heading("total", text="المبلغ")
            summary_tree.column("spec", width=220)
            summary_tree.column("qty", width=60, anchor="center")
            summary_tree.column("total", width=80, anchor="center")
            for ln in lines:
                prefix = "\u21a9 " if is_return else ""
                spec = f"{prefix}{ln['item_type']} - {ln['school']} ({ln['color']}/{ln['size']})"
                lt = float(ln['unit_price']) * int(ln['qty'])
                summary_tree.insert("", tk.END, values=(spec, ln['qty'], f"{lt:.2f}"))
        summary_tree.pack(fill=tk.BOTH, expand=True)

        ttk.Label(frm, text=total_label,
                  font=("Segoe UI", 14, "bold")).pack(anchor="center", pady=(8, 12))

        btn_row = ttk.Frame(frm)
        btn_row.pack(fill=tk.X)

        def _do_finalize():
            overlay.destroy()
            self._execute_finalize(lines, customer, is_reservation, total)

        ttk.Button(btn_row, text="تأكيد", command=_do_finalize,
                   style="Success.TButton").pack(side=tk.LEFT, padx=(0, 8))
        ttk.Button(btn_row, text="إلغاء", command=overlay.destroy,
                   style="Secondary.TButton").pack(side=tk.LEFT)

    def _execute_finalize(self, lines, customer, is_reservation, total):
        """Execute the actual finalize after confirmation."""
        is_return = self.active_bill == 4
        is_exchange = self.active_bill == 6

        if is_return:
            try:
                bill_id = self.db.create_return_bill(customer, lines)
            except Exception as ex:
                messagebox.showerror("فشل المرتجع", str(ex), parent=self)
                return
            self.bills[4].clear()
            self._sync_bill_table()
            self._refresh_size_grid_if_current(preserve_size=self._sel_size)
            ToastNotification.show(self.winfo_toplevel(),
                                   f"تم إنشاء فاتورة مرتجع #{bill_id} (استرداد: {total:.2f})", toast_type="success")
            if messagebox.askyesno("طباعة", f"طباعة فاتورة المرتجع #{bill_id}؟"):
                self._print_return_bill(bill_id, total)

        elif is_exchange:
            # Copy lines before clearing since lines is a reference to self.bills[6]
            lines_copy = [dict(ln) for ln in lines]
            ret_lines = [ln for ln in lines_copy if ln.get("direction") == "return"]
            take_lines = [ln for ln in lines_copy if ln.get("direction") == "take"]
            try:
                bill_id = self.db.create_exchange_bill(customer, ret_lines, take_lines)
            except Exception as ex:
                messagebox.showerror("فشل الاستبدال", str(ex), parent=self)
                return
            self.bills[6].clear()
            self._exchange_mode = "return"
            self._exchange_mode_var.set("return")
            self._exchange_diff_var.set("فرق السعر: 0.00")
            self._sync_bill_table()
            self._refresh_size_grid_if_current(preserve_size=self._sel_size)
            ToastNotification.show(self.winfo_toplevel(),
                                   f"تم إنشاء فاتورة استبدال #{bill_id}", toast_type="success")
            if messagebox.askyesno("طباعة", f"طباعة فاتورة الاستبدال #{bill_id}؟"):
                self._print_exchange_bill(bill_id, lines_copy, total)

        elif is_reservation:
            try:
                paid = float((self._paid_var.get() or "0").strip())
            except Exception:
                paid = 0.0
            note = self._res_note_var.get().strip()
            for ln in lines:
                ln["note"] = note
            try:
                ids = self.db.create_reservation(customer, lines, paid_amount=paid)
                self._print_reservation_receipt(ids, lines, customer, paid, total)
                ToastNotification.show(self.winfo_toplevel(),
                                       f"تم إنشاء {len(ids)} حجز/حجوزات بنجاح", toast_type="success")
                self.bills[5].clear()
                self._paid_var.set("0")
                self._res_note_var.set("")
                self._sync_bill_table()
                self._refresh_size_grid_if_current(preserve_size=self._sel_size)
            except Exception as ex:
                messagebox.showerror("فشل الحجز", str(ex), parent=self)
        else:
            try:
                bill_id = self.db.create_bill(customer, lines)
            except Exception as ex:
                messagebox.showerror("فشل الحفظ", str(ex), parent=self)
                return

            self.bills[self.active_bill].clear()
            self._sync_bill_table()
            self._refresh_size_grid_if_current(preserve_size=self._sel_size)
            self._refresh_favorites()
            warehouse_target = _extract_warehouse_target(customer)
            branch_target = _extract_branch_target(customer)
            if warehouse_target:
                ToastNotification.show(
                    self.winfo_toplevel(),
                    f"تم إنشاء {WAREHOUSE_RETURN_LABEL} #{bill_id} وسيتم إرساله للمراجعة في المخزن",
                    toast_type="success",
                )
                print_title = "طباعة مستند الإرجاع"
                print_body = f"طباعة مستند {WAREHOUSE_RETURN_LABEL} #{bill_id}؟"
            elif branch_target:
                ToastNotification.show(
                    self.winfo_toplevel(),
                    f"تم إنشاء طلب تحويل #{bill_id} إلى {branch_target} عبر المخزن",
                    toast_type="success",
                )
                print_title = "طباعة مستند التحويل"
                print_body = f"طباعة مستند التحويل إلى {branch_target} #{bill_id}؟"
            else:
                ToastNotification.show(
                    self.winfo_toplevel(),
                    f"تم إنشاء الفاتورة #{bill_id} (الإجمالي: {total:.2f})",
                    toast_type="success",
                )
                print_title = "طباعة الفاتورة"
                print_body = f"طباعة الفاتورة #{bill_id}؟"

            if messagebox.askyesno(print_title, print_body):
                self._direct_print_bill(bill_id)

    def _direct_print_bill(self, bill_id: int, copies: int = 2):
        try:
            bill = next(b for b in self.db.list_bills() if int(b["id"]) == int(bill_id))
            items = self.db.list_bill_items(bill_id)
        except StopIteration:
            messagebox.showerror("فشل الطباعة", "لم يتم العثور على الفاتورة.")
            return
        tmp_dir = tempfile.gettempdir()
        path = os.path.join(tmp_dir, f"bill_{bill_id}.html")
        save_bill_as_html(path, bill, items)
        _print_html_auto(path, copies=max(1, int(copies)), parent=self)

    def _print_reservation_receipt(self, ids, lines, customer, paid, total):
        """Auto-print a receipt for a newly created reservation."""
        res_nums = ", ".join(str(i) for i in ids)
        rows_html = ""
        for ln in lines:
            line_total = float(ln["unit_price"]) * int(ln["qty"])
            rows_html += (
                f'<tr><td>{_html(ln["item_type"])} - {_html(ln["school"])}'
                f'<br><small>{_html(ln["color"])} / {_html(ln["size"])}</small></td>'
                f'<td style="text-align:center">{ln["qty"]}</td>'
                f'<td style="text-align:left">{float(ln["unit_price"]):.2f}</td>'
                f'<td style="text-align:left">{line_total:.2f}</td></tr>\n'
            )
        customer_line = f"<div>العميل: {_html(customer)}</div>" if customer else ""
        remaining = total - paid
        html = f"""<!DOCTYPE html>
<html lang="ar" dir="rtl">
<head><meta charset="utf-8"><title>حجز #{res_nums}</title>
<style>
  @page {{ size: 80mm auto; margin: 2mm; }}
  * {{ box-sizing: border-box; }}
  body {{ font-family: "Segoe UI", Tahoma, Arial, sans-serif; font-size: 11px; margin: 0; padding: 0; width: 76mm; direction: rtl; }}
  .receipt {{ padding: 2mm; }}
  .center {{ text-align: center; }}
  .sep {{ border: none; border-top: 1px dashed #000; margin: 4px 0; }}
  h2 {{ font-size: 14px; margin: 4px 0; text-align: center; }}
  .info {{ font-size: 11px; margin: 2px 0; }}
  table {{ border-collapse: collapse; width: 100%; font-size: 10px; }}
  th {{ border-bottom: 1px solid #000; padding: 3px 2px; text-align: right; }}
  td {{ padding: 3px 2px; vertical-align: top; border-bottom: 1px dotted #ccc; }}
  .total-row {{ font-size: 13px; font-weight: bold; text-align: center; margin: 4px 0; }}
  .footer {{ text-align: center; font-size: 10px; margin-top: 6px; }}
</style></head>
<body><div class="receipt">
  <h2>حجز #{res_nums}</h2>
  <hr class="sep">
  <div class="info">التاريخ: {now_iso()[:16].replace("T", " ")}</div>
  {customer_line}
  <hr class="sep">
  <table><thead><tr><th>الصنف</th><th>الكمية</th><th>السعر</th><th>المجموع</th></tr></thead>
  <tbody>{rows_html}</tbody></table>
  <hr class="sep">
  <div class="total-row">الإجمالي: {total:.2f}</div>
  <div class="info" style="text-align:center">المدفوع: {paid:.2f}</div>
  <div class="info" style="text-align:center;font-weight:bold">المتبقي: {remaining:.2f}</div>
  <hr class="sep">
  <div class="footer">شكراً لتعاملكم معنا</div>
</div></body></html>"""
        path = os.path.join(tempfile.gettempdir(), f"reservation_{res_nums}.html")
        with open(path, "w", encoding="utf-8") as f:
            f.write(html)
        _print_html_auto(path, copies=2, parent=self)

    def _print_return_bill(self, bill_id: int, total: float):
        items = self.db.list_bill_items(bill_id)
        rows_html = ""
        for it in items:
            rows_html += (
                f'<tr><td>{_html(it["item_type"])} - {_html(it["school"])}'
                f'<br><small>{_html(it["color"])} / {_html(it["size"])}</small></td>'
                f'<td style="text-align:center">{it["qty"]}</td>'
                f'<td style="text-align:left">{float(it["unit_price"]):.2f}</td>'
                f'<td style="text-align:left">{float(it["line_total"]):.2f}</td></tr>\n'
            )
        html = f"""<!DOCTYPE html>
<html lang="ar" dir="rtl"><head><meta charset="utf-8"><title>مرتجع #{bill_id}</title>
<style>
@page {{ size: 80mm auto; margin: 2mm; }}
body {{ font-family: "Segoe UI", Tahoma, sans-serif; font-size: 11px; width: 76mm; direction: rtl; margin:0; padding:2mm; }}
h2 {{ font-size: 14px; text-align: center; margin: 4px 0; }}
.sep {{ border:none; border-top:1px dashed #000; margin:4px 0; }}
table {{ border-collapse: collapse; width: 100%; font-size: 10px; }}
th {{ border-bottom: 1px solid #000; padding: 3px 2px; text-align: right; }}
td {{ padding: 3px 2px; border-bottom: 1px dotted #ccc; }}
.total {{ font-size: 13px; font-weight: bold; text-align: center; margin: 6px 0; }}
</style></head><body>
<h2>فاتورة مرتجع #{bill_id}</h2>
<hr class="sep">
<div>التاريخ: {now_iso()[:16].replace("T"," ")}</div>
<hr class="sep">
<table><thead><tr><th>الصنف</th><th>الكمية</th><th>السعر</th><th>المجموع</th></tr></thead>
<tbody>{rows_html}</tbody></table>
<hr class="sep">
<div class="total">مبلغ الاسترداد: {total:.2f}</div>
</body></html>"""
        path = os.path.join(tempfile.gettempdir(), f"return_{bill_id}.html")
        with open(path, "w", encoding="utf-8") as f:
            f.write(html)
        _print_html_auto(path, copies=2, parent=self)

    def _print_exchange_bill(self, bill_id: int, lines: list, diff: float):
        ret_rows = ""
        take_rows = ""
        for ln in lines:
            lt = float(ln["unit_price"]) * int(ln["qty"])
            row = (
                f'<tr><td>{_html(ln["item_type"])} - {_html(ln["school"])}'
                f'<br><small>{_html(ln["color"])} / {_html(ln["size"])}</small></td>'
                f'<td style="text-align:center">{ln["qty"]}</td>'
                f'<td style="text-align:left">{float(ln["unit_price"]):.2f}</td>'
                f'<td style="text-align:left">{lt:.2f}</td></tr>\n'
            )
            if ln.get("direction") == "return":
                ret_rows += row
            else:
                take_rows += row
        if diff > 0:
            diff_text = f"يدفع العميل: {diff:.2f}"
        elif diff < 0:
            diff_text = f"يسترد العميل: {abs(diff):.2f}"
        else:
            diff_text = "لا فرق في السعر"
        html = f"""<!DOCTYPE html>
<html lang="ar" dir="rtl"><head><meta charset="utf-8"><title>استبدال #{bill_id}</title>
<style>
@page {{ size: 80mm auto; margin: 2mm; }}
body {{ font-family: "Segoe UI", Tahoma, sans-serif; font-size: 11px; width: 76mm; direction: rtl; margin:0; padding:2mm; }}
h2 {{ font-size: 14px; text-align: center; margin: 4px 0; }}
h3 {{ font-size: 12px; margin: 4px 0; }}
.sep {{ border:none; border-top:1px dashed #000; margin:4px 0; }}
table {{ border-collapse: collapse; width: 100%; font-size: 10px; }}
th {{ border-bottom: 1px solid #000; padding: 3px 2px; text-align: right; }}
td {{ padding: 3px 2px; border-bottom: 1px dotted #ccc; }}
.total {{ font-size: 13px; font-weight: bold; text-align: center; margin: 6px 0; }}
</style></head><body>
<h2>فاتورة استبدال #{bill_id}</h2>
<hr class="sep">
<div>التاريخ: {now_iso()[:16].replace("T"," ")}</div>
<hr class="sep">
<h3>الأصناف المرتجعة</h3>
<table><thead><tr><th>الصنف</th><th>الكمية</th><th>السعر</th><th>المجموع</th></tr></thead>
<tbody>{ret_rows if ret_rows else "<tr><td colspan='4' style='text-align:center'>-</td></tr>"}</tbody></table>
<hr class="sep">
<h3>الأصناف المأخوذة</h3>
<table><thead><tr><th>الصنف</th><th>الكمية</th><th>السعر</th><th>المجموع</th></tr></thead>
<tbody>{take_rows if take_rows else "<tr><td colspan='4' style='text-align:center'>-</td></tr>"}</tbody></table>
<hr class="sep">
<div class="total">{diff_text}</div>
</body></html>"""
        path = os.path.join(tempfile.gettempdir(), f"exchange_{bill_id}.html")
        with open(path, "w", encoding="utf-8") as f:
            f.write(html)
        _print_html_auto(path, copies=2, parent=self)


# ------------------- Factory Item Dialog -------------------

class FactoryItemDialog(tk.Toplevel):
    """
    Create an ad-hoc bill line that ships directly from the factory.
    """
    def __init__(self, master: tk.Widget, db: SqliteDatabase, preset: Optional[Dict[str, str]] = None):
        super().__init__(master)
        self.db = db
        self.title("إضافة بند من المصنع")
        self.transient(master)
        self.grab_set()
        self.geometry("520x420")
        self.resizable(False, False)

        preset = preset or {}
        frm = ttk.Frame(self, padding=10); frm.pack(fill=tk.BOTH, expand=True)

        self.t = LabeledCombobox(frm, "النوع",   db, "item_type"); self.t.set(preset.get("item_type",""))
        self.s = LabeledCombobox(frm, "المدرسة", db, "school");    self.s.set(preset.get("school",""))
        self.c = LabeledCombobox(frm, "اللون",   db, "color");     self.c.set(preset.get("color",""))
        self.z = LabeledCombobox(frm, "المقاس",  db, "size");      self.z.set(preset.get("size",""))
        for i, w in enumerate((self.t, self.s, self.c, self.z)):
            w.grid(row=i//2, column=i%2, padx=6, pady=6, sticky="ew")
        for w in (self.t.cb, self.s.cb, self.c.cb, self.z.cb):
            w.bind("<<ComboboxSelected>>", lambda e: self._auto_fill_price(), add="+")
            w.bind("<FocusOut>",           lambda e: self._auto_fill_price(), add="+")
            w.bind("<KeyRelease>",         lambda e: self._auto_fill_price(), add="+")
        frm.columnconfigure(0, weight=1); frm.columnconfigure(1, weight=1)

        self.qv = tk.StringVar(value=preset.get("qty","1"))
        self.pv = tk.StringVar(value=preset.get("unit_price",""))
        grid2 = ttk.Frame(frm); grid2.grid(row=2, column=0, columnspan=2, sticky="ew", padx=6, pady=6)
        ttk.Label(grid2, text="الكمية:").grid(row=0, column=0, sticky="e", padx=4, pady=4)
        ttk.Entry(grid2, textvariable=self.qv, width=12).grid(row=0, column=1, sticky="w")
        ttk.Label(grid2, text="سعر الوحدة:").grid(row=0, column=2, sticky="e", padx=12, pady=4)
        ttk.Entry(grid2, textvariable=self.pv, width=12).grid(row=0, column=3, sticky="w")

        btns = ttk.Frame(frm); btns.grid(row=3, column=0, columnspan=2, sticky="e", padx=6, pady=(10,0))
        ttk.Button(btns, text="إلغاء", command=self.destroy).pack(side=tk.RIGHT)
        ttk.Button(btns, text="إضافة", command=self._ok).pack(side=tk.RIGHT, padx=8)

        self.after(10, lambda: (self.t.cb.focus_set(), self._auto_fill_price()))

        self._result: Optional[Dict[str, str]] = None

    def _auto_fill_price(self):
        if (self.pv.get() or "").strip():
            return
        it = (self.t.get() or "").strip()
        sc = (self.s.get() or "").strip()
        cl = (self.c.get() or "").strip()
        sz = (self.z.get() or "").strip()
        if not (it and sc and cl and sz):
            return
        try:
            p = self.db.last_price_for_specs(it, sc, cl, sz)
            if p is not None:
                self.pv.set(f"{p:.2f}")
        except Exception:
            pass

    def _ok(self):
        try:
            item_type = self.t.get() or self._err("النوع مطلوب")
            school    = self.s.get() or self._err("المدرسة مطلوبة")
            color     = self.c.get() or self._err("اللون مطلوب")
            size      = self.z.get() or self._err("المقاس مطلوب")
            qty       = int((self.qv.get() or "0").strip());   assert qty > 0
            price     = float((self.pv.get() or "").strip());  assert price >= 0.0
        except AssertionError:
            messagebox.showerror("بيانات غير صالحة", "تحقق من الكمية (>0) والسعر (>=0).", parent=self); return
        except Exception as ex:
            messagebox.showerror("بيانات ناقصة", str(ex), parent=self); return

        self._result = {
            "item_type": item_type, "school": school, "color": color, "size": size,
            "unit_price": price, "qty": qty,
            "allow_factory_fill": True,
        }
        self.destroy()

    @staticmethod
    def _err(msg: str) -> None:
        raise RuntimeError(msg)

    def run(self) -> Optional[Dict[str, str]]:
        self.wait_window()
        return self._result

# ------------------- Statistics Frame -------------------

class StatisticsFrame(ttk.Frame):
    """Statistics and reporting tab."""

    def __init__(self, master, db: SqliteDatabase):
        super().__init__(master, padding=6)
        self.db = db
        self._build()

    def _build(self):
        ttk.Label(self, text="الإحصائيات والتقارير", font=("Segoe UI", 12, "bold")).pack(anchor="w", pady=(0, 6))

        # Filter bar with dropdown filters + date range
        filter_bar = ttk.Frame(self)
        filter_bar.pack(fill=tk.X, pady=(0, 6))

        self._sf_school = LabeledCombobox(filter_bar, "المدرسة", self.db, "school")
        self._sf_school.set_supplier(lambda: self.db.get_distinct_filtered("school", self._build_stats_constraints("school")))
        self._sf_school.grid(row=0, column=0, padx=4, pady=4, sticky="ew")

        self._sf_item = LabeledCombobox(filter_bar, "النوع", self.db, "item_type")
        self._sf_item.set_supplier(lambda: self.db.get_distinct_filtered("item_type", self._build_stats_constraints("item_type")))
        self._sf_item.grid(row=0, column=1, padx=4, pady=4, sticky="ew")

        self._sf_color = LabeledCombobox(filter_bar, "اللون", self.db, "color")
        self._sf_color.set_supplier(lambda: self.db.get_distinct_filtered("color", self._build_stats_constraints("color")))
        self._sf_color.grid(row=0, column=2, padx=4, pady=4, sticky="ew")

        df_frame = ttk.Frame(filter_bar)
        df_frame.grid(row=0, column=3, padx=4, pady=4, sticky="ew")
        ttk.Label(df_frame, text="من:").pack(side=tk.LEFT)
        self._df = DateField(df_frame, "")
        self._df.pack(side=tk.LEFT, padx=(4, 0))

        dt_frame = ttk.Frame(filter_bar)
        dt_frame.grid(row=0, column=4, padx=4, pady=4, sticky="ew")
        ttk.Label(dt_frame, text="إلى:").pack(side=tk.LEFT)
        self._dt = DateField(dt_frame, "")
        self._dt.pack(side=tk.LEFT, padx=(4, 0))

        ttk.Button(filter_bar, text="تحديث", command=self._refresh_all).grid(row=0, column=5, padx=4, pady=4)
        ttk.Button(filter_bar, text="مسح", command=self._clear_stats_filters).grid(row=0, column=6, padx=4, pady=4)

        for c in range(7):
            filter_bar.columnconfigure(c, weight=1 if c < 5 else 0)

        self._sf_school.cb.bind("<<ComboboxSelected>>", lambda e: self._on_stats_filter_changed(), add="+")
        self._sf_item.cb.bind("<<ComboboxSelected>>", lambda e: self._on_stats_filter_changed(), add="+")
        self._sf_color.cb.bind("<<ComboboxSelected>>", lambda e: self._on_stats_filter_changed(), add="+")

        # ---- Section 1: Money flow ----
        mf = ttk.LabelFrame(self, text="التدفق المالي")
        mf.pack(fill=tk.X, padx=4, pady=(0, 8))

        self._sales_count_var = tk.StringVar(value="0")
        self._sales_total_var = tk.StringVar(value="0.00")
        self._res_count_var   = tk.StringVar(value="0")
        self._res_total_var   = tk.StringVar(value="0.00")
        self._res_paid_var    = tk.StringVar(value="0.00")

        r1 = ttk.Frame(mf); r1.pack(fill=tk.X, padx=8, pady=4)
        ttk.Label(r1, text="عدد الفواتير:").pack(side=tk.LEFT)
        ttk.Label(r1, textvariable=self._sales_count_var, font=("Segoe UI", 10, "bold")).pack(side=tk.LEFT, padx=(4,20))
        ttk.Label(r1, text="إجمالي المبيعات:").pack(side=tk.LEFT)
        ttk.Label(r1, textvariable=self._sales_total_var, font=("Segoe UI", 10, "bold")).pack(side=tk.LEFT, padx=(4,0))

        r2 = ttk.Frame(mf); r2.pack(fill=tk.X, padx=8, pady=4)
        ttk.Label(r2, text="عدد الحجوزات:").pack(side=tk.LEFT)
        ttk.Label(r2, textvariable=self._res_count_var, font=("Segoe UI", 10, "bold")).pack(side=tk.LEFT, padx=(4,20))
        ttk.Label(r2, text="إجمالي الحجوزات:").pack(side=tk.LEFT)
        ttk.Label(r2, textvariable=self._res_total_var, font=("Segoe UI", 10, "bold")).pack(side=tk.LEFT, padx=(4,20))
        ttk.Label(r2, text="المدفوع:").pack(side=tk.LEFT)
        ttk.Label(r2, textvariable=self._res_paid_var, font=("Segoe UI", 10, "bold")).pack(side=tk.LEFT, padx=(4,0))

        ttk.Button(mf, text="طباعة", command=self._print_money_flow).pack(anchor="e", padx=8, pady=4)

        # ---- Section 2: Item movement ----
        mv = ttk.LabelFrame(self, text="حركة الأصناف")
        mv.pack(fill=tk.BOTH, expand=True, padx=4, pady=(0, 8))

        # Row filter checkboxes — only show items where checked column > 0
        mv_col_row = ttk.Frame(mv)
        mv_col_row.pack(fill=tk.X, padx=8, pady=(4, 0))
        ttk.Label(mv_col_row, text="إظهار فقط:").pack(side=tk.RIGHT, padx=(0, 4))
        self._mv_flt_received = tk.BooleanVar(value=False)
        self._mv_flt_sold = tk.BooleanVar(value=False)
        self._mv_flt_reserved = tk.BooleanVar(value=False)
        self._mv_flt_remaining = tk.BooleanVar(value=False)
        for var, text in [
            (self._mv_flt_received, "وارد"),
            (self._mv_flt_sold, "مباع"),
            (self._mv_flt_reserved, "محجوز"),
            (self._mv_flt_remaining, "متبقي"),
        ]:
            ttk.Checkbutton(mv_col_row, text=text, variable=var,
                            command=self._refresh_all).pack(side=tk.RIGHT, padx=4)

        mv_wrap = ttk.Frame(mv)
        mv_wrap.pack(fill=tk.BOTH, expand=True, padx=4, pady=4)

        self._mv_table = ttk.Treeview(
            mv_wrap,
            columns=("type", "school", "color", "size", "received", "sold", "reserved", "remaining"),
            show="headings",
            height=6,
        )
        for col, txt, w in [
            ("type", "النوع", 100), ("school", "المدرسة", 120), ("color", "اللون", 70),
            ("size", "المقاس", 60), ("received", "وارد", 70), ("sold", "مباع", 70),
            ("reserved", "محجوز", 70), ("remaining", "متبقي", 70),
        ]:
            self._mv_table.heading(col, text=txt)
            self._mv_table.column(col, width=w, anchor="center")
        mv_ysb = ttk.Scrollbar(mv_wrap, orient="vertical", command=self._mv_table.yview)
        self._mv_table.configure(yscrollcommand=mv_ysb.set)
        self._mv_table.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        mv_ysb.pack(side=tk.RIGHT, fill=tk.Y)
        _bind_mousewheel(self._mv_table)

        # Context menu on movement table
        _add_context_menu(self._mv_table, self)

        ttk.Button(mv, text="طباعة", command=self._print_movements).pack(anchor="e", padx=8, pady=4)

        # ---- Section 3: Income ----
        income_frame = ttk.LabelFrame(self, text="الوارد")
        income_frame.pack(fill=tk.X, padx=4, pady=4)

        self._income_count_var = tk.StringVar(value="0")
        self._income_qty_var = tk.StringVar(value="0")
        self._income_total_var = tk.StringVar(value="0.00")

        ir1 = ttk.Frame(income_frame)
        ir1.pack(fill=tk.X, padx=8, pady=4)
        ttk.Label(ir1, text="عدد فواتير الوارد:").pack(side=tk.RIGHT)
        ttk.Label(ir1, textvariable=self._income_count_var, font=("Segoe UI", 10, "bold")).pack(side=tk.RIGHT, padx=(0, 6))

        ir2 = ttk.Frame(income_frame)
        ir2.pack(fill=tk.X, padx=8, pady=4)
        ttk.Label(ir2, text="إجمالي القطع الواردة:").pack(side=tk.RIGHT)
        ttk.Label(ir2, textvariable=self._income_qty_var, font=("Segoe UI", 10, "bold")).pack(side=tk.RIGHT, padx=(0, 6))

        ir3 = ttk.Frame(income_frame)
        ir3.pack(fill=tk.X, padx=8, pady=4)
        ttk.Label(ir3, text="إجمالي قيمة الوارد:").pack(side=tk.RIGHT)
        ttk.Label(ir3, textvariable=self._income_total_var, font=("Segoe UI", 10, "bold")).pack(side=tk.RIGHT, padx=(0, 6))

        ttk.Button(income_frame, text="طباعة", command=self._print_income).pack(anchor="w", padx=8, pady=4)

        self._refresh_all()

    def _refresh_all(self):
        df = self._df.get() or None
        dt = self._dt.get() or None
        school = self._sf_school.get() or None
        item_type = self._sf_item.get() or None
        color = self._sf_color.get() or None

        # Money flow - pass all filters
        try:
            stats = self.db.get_sales_stats(df, dt, school=school, item_type=item_type, color=color)
            self._sales_count_var.set(str(stats["sales_count"]))
            self._sales_total_var.set(f"{stats['sales_total']:.2f}")
            self._res_count_var.set(str(stats["res_count"]))
            self._res_total_var.set(f"{stats['res_total']:.2f}")
            self._res_paid_var.set(f"{stats['res_paid']:.2f}")
        except Exception:
            pass

        # Item movements - pass all filters + row filtering by checkboxes
        flt_recv = self._mv_flt_received.get()
        flt_sold = self._mv_flt_sold.get()
        flt_resv = self._mv_flt_reserved.get()
        flt_rem  = self._mv_flt_remaining.get()
        any_flt  = flt_recv or flt_sold or flt_resv or flt_rem
        try:
            self._mv_table.delete(*self._mv_table.get_children())
            for r in self.db.get_item_movement_stats(date_from=df, date_to=dt, school=school, item_type=item_type, color=color):
                received = int(r.get("received") or 0)
                sold = int(r.get("sold") or 0)
                reserved = int(r.get("reserved") or 0)
                adjusted = int(r.get("adjusted") or 0)
                remaining = received - sold - reserved - adjusted
                # If any checkbox is active, only show rows matching at least one
                if any_flt:
                    if not ((flt_recv and received > 0) or
                            (flt_sold and sold > 0) or
                            (flt_resv and reserved > 0) or
                            (flt_rem and remaining > 0)):
                        continue
                self._mv_table.insert("", tk.END, values=(
                    r["item_type"], r["school"], r["color"], r["size"],
                    received, sold, reserved, remaining
                ))
            _apply_zebra_tags(self._mv_table)
        except Exception:
            pass

        # Income stats
        try:
            inc = self.db.get_income_stats(date_from=df, date_to=dt, school=school, item_type=item_type, color=color)
            self._income_count_var.set(str(inc["bill_count"]))
            self._income_qty_var.set(str(inc["total_qty"]))
            self._income_total_var.set(f"{inc['total_value']:.2f}")
        except Exception:
            pass

    def _build_stats_constraints(self, exclude_field: str) -> Dict[str, Any]:
        c: Dict[str, Any] = {}
        if exclude_field != "school":
            v = self._sf_school.get()
            if v:
                c["school"] = v
        if exclude_field != "item_type":
            v = self._sf_item.get()
            if v:
                c["item_type"] = v
        if exclude_field != "color":
            v = self._sf_color.get()
            if v:
                c["color"] = v
        return c

    def _on_stats_filter_changed(self):
        self._sf_school.refresh_values()
        self._sf_item.refresh_values()
        self._sf_color.refresh_values()
        self._refresh_all()

    def _clear_stats_filters(self):
        self._sf_school.set("")
        self._sf_item.set("")
        self._sf_color.set("")
        self._df.set("")
        self._dt.set("")
        self._sf_school.refresh_values()
        self._sf_item.refresh_values()
        self._sf_color.refresh_values()
        self._refresh_all()

    def _print_money_flow(self):
        try:
            df = self._df.get() or None
            dt = self._dt.get() or None
            stats = self.db.get_sales_stats(df, dt)
            period = ""
            if df or dt:
                period = f"من {df or '---'} إلى {dt or '---'}"
            html = f"""<!DOCTYPE html>
<html lang='ar' dir='rtl'>
<head><meta charset='utf-8'><title>التدفق المالي</title>
<style>body{{font-family:Tahoma,Arial;margin:20px}} table{{border-collapse:collapse;width:100%}} th,td{{border:1px solid #777;padding:8px;text-align:right}}</style>
</head><body>
<h2>التدفق المالي {period}</h2>
<table><thead><tr><th>البيان</th><th>القيمة</th></tr></thead><tbody>
<tr><td>عدد الفواتير</td><td>{stats['sales_count']}</td></tr>
<tr><td>إجمالي المبيعات</td><td>{stats['sales_total']:.2f}</td></tr>
<tr><td>عدد الحجوزات</td><td>{stats['res_count']}</td></tr>
<tr><td>إجمالي الحجوزات</td><td>{stats['res_total']:.2f}</td></tr>
<tr><td>المبلغ المدفوع من الحجوزات</td><td>{stats['res_paid']:.2f}</td></tr>
</tbody></table>
</body></html>"""
            import tempfile, os
            path = os.path.join(tempfile.gettempdir(), "money_flow.html")
            with open(path, "w", encoding="utf-8") as f:
                f.write(html)
            _print_html_auto(path, copies=1, parent=self)
        except Exception as ex:
            messagebox.showerror("خطأ", str(ex), parent=self)

    def _print_movements(self):
        try:
            rows_html = ""
            for iid in self._mv_table.get_children():
                vals = self._mv_table.item(iid, "values")
                rows_html += "<tr>" + "".join(f"<td>{_html(str(v))}</td>" for v in vals) + "</tr>"
            html = f"""<!DOCTYPE html>
<html lang='ar' dir='rtl'>
<head><meta charset='utf-8'><title>حركة الأصناف</title>
<style>body{{font-family:Tahoma,Arial;margin:20px}} table{{border-collapse:collapse;width:100%}} th,td{{border:1px solid #777;padding:6px;text-align:center}}</style>
</head><body>
<h2>حركة الأصناف</h2>
<table><thead><tr><th>النوع</th><th>المدرسة</th><th>اللون</th><th>المقاس</th><th>وارد</th><th>مباع</th><th>محجوز</th><th>متبقي</th></tr></thead>
<tbody>{rows_html}</tbody></table>
</body></html>"""
            import tempfile, os
            path = os.path.join(tempfile.gettempdir(), "item_movements.html")
            with open(path, "w", encoding="utf-8") as f:
                f.write(html)
            _print_html_auto(path, copies=1, parent=self)
        except Exception as ex:
            messagebox.showerror("خطأ", str(ex), parent=self)

    def _print_income(self):
        html = f"""<!DOCTYPE html>
<html lang="ar" dir="rtl"><head><meta charset="utf-8"><title>إحصائيات الوارد</title>
<style>@page {{ size: 80mm auto; margin: 2mm; }}
body {{ font-family: "Segoe UI", Tahoma, sans-serif; font-size: 11px; width: 76mm; direction: rtl; margin:0; padding:2mm; }}
h2 {{ font-size: 14px; text-align: center; margin: 4px 0; }}
.sep {{ border:none; border-top:1px dashed #000; margin:4px 0; }}
.row {{ margin: 3px 0; }}
</style></head><body>
<h2>إحصائيات الوارد</h2>
<hr class="sep">
<div class="row">عدد فواتير الوارد: {self._income_count_var.get()}</div>
<div class="row">إجمالي القطع الواردة: {self._income_qty_var.get()}</div>
<div class="row">إجمالي قيمة الوارد: {self._income_total_var.get()}</div>
</body></html>"""
        tmp = os.path.join(tempfile.gettempdir(), "income_stats.html")
        with open(tmp, "w", encoding="utf-8") as f:
            f.write(html)
        _print_html_auto(tmp, copies=1, parent=self)


# ------------------- Shifts Summary Frame -------------------

class ShiftsSummaryFrame(ttk.Frame):
    """Tab showing history and details of all shifts."""

    def __init__(self, master, db: "SqliteDatabase"):
        super().__init__(master, padding=6)
        self.db = db
        self._build()

    def _build(self):
        ttk.Label(self, text="ملخص الورديات", font=("Segoe UI", 12, "bold")).pack(anchor="w", pady=(0, 6))

        # ---- Filter bar ----
        filter_bar = ttk.Frame(self)
        filter_bar.pack(fill=tk.X, pady=(0, 6))

        ttk.Label(filter_bar, text="من:").pack(side=tk.LEFT)
        self._df = DateField(filter_bar, "")
        self._df.pack(side=tk.LEFT, padx=(4, 12))

        ttk.Label(filter_bar, text="إلى:").pack(side=tk.LEFT)
        self._dt = DateField(filter_bar, "")
        self._dt.pack(side=tk.LEFT, padx=(4, 12))

        ttk.Button(filter_bar, text="تحديث", command=self._refresh_all).pack(side=tk.LEFT, padx=4)
        ttk.Button(filter_bar, text="مسح", command=self._clear_filters).pack(side=tk.LEFT, padx=4)
        ttk.Button(filter_bar, text="طباعة الكل", command=self._print_all_shifts).pack(side=tk.RIGHT, padx=4)

        # ---- Shifts table ----
        table_frame = ttk.Frame(self)
        table_frame.pack(fill=tk.BOTH, expand=True, pady=(0, 6))

        cols = ("id", "started", "ended", "status", "sales_count", "sales_total",
                "res_count", "res_paid", "deliver_total", "inflow_qty", "cash")
        self._tree = ttk.Treeview(table_frame, columns=cols, show="headings", height=10)
        for col, txt, w in [
            ("id", "#", 40), ("started", "البداية", 130), ("ended", "النهاية", 130),
            ("status", "الحالة", 70), ("sales_count", "فواتير", 60),
            ("sales_total", "مبيعات", 90), ("res_count", "حجوزات", 60),
            ("res_paid", "مدفوع حجز", 90), ("deliver_total", "تسليمات", 90),
            ("inflow_qty", "وارد (قطع)", 80), ("cash", "إجمالي نقدية", 100),
        ]:
            self._tree.heading(col, text=txt)
            self._tree.column(col, width=w, anchor="center")
        ysb = ttk.Scrollbar(table_frame, orient="vertical", command=self._tree.yview)
        self._tree.configure(yscrollcommand=ysb.set)
        self._tree.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        ysb.pack(side=tk.RIGHT, fill=tk.Y)
        _bind_mousewheel(self._tree)
        _add_context_menu(self._tree, self)

        self._tree.bind("<<TreeviewSelect>>", self._on_shift_selected)

        # ---- Detail panel ----
        self._detail_frame = ttk.LabelFrame(self, text="تفاصيل الوردية")
        self._detail_frame.pack(fill=tk.X, pady=(0, 4))

        detail_top = ttk.Frame(self._detail_frame)
        detail_top.pack(fill=tk.X, padx=8, pady=4)

        self._det_id_var = tk.StringVar(value="-")
        self._det_period_var = tk.StringVar(value="-")
        self._det_status_var = tk.StringVar(value="-")

        ttk.Label(detail_top, text="الوردية:").pack(side=tk.LEFT)
        ttk.Label(detail_top, textvariable=self._det_id_var, font=("Segoe UI", 10, "bold")).pack(side=tk.LEFT, padx=(4, 16))
        ttk.Label(detail_top, text="الفترة:").pack(side=tk.LEFT)
        ttk.Label(detail_top, textvariable=self._det_period_var, font=("Segoe UI", 10, "bold")).pack(side=tk.LEFT, padx=(4, 16))
        ttk.Label(detail_top, text="الحالة:").pack(side=tk.LEFT)
        ttk.Label(detail_top, textvariable=self._det_status_var, font=("Segoe UI", 10, "bold")).pack(side=tk.LEFT, padx=(4, 0))

        detail_nums = ttk.Frame(self._detail_frame)
        detail_nums.pack(fill=tk.X, padx=8, pady=4)

        self._det_sales_var = tk.StringVar(value="-")
        self._det_res_var = tk.StringVar(value="-")
        self._det_deliver_var = tk.StringVar(value="-")
        self._det_inflow_var = tk.StringVar(value="-")
        self._det_return_var = tk.StringVar(value="-")
        self._det_exchange_var = tk.StringVar(value="-")
        self._det_cash_var = tk.StringVar(value="-")

        for lbl, var in [
            ("المبيعات:", self._det_sales_var),
            ("الحجوزات:", self._det_res_var),
            ("التسليمات:", self._det_deliver_var),
            ("الوارد:", self._det_inflow_var),
        ]:
            ttk.Label(detail_nums, text=lbl).pack(side=tk.LEFT)
            ttk.Label(detail_nums, textvariable=var, font=("Segoe UI", 10, "bold")).pack(side=tk.LEFT, padx=(4, 16))

        detail_nums2 = ttk.Frame(self._detail_frame)
        detail_nums2.pack(fill=tk.X, padx=8, pady=(0, 4))
        for lbl, var in [
            ("المرتجعات:", self._det_return_var),
            ("الاستبدالات:", self._det_exchange_var),
        ]:
            ttk.Label(detail_nums2, text=lbl).pack(side=tk.LEFT)
            ttk.Label(detail_nums2, textvariable=var, font=("Segoe UI", 10, "bold")).pack(side=tk.LEFT, padx=(4, 16))

        cash_row = ttk.Frame(self._detail_frame)
        cash_row.pack(fill=tk.X, padx=8, pady=(0, 4))
        ttk.Label(cash_row, text="إجمالي النقدية المحصلة:", font=("Segoe UI", 10, "bold")).pack(side=tk.LEFT)
        ttk.Label(cash_row, textvariable=self._det_cash_var, font=("Segoe UI", 12, "bold")).pack(side=tk.LEFT, padx=(4, 0))

        detail_btns = ttk.Frame(self._detail_frame)
        detail_btns.pack(fill=tk.X, padx=8, pady=(0, 6))
        ttk.Button(detail_btns, text="طباعة ملخص الوردية", command=self._print_selected_shift).pack(side=tk.LEFT)

        # ---- Totals bar ----
        totals = ttk.LabelFrame(self, text="الإجماليات")
        totals.pack(fill=tk.X)

        totals_row = ttk.Frame(totals)
        totals_row.pack(fill=tk.X, padx=8, pady=4)

        self._tot_shifts_var = tk.StringVar(value="0")
        self._tot_sales_var = tk.StringVar(value="0.00")
        self._tot_cash_var = tk.StringVar(value="0.00")

        ttk.Label(totals_row, text="عدد الورديات:").pack(side=tk.LEFT)
        ttk.Label(totals_row, textvariable=self._tot_shifts_var, font=("Segoe UI", 10, "bold")).pack(side=tk.LEFT, padx=(4, 20))
        ttk.Label(totals_row, text="إجمالي المبيعات:").pack(side=tk.LEFT)
        ttk.Label(totals_row, textvariable=self._tot_sales_var, font=("Segoe UI", 10, "bold")).pack(side=tk.LEFT, padx=(4, 20))
        ttk.Label(totals_row, text="إجمالي النقدية:").pack(side=tk.LEFT)
        ttk.Label(totals_row, textvariable=self._tot_cash_var, font=("Segoe UI", 12, "bold")).pack(side=tk.LEFT, padx=(4, 0))

        self._shifts_data: List[Dict[str, Any]] = []
        self._refresh_all()

    def _refresh_all(self):
        df = self._df.get() or None
        dt = self._dt.get() or None
        try:
            self._shifts_data = self.db.get_all_shifts(date_from=df, date_to=dt)
        except Exception:
            self._shifts_data = []

        self._tree.delete(*self._tree.get_children())
        total_sales = 0.0
        total_cash = 0.0
        for s in self._shifts_data:
            started = s["started_at"][:16].replace("T", " ")
            ended = s["ended_at"][:16].replace("T", " ") if s["ended_at"] else "-"
            status = "مفتوحة" if s["status"] == "OPEN" else "مغلقة"
            self._tree.insert("", tk.END, values=(
                s["id"], started, ended, status,
                s["sales_count"], f"{s['sales_total']:.2f}",
                s["res_count"], f"{s['res_paid']:.2f}",
                f"{s['deliver_total']:.2f}",
                s["inflow_total_qty"],
                f"{s['cash_collected']:.2f}",
            ))
            total_sales += s["sales_total"]
            total_cash += s["cash_collected"]
        _apply_zebra_tags(self._tree)

        self._tot_shifts_var.set(str(len(self._shifts_data)))
        self._tot_sales_var.set(f"{total_sales:.2f}")
        self._tot_cash_var.set(f"{total_cash:.2f}")

        # Clear detail
        self._det_id_var.set("-")
        self._det_period_var.set("-")
        self._det_status_var.set("-")
        self._det_sales_var.set("-")
        self._det_res_var.set("-")
        self._det_deliver_var.set("-")
        self._det_inflow_var.set("-")
        self._det_return_var.set("-")
        self._det_exchange_var.set("-")
        self._det_cash_var.set("-")

    def _on_shift_selected(self, _event=None):
        sel = self._tree.selection()
        if not sel:
            return
        vals = self._tree.item(sel[0], "values")
        shift_id = int(vals[0])
        s = next((x for x in self._shifts_data if x["id"] == shift_id), None)
        if not s:
            return

        started = s["started_at"][:16].replace("T", " ")
        ended = s["ended_at"][:16].replace("T", " ") if s["ended_at"] else "الآن"
        status = "مفتوحة" if s["status"] == "OPEN" else "مغلقة"

        self._det_id_var.set(f"#{s['id']}")
        self._det_period_var.set(f"{started}  →  {ended}")
        self._det_status_var.set(status)
        self._det_sales_var.set(f"{s['sales_count']} فاتورة - {s['sales_total']:.2f}")
        self._det_res_var.set(f"{s['res_count']} حجز - مدفوع {s['res_paid']:.2f} من {s['res_total']:.2f}")
        self._det_deliver_var.set(f"{s['deliver_count']} عملية - {s['deliver_total']:.2f}")
        self._det_inflow_var.set(f"{s['inflow_count']} عملية - {s['inflow_total_qty']} قطعة")
        self._det_return_var.set(f"{s.get('return_count', 0)} فاتورة - {s.get('return_total', 0.0):.2f}")
        self._det_exchange_var.set(f"{s.get('exchange_count', 0)} فاتورة - صافي {s.get('exchange_total', 0.0):.2f}")
        self._det_cash_var.set(f"{s['cash_collected']:.2f}")

    def _clear_filters(self):
        self._df.set("")
        self._dt.set("")
        self._refresh_all()

    def _get_selected_shift(self) -> Optional[Dict[str, Any]]:
        sel = self._tree.selection()
        if not sel:
            messagebox.showinfo("تنبيه", "اختر وردية أولاً", parent=self)
            return None
        shift_id = int(self._tree.item(sel[0], "values")[0])
        return next((x for x in self._shifts_data if x["id"] == shift_id), None)

    def _print_selected_shift(self):
        s = self._get_selected_shift()
        if not s:
            return
        started = s["started_at"][:16].replace("T", " ")
        ended = s["ended_at"][:16].replace("T", " ") if s["ended_at"] else "الآن"
        status = "مفتوحة" if s["status"] == "OPEN" else "مغلقة"
        html = f"""<!DOCTYPE html>
<html lang="ar" dir="rtl"><head><meta charset="utf-8"><title>ملخص وردية</title>
<style>
@page {{ size: 80mm auto; margin: 2mm; }}
body {{ font-family: "Segoe UI", Tahoma, sans-serif; font-size: 11px; width: 76mm; direction: rtl; margin:0; padding:2mm; }}
h2 {{ font-size: 14px; text-align: center; margin: 4px 0; }}
.sep {{ border:none; border-top:1px dashed #000; margin:4px 0; }}
.total {{ font-size: 13px; font-weight: bold; text-align: center; margin: 6px 0; }}
</style></head><body>
<h2>ملخص الوردية #{s['id']}</h2>
<hr class="sep">
<div>الحالة: {status}</div>
<div>بداية: {started}</div>
<div>نهاية: {ended}</div>
<hr class="sep">
<div><b>المبيعات:</b> {s['sales_count']} فاتورة - {s['sales_total']:.2f}</div>
<hr class="sep">
<div><b>الحجوزات:</b> {s['res_count']} - إجمالي {s['res_total']:.2f} - مدفوع {s['res_paid']:.2f}</div>
<hr class="sep">
<div><b>التسليمات:</b> {s['deliver_count']} عملية - {s['deliver_total']:.2f}</div>
<hr class="sep">
<div><b>الوارد:</b> {s['inflow_count']} عملية - {s['inflow_total_qty']} قطعة</div>
<hr class="sep">
<div class="total">النقدية المحصلة: {s['cash_collected']:.2f}</div>
</body></html>"""
        tmp = os.path.join(tempfile.gettempdir(), f"shift_summary_{s['id']}.html")
        with open(tmp, "w", encoding="utf-8") as f:
            f.write(html)
        _print_html_auto(tmp, copies=1, parent=self)

    def _print_all_shifts(self):
        if not self._shifts_data:
            messagebox.showinfo("تنبيه", "لا توجد ورديات لطباعتها", parent=self)
            return
        rows_html = ""
        total_sales = 0.0
        total_cash = 0.0
        for s in self._shifts_data:
            started = s["started_at"][:16].replace("T", " ")
            ended = s["ended_at"][:16].replace("T", " ") if s["ended_at"] else "-"
            status = "مفتوحة" if s["status"] == "OPEN" else "مغلقة"
            rows_html += f"<tr><td>{s['id']}</td><td>{started}</td><td>{ended}</td><td>{status}</td>"
            rows_html += f"<td>{s['sales_count']}</td><td>{s['sales_total']:.2f}</td>"
            rows_html += f"<td>{s['res_paid']:.2f}</td><td>{s['deliver_total']:.2f}</td>"
            rows_html += f"<td>{s['inflow_total_qty']}</td><td>{s['cash_collected']:.2f}</td></tr>\n"
            total_sales += s["sales_total"]
            total_cash += s["cash_collected"]
        html = f"""<!DOCTYPE html>
<html lang="ar" dir="rtl"><head><meta charset="utf-8"><title>ملخص الورديات</title>
<style>
body {{ font-family: Tahoma, Arial; margin: 20px; direction: rtl; }}
h2 {{ text-align: center; }}
table {{ border-collapse: collapse; width: 100%; font-size: 11px; }}
th {{ border-bottom: 2px solid #000; padding: 6px; text-align: center; background: #f0f0f0; }}
td {{ padding: 5px; border-bottom: 1px solid #ccc; text-align: center; }}
.totals {{ font-weight: bold; font-size: 13px; margin-top: 12px; }}
</style></head><body>
<h2>ملخص جميع الورديات ({len(self._shifts_data)} وردية)</h2>
<table><thead><tr>
<th>#</th><th>البداية</th><th>النهاية</th><th>الحالة</th>
<th>فواتير</th><th>مبيعات</th><th>مدفوع حجز</th><th>تسليمات</th>
<th>وارد</th><th>نقدية</th>
</tr></thead><tbody>
{rows_html}
</tbody></table>
<div class="totals">إجمالي المبيعات: {total_sales:.2f} | إجمالي النقدية: {total_cash:.2f}</div>
</body></html>"""
        tmp = os.path.join(tempfile.gettempdir(), "all_shifts_summary.html")
        with open(tmp, "w", encoding="utf-8") as f:
            f.write(html)
        _print_html_auto(tmp, copies=1, parent=self)


# ------------------- Reservations Frame -------------------

class ReservationsFrame(ttk.Frame):
    """Standalone tab for managing reservations."""
    def __init__(self, master, db: SqliteDatabase):
        super().__init__(master, padding=6)
        self.db = db
        self._build()

    def _build(self):
        # ---- Filter row ----
        flt = ttk.Frame(self)
        flt.pack(fill=tk.X, padx=4, pady=(0, 6))

        self._rf_school = LabeledCombobox(flt, "المدرسة", self.db, "school")
        self._rf_school.pack(side=tk.RIGHT, padx=4)
        self._rf_school.set_supplier(lambda: self.db.get_distinct_filtered("school", self._build_res_filter_constraints("school")))

        self._rf_item = LabeledCombobox(flt, "النوع", self.db, "item_type")
        self._rf_item.pack(side=tk.RIGHT, padx=4)
        self._rf_item.set_supplier(lambda: self.db.get_distinct_filtered("item_type", self._build_res_filter_constraints("item_type")))

        self._rf_color = LabeledCombobox(flt, "اللون", self.db, "color")
        self._rf_color.pack(side=tk.RIGHT, padx=4)
        self._rf_color.set_supplier(lambda: self.db.get_distinct_filtered("color", self._build_res_filter_constraints("color")))

        # Status filter
        ttk.Label(flt, text="الحالة:").pack(side=tk.RIGHT, padx=(4, 0))
        self._rf_status = ttk.Combobox(flt, values=["", "معلق", "تم التسليم"], state="readonly", width=12)
        self._rf_status.pack(side=tk.RIGHT, padx=4)
        self._rf_status.bind("<<ComboboxSelected>>", lambda e: self._refresh())

        # Date filters
        self._rf_df = DateField(flt, "من:")
        self._rf_df.pack(side=tk.RIGHT, padx=4)
        self._rf_dt = DateField(flt, "إلى:")
        self._rf_dt.pack(side=tk.RIGHT, padx=4)

        ttk.Button(flt, text="تحديث", command=self._refresh).pack(side=tk.RIGHT, padx=4)
        ttk.Button(flt, text="مسح", command=self._clear_filters).pack(side=tk.RIGHT, padx=4)

        for w in (self._rf_school.cb, self._rf_item.cb, self._rf_color.cb):
            w.bind("<<ComboboxSelected>>", lambda e: self._on_filter_changed(), add="+")

        # ---- Reservations table ----
        tbl_wrap = ttk.Frame(self)
        tbl_wrap.pack(fill=tk.BOTH, expand=True, padx=4, pady=4)

        self._res_table = ttk.Treeview(
            tbl_wrap,
            columns=("id", "created", "customer", "item", "school", "color", "size", "qty", "total", "paid", "status"),
            show="headings",
            height=12,
        )
        for col, txt, w in [
            ("id", "رقم", 40), ("created", "التاريخ", 100), ("customer", "العميل", 100),
            ("item", "النوع", 90), ("school", "المدرسة", 110), ("color", "اللون", 70),
            ("size", "المقاس", 55), ("qty", "الكمية", 55), ("total", "الإجمالي", 80),
            ("paid", "المدفوع", 80), ("status", "الحالة", 80),
        ]:
            self._res_table.heading(col, text=txt)
            self._res_table.column(col, width=w, anchor="center")
        rv_ysb = ttk.Scrollbar(tbl_wrap, orient="vertical", command=self._res_table.yview)
        self._res_table.configure(yscrollcommand=rv_ysb.set)
        self._res_table.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        rv_ysb.pack(side=tk.RIGHT, fill=tk.Y)
        _bind_mousewheel(self._res_table)

        # Context menu
        _add_context_menu(self._res_table, self)

        # ---- Buttons row ----
        btns = ttk.Frame(self)
        btns.pack(fill=tk.X, padx=4, pady=4)
        ttk.Button(btns, text="تسليم الحجز", command=self._deliver_reservation).pack(side=tk.RIGHT, padx=(4, 0))
        ttk.Button(btns, text="طباعة القائمة", command=self._print_reservations).pack(side=tk.RIGHT)

        self._refresh()

    def _build_res_filter_constraints(self, exclude_field: str) -> Dict[str, Any]:
        c: Dict[str, Any] = {}
        if exclude_field != "school":
            v = self._rf_school.get()
            if v: c["school"] = v
        if exclude_field != "item_type":
            v = self._rf_item.get()
            if v: c["item_type"] = v
        if exclude_field != "color":
            v = self._rf_color.get()
            if v: c["color"] = v
        return c

    def _on_filter_changed(self):
        self._rf_school.refresh_values()
        self._rf_item.refresh_values()
        self._rf_color.refresh_values()
        self._refresh()

    def _clear_filters(self):
        self._rf_school.set("")
        self._rf_item.set("")
        self._rf_color.set("")
        self._rf_status.set("")
        self._rf_df.set("")
        self._rf_dt.set("")
        self._rf_school.refresh_values()
        self._rf_item.refresh_values()
        self._rf_color.refresh_values()
        self._refresh()

    def _refresh(self):
        df = self._rf_df.get() or None
        dt = self._rf_dt.get() or None
        school = self._rf_school.get() or None
        item_type = self._rf_item.get() or None
        color = self._rf_color.get() or None
        status = self._rf_status.get() or None
        try:
            self._res_table.delete(*self._res_table.get_children())
            for r in self.db.list_reservations(status=status, date_from=df, date_to=dt,
                                                school=school, item_type=item_type, color=color):
                self._res_table.insert("", tk.END, values=(
                    r["id"], r["created_at"][:10], r.get("customer", ""),
                    r["item_type"], r["school"], r["color"], r["size"],
                    r["qty"], f"{float(r['total_amount']):.2f}",
                    f"{float(r['paid_amount']):.2f}", r["status"]
                ))
            _apply_zebra_tags(self._res_table)
        except Exception:
            pass

    def _deliver_reservation(self):
        """Mark selected reservation as delivered and collect remaining payment."""
        sel = self._res_table.selection()
        if not sel:
            messagebox.showwarning("تنبيه", "اختر حجزاً من الجدول أولاً.", parent=self)
            return
        vals = self._res_table.item(sel[0], "values")
        res_id = int(vals[0])
        status = vals[10]
        if status == "تم التسليم":
            messagebox.showinfo("تنبيه", "تم تسليم هذا الحجز مسبقاً.", parent=self)
            return
        total = float(vals[8])
        paid = float(vals[9])
        remaining = total - paid

        dlg = tk.Toplevel(self)
        dlg.title("تسليم حجز")
        dlg.geometry("350x220")
        dlg.transient(self.winfo_toplevel())
        dlg.grab_set()

        frm = ttk.Frame(dlg, padding=12)
        frm.pack(fill=tk.BOTH, expand=True)

        ttk.Label(frm, text=f"حجز رقم: {res_id}", font=("Segoe UI", 11, "bold")).pack(anchor="w", pady=(0, 4))
        ttk.Label(frm, text=f"الإجمالي: {total:.2f}").pack(anchor="w")
        ttk.Label(frm, text=f"المدفوع سابقاً: {paid:.2f}").pack(anchor="w")
        ttk.Label(frm, text=f"المتبقي: {remaining:.2f}", font=("Segoe UI", 10, "bold")).pack(anchor="w", pady=(0, 8))

        ttk.Label(frm, text="المبلغ المحصّل الآن:").pack(anchor="w")
        collect_var = tk.StringVar(value=f"{remaining:.2f}" if remaining > 0 else "0")
        ttk.Entry(frm, textvariable=collect_var, width=15).pack(anchor="w", pady=(0, 8))

        def _confirm():
            try:
                collected = _parse_money_amount(collect_var.get())
            except ValueError:
                messagebox.showerror("خطأ", "أدخل مبلغاً صحيحاً.", parent=dlg)
                return
            try:
                self.db.deliver_reservation(res_id, collected)
                dlg.destroy()
                ToastNotification.show(self.winfo_toplevel(), "تم تسليم الحجز بنجاح", toast_type="success")
                self._refresh()
                try:
                    ac = getattr(self.winfo_toplevel(), "_app_controller", None)
                    if ac is not None and hasattr(ac, "refresh_dashboard"):
                        ac.refresh_dashboard()
                except Exception:
                    pass
            except Exception as ex:
                messagebox.showerror("خطأ", str(ex), parent=dlg)

        ttk.Button(frm, text="تأكيد التسليم", command=_confirm).pack(anchor="e")

    def _print_reservations(self):
        try:
            rows_html = ""
            for iid in self._res_table.get_children():
                vals = self._res_table.item(iid, "values")
                rows_html += "<tr>" + "".join(f"<td>{_html(str(v))}</td>" for v in vals) + "</tr>"
            html = f"""<!DOCTYPE html>
<html lang='ar' dir='rtl'>
<head><meta charset='utf-8'><title>الحجوزات</title>
<style>body{{font-family:Tahoma,Arial;margin:20px}} table{{border-collapse:collapse;width:100%}} th,td{{border:1px solid #777;padding:6px;text-align:center}}</style>
</head><body>
<h2>الحجوزات</h2>
<table><thead><tr><th>رقم</th><th>التاريخ</th><th>العميل</th><th>النوع</th><th>المدرسة</th><th>اللون</th><th>المقاس</th><th>الكمية</th><th>الإجمالي</th><th>المدفوع</th><th>الحالة</th></tr></thead>
<tbody>{rows_html}</tbody></table>
</body></html>"""
            import tempfile, os
            path = os.path.join(tempfile.gettempdir(), "reservations.html")
            with open(path, "w", encoding="utf-8") as f:
                f.write(html)
            _print_html_auto(path, copies=1, parent=self)
        except Exception as ex:
            messagebox.showerror("خطأ", str(ex), parent=self)


# ------------------- Inventory Window -------------------

class InventoryWindow(tk.Toplevel):
    def __init__(self, master, db: SqliteDatabase):
        super().__init__(master)
        self.db = db
        self.title("المخزون")
        self.geometry("1000x560")

        self.multi: Dict[str, List[Any]] = {
            "item_type": [], "school": [], "color": [], "size": [],
        }

        self._multi_btns: Dict[str, ttk.Button] = {}
        self._field_widgets: Dict[str, tk.Widget] = {}

        self._build()

    def _build(self):
        restricted = []
        if not self.db.is_manager_feature_enabled("allow_inventory_delete"):
            restricted.append("الحذف")
        if not self.db.is_manager_feature_enabled("allow_inventory_price_edit"):
            restricted.append("السعر")
        if not self.db.is_manager_feature_enabled("allow_inventory_specs_edit"):
            restricted.append("المواصفات")
        if restricted:
            ttk.Label(
                self,
                text=f"صلاحيات مقيدة من المدير في هذه النافذة: {', '.join(restricted)}",
                foreground="#b91c1c",
                font=("Segoe UI", 10, "bold"),
            ).pack(fill=tk.X, padx=8, pady=(8, 0))

        filters = ttk.LabelFrame(self, text="تصنيف")
        filters.pack(fill=tk.X, padx=8, pady=8)

        self.f_type   = LabeledCombobox(filters, "النوع",    self.db, "item_type")
        self.f_school = LabeledCombobox(filters, "المدرسة", self.db, "school")
        self.f_color  = LabeledCombobox(filters, "اللون",    self.db, "color")
        self.f_size   = LabeledCombobox(filters, "المقاس",   self.db, "size")

        def _constraints(exclude=None):
            d = {
                "item_type": self.f_type.get(),
                "school":    self.f_school.get(),
                "color":     self.f_color.get(),
                "size":      self.f_size.get(),
            }
            if exclude:
                d.pop(exclude, None)
            return d

        self.f_type.set_supplier(  lambda: self.db.get_distinct_filtered("item_type", _constraints(exclude="item_type")))
        self.f_school.set_supplier(lambda: self.db.get_distinct_filtered("school",    _constraints(exclude="school")))
        self.f_color.set_supplier( lambda: self.db.get_distinct_filtered("color",     _constraints(exclude="color")))
        self.f_size.set_supplier(  lambda: self.db.get_distinct_filtered("size",      _constraints(exclude="size")))

        def _refresh_all_suppliers(*_):
            self.f_type.refresh_values()
            self.f_school.refresh_values()
            self.f_color.refresh_values()
            self.f_size.refresh_values()

        for w in (self.f_type, self.f_school, self.f_color, self.f_size):
            for ev in ("<<ComboboxSelected>>", "<KeyRelease>"):
                w.cb.bind(ev, lambda e: (_refresh_all_suppliers(), self._schedule_refresh()), add="+")

        widgets = [self.f_type, self.f_school, self.f_color, self.f_size]
        fields  = ["item_type","school","color","size"]

        for i, (w, fld) in enumerate(zip(widgets, fields)):
            r, c = divmod(i, 3)
            w.grid(row=r*2, column=c, padx=6, pady=(6, 0), sticky="ew")
            filters.columnconfigure(c, weight=1)

            from functools import partial

            def _safe_open(field):
                try:
                    self._open_multi_dialog(field)
                except Exception as ex:
                    messagebox.showerror("خطأ", str(ex), parent=self)

            btn = ttk.Button(filters, text="اختيار متعدد...", command=partial(_safe_open, fld))
            btn.grid(row=r*2+1, column=c, sticky="w", padx=6, pady=(2, 8))
            self._multi_btns[fld] = btn

            if hasattr(w, "cb"):
                self._field_widgets[fld] = w.cb
            elif hasattr(w, "entry"):
                self._field_widgets[fld] = w.entry
            else:
                self._field_widgets[fld] = w

        for cb in (self.f_type.cb, self.f_school.cb, self.f_color.cb, self.f_size.cb):
            cb.bind("<<ComboboxSelected>>", lambda e: self._schedule_refresh(), add="+")
            cb.bind("<KeyRelease>",         lambda e: self._schedule_refresh(), add="+")

        btns = ttk.Frame(filters)
        btns.grid(row=4, column=0, columnspan=3, sticky="w", padx=6, pady=(0, 6))
        ttk.Button(btns, text="بحث", command=self._refresh).pack(side=tk.LEFT)
        ttk.Button(btns, text="مسح", command=self._clear_all).pack(side=tk.LEFT, padx=8)
        ttk.Button(btns, text="تصدير إلى إكسل", command=self._export_excel).pack(side=tk.LEFT, padx=8)
        ttk.Button(btns, text="طباعة جداول المقاسات", command=self._print_size_sheets).pack(side=tk.LEFT, padx=8)
        size_ranges_btn = ttk.Button(btns, text="تعديل نطاقات المقاسات...", command=self._edit_size_ranges_dialog)
        size_ranges_btn.pack(side=tk.LEFT, padx=8)
        if not self.db.is_manager_feature_enabled("allow_size_profile_edit"):
            size_ranges_btn.configure(state="disabled")

        table_wrap = ttk.Frame(self)
        table_wrap.pack(fill=tk.BOTH, expand=True, padx=8, pady=(0, 4))
        self.table = ttk.Treeview(
            table_wrap,
            columns=("id", "type", "school", "color", "size", "price", "count", "value"),
            show="headings",
            selectmode="extended",
        )
        for col, txt, w in [
            ("id","المعرّف",60), ("type","النوع",140), ("school","المدرسة",160),
            ("color","اللون",80), ("size","المقاس",70),
            ("price","السعر",80), ("count","الكمية",70), ("value","القيمة",90),
        ]:
            self.table.heading(col, text=txt)
            self.table.column(col, width=w, anchor="center")
        ysb = ttk.Scrollbar(table_wrap, orient="vertical", command=self.table.yview)
        xsb = ttk.Scrollbar(table_wrap, orient="horizontal", command=self.table.xview)
        self.table.configure(yscrollcommand=ysb.set, xscrollcommand=xsb.set)
        self.table.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        ysb.pack(side=tk.RIGHT, fill=tk.Y)
        xsb.pack(side=tk.BOTTOM, fill=tk.X)
        _bind_mousewheel(self.table)

        # Context menu
        _add_context_menu(self.table, self)

        bar = ttk.Frame(self)
        bar.pack(fill=tk.X, padx=8, pady=(0, 8))
        self.sum_qty = tk.StringVar(value="0")
        self.sum_val = tk.StringVar(value="0.00")
        ttk.Label(bar, text="إجمالي الكمية:").pack(side=tk.LEFT)
        ttk.Label(bar, textvariable=self.sum_qty, font=("Segoe UI", 10, "bold")).pack(side=tk.LEFT, padx=(4, 12))
        ttk.Label(bar, text="إجمالي القيمة:").pack(side=tk.LEFT)
        ttk.Label(bar, textvariable=self.sum_val, font=("Segoe UI", 10, "bold")).pack(side=tk.LEFT, padx=(4, 12))
        edit_price_btn = ttk.Button(bar, text="تعديل السعر...", command=self._edit_price_dialog)
        edit_price_btn.pack(side=tk.RIGHT)
        edit_specs_btn = ttk.Button(bar, text="تعديل المواصفات...", command=self._edit_specs_dialog)
        edit_specs_btn.pack(side=tk.RIGHT, padx=(8, 0))
        remove_btn = ttk.Button(bar, text="حذف المحدد...", command=self._remove_selected_dialog)
        remove_btn.pack(side=tk.RIGHT, padx=(8, 0))
        if not self.db.is_manager_feature_enabled("allow_inventory_price_edit"):
            edit_price_btn.configure(state="disabled")
        if not self.db.is_manager_feature_enabled("allow_inventory_specs_edit"):
            edit_specs_btn.configure(state="disabled")
        if not self.db.is_manager_feature_enabled("allow_inventory_delete"):
            remove_btn.configure(state="disabled")

        self._refresh()

    def _get_selected_profile_keys(self):
        sel = self.table.selection()
        if not sel:
            return None
        keys = set()
        for iid in sel:
            vals = self.table.item(iid, "values")
            keys.add((vals[1], vals[2], vals[3]))
        if len(keys) != 1:
            messagebox.showwarning("تحديد غير صالح",
                "يجب أن تكون الصفوف المحددة من نفس (النوع، المدرسة، اللون).", parent=self)
            return None
        return keys.pop()

    def _edit_size_ranges_dialog(self):
        if not self.db.is_manager_feature_enabled("allow_size_profile_edit"):
            messagebox.showwarning("مقيد", _feature_restricted_message("تعديل نطاقات المقاسات من نقطة البيع غير مسموح به حالياً."), parent=self)
            return
        picked = self._get_selected_profile_keys()
        if picked:
            item_type, school, color = picked
        else:
            item_type = (self.f_type.get() or "").strip()
            school    = (self.f_school.get() or "").strip()
            color     = (self.f_color.get() or "").strip()
            if not (item_type and school and color):
                messagebox.showwarning("حدد الصنف",
                    "حدد صفوفاً من الجدول أو اختر (النوع، المدرسة، اللون) أولاً.", parent=self)
                return

        dlg = tk.Toplevel(self)
        dlg.title("تعديل نطاقات المقاسات")
        dlg.transient(self)
        dlg.grab_set()

        frm = ttk.Frame(dlg, padding=12)
        frm.pack(fill=tk.BOTH, expand=True)

        ttk.Label(frm, text=f"{item_type} / {school} / {color}", font=("Segoe UI", 10, "bold"))\
            .grid(row=0, column=0, columnspan=4, sticky="w", pady=(0, 10))

        r1_var = tk.StringVar()
        r2_var = tk.StringVar()
        has_alpha_var = tk.BooleanVar(value=False)

        ttk.Label(frm, text="النطاق الأول").grid(row=0, column=0, sticky="w")
        r1_combo = ttk.Combobox(frm, textvariable=r1_var, values=[""] + NUMERIC_RANGE_LABELS,
                                state="readonly", width=12)
        r1_combo.grid(row=0, column=1, columnspan=2, sticky="w")

        ttk.Label(frm, text="النطاق الثاني (اختياري)").grid(row=1, column=0, sticky="w")
        r2_combo = ttk.Combobox(frm, textvariable=r2_var, values=[""] + NUMERIC_RANGE_LABELS,
                                state="readonly", width=12)
        r2_combo.grid(row=1, column=1, columnspan=2, sticky="w")

        profile = self.db.get_size_profile(item_type, school, color)
        if profile:
            r1s, r1e, r2s, r2e, has_alpha = profile
            if r1s is not None and r1e is not None:
                r1_var.set(f"{r1s} \\u2192 {r1e}")
            if r2s is not None and r2e is not None:
                r2_var.set(f"{r2s} \\u2192 {r2e}")
            has_alpha_var.set(bool(has_alpha))

        ttk.Checkbutton(frm, text="تفعيل المقاسات بالحروف (S / M / L ...)",
                        variable=has_alpha_var).grid(row=2, column=0, columnspan=3, sticky="w")

        btns = ttk.Frame(frm)
        btns.grid(row=6, column=0, columnspan=4, sticky="e", pady=(10,0))

        def _parse_range_label(v: str):
            v = (v or "").strip()
            if not v:
                return None, None
            a, b = v.split("\\u2192")
            return int(a.strip()), int(b.strip())

        def on_save():
            try:
                r1s, r1e = _parse_range_label(r1_var.get())
                r2s, r2e = _parse_range_label(r2_var.get())
                self.db.upsert_size_profile(item_type, school, color,
                    r1_start=r1s, r1_end=r1e, r2_start=r2s, r2_end=r2e,
                    has_alpha=has_alpha_var.get())
                dlg.destroy()
                ToastNotification.show(self.winfo_toplevel(), "تم حفظ نطاقات المقاسات", toast_type="success")
            except Exception as ex:
                messagebox.showerror("خطأ", str(ex), parent=dlg)

        ttk.Button(btns, text="إلغاء", command=dlg.destroy).pack(side=tk.RIGHT)
        ttk.Button(btns, text="حفظ", command=on_save).pack(side=tk.RIGHT, padx=6)

    def _open_multi_dialog(self, field: str):
        try:
            values = self.db.get_distinct(field)
        except Exception:
            values = []

        dlg = MultiSelectDialog(self, title="اختيار متعدد", values=values,
                                preselected=[str(x) for x in self.multi[field]])
        picked = dlg.run()
        if picked is None:
            return

        self.multi[field] = [str(x) for x in picked]

        btn = self._multi_btns.get(field)
        if btn:
            btn.configure(text=f"اختيار متعدد... ({len(self.multi[field])})" if self.multi[field] else "اختيار متعدد...")

        self._enforce_single_active_field()
        self._refresh()

    def _enforce_single_active_field(self):
        active_fields = [k for k, v in self.multi.items() if v]
        if len(active_fields) >= 1:
            active = active_fields[0]
            for fld, w in self._field_widgets.items():
                is_active = (fld == active)
                if self.multi[active]:
                    state = ("normal" if is_active else "disabled")
                else:
                    state = "normal"
                try:
                    w.configure(state=state)
                except Exception:
                    pass

            for fld, b in self._multi_btns.items():
                if self.multi[active]:
                    b.configure(state=("normal" if fld == active else "disabled"))
                else:
                    b.configure(state="normal")

            if self.multi[active]:
                for fld in self._field_widgets:
                    if fld != active:
                        self._clear_single_field_text(fld)
        else:
            for w in self._field_widgets.values():
                try: w.configure(state="normal")
                except Exception: pass
            for b in self._multi_btns.values():
                b.configure(state="normal")

    def _clear_single_field_text(self, field: str):
        if field == "item_type": self.f_type.set("")
        elif field == "school": self.f_school.set("")
        elif field == "color": self.f_color.set("")
        elif field == "size": self.f_size.set("")

    def _print_size_sheets(self):
        try:
            rows = self.db.current_inventory(self._filters())
        except Exception as ex:
            messagebox.showerror("فشل", str(ex), parent=self)
            return
        if not rows:
            messagebox.showinfo("لا توجد بيانات", "لا توجد صفوف مطابقة للطباعة.", parent=self)
            return
        from collections import defaultdict, OrderedDict

        school_groups = OrderedDict()
        for r in rows:
            school = (r.get("school") or "").strip()
            item   = (r.get("item_type") or "").strip()
            color  = (r.get("color") or "").strip()
            school_groups.setdefault(school, OrderedDict())
            school_groups[school].setdefault((item, color), []).append(r)

        def build_size_ranges_from_profile(profile):
            numeric_tables = []
            alpha_labels = []
            if profile is None:
                return numeric_tables, alpha_labels
            r1_start, r1_end, r2_start, r2_end, has_alpha = profile
            merged = merged_numeric_size_labels_from_profile(
                r1_start, r1_end, r2_start, r2_end,
            )
            if merged:
                numeric_tables.append(merged[:])
            if has_alpha:
                alpha_labels = ALPHA_SIZES[:]
            return numeric_tables, alpha_labels

        tables_html = []

        for sch, item_groups in school_groups.items():
            for (t, clr), items in item_groups.items():
                size_counts = defaultdict(int)
                for r in items:
                    sz = _normalize_size_label(r.get("size") or "")
                    size_counts[sz] += int(r.get("count") or 0)

                profile = self.db.get_size_profile(t, sch, clr)
                numeric_tables, alpha_labels = build_size_ranges_from_profile(profile)

                if not numeric_tables and not alpha_labels:
                    all_sizes = sorted({_normalize_size_label(r.get("size") or "") for r in items})
                    numeric = [s for s in all_sizes if s.isdigit()]
                    alpha   = [s for s in all_sizes if not s.isdigit()]
                    if numeric:
                        numeric_tables = [numeric]
                    if alpha:
                        alpha_labels = alpha

                def row_counts(labels):
                    out = []
                    for lbl in labels:
                        v = int(size_counts.get(lbl, 0))
                        out.append("" if v == 0 else str(v))
                    return out

                head = f"<div class='hdr'><span>النوع: {_html(t)}</span><span>المدرسة: {_html(sch)}</span><span>اللون: {_html(clr)}</span></div>"

                def build_table(chunk):
                    return f"<table class='grid'><tbody><tr>{''.join(f'<th>{_html(x)}</th>' for x in chunk)}</tr><tr>{''.join(f'<td class=num>{v}</td>' for v in row_counts(chunk))}</tr><tr>{''.join('<td>&nbsp;</td>' for _ in chunk)}</tr></tbody></table>"

                tables = []
                for numeric_labels in numeric_tables:
                    for i in range(0, len(numeric_labels), 15):
                        tables.append(build_table(numeric_labels[i:i + 15]))
                if alpha_labels:
                    tables.append("<div style='margin-top:6px;font-weight:600'>المقاسات بالحروف</div>" + build_table(alpha_labels))

                tables_html.append(f'<section class="sheet">{head}{"".join(tables)}</section>')

        html = f"""<!DOCTYPE html>
<html lang='ar' dir='rtl'><head><meta charset='utf-8'><title>جداول المقاسات</title>
<style>@page{{size:A4;margin:12mm}}*{{box-sizing:border-box}}body{{font-family:'Segoe UI',Tahoma,Arial,'Noto Sans Arabic',sans-serif;margin:0;direction:rtl}}
.sheet{{page-break-inside:avoid;margin-bottom:10mm}}.hdr{{display:flex;justify-content:space-between;font-weight:600;margin:6px 2px 8px}}
.grid{{border-collapse:collapse;width:100%;table-layout:fixed;margin-bottom:6px}}.grid th,.grid td{{border:1px solid #555;padding:6px 4px;text-align:center}}.grid th{{background:#eee}}.num{{font-variant-numeric:tabular-nums}}
</style></head><body>{''.join(tables_html)}<script>window.onload=function(){{try{{window.print();}}catch(e){{}}}};</script></body></html>"""

        import tempfile, os
        path = os.path.join(tempfile.gettempdir(), f"size_sheets_{datetime.now().strftime('%Y%m%d_%H%M%S')}.html")
        with open(path, "w", encoding="utf-8") as f:
            f.write(html)
        _print_html_auto(path, copies=1, parent=self)

    def _schedule_refresh(self, delay_ms: int = 250):
        if hasattr(self, "_inv_job") and self._inv_job:
            self.after_cancel(self._inv_job)
        self._inv_job = self.after(delay_ms, self._refresh)

    def _filters(self) -> Dict[str, Any]:
        active_multi = [k for k, v in self.multi.items() if v]
        if active_multi:
            fld = active_multi[0]
            f: Dict[str, Any] = {
                "item_type": None, "school": None, "color": None, "size": None,
            }
            f[fld] = self.multi[fld][:]
            return f

        return {
            "item_type": self.f_type.get() or None,
            "school": self.f_school.get() or None,
            "color": self.f_color.get() or None,
            "size": self.f_size.get() or None,
        }

    def _refresh(self):
        try:
            rows = self.db.current_inventory(self._filters())
        except Exception as ex:
            messagebox.showerror("فشل البحث", str(ex), parent=self)
            return

        self.table.delete(*self.table.get_children())
        total_qty = 0
        total_value = 0.0
        for r in rows:
            self.table.insert(
                "", tk.END,
                values=(r["id"], r["item_type"], r["school"], r["color"], r["size"],
                        f"{float(r['unit_price']):.2f}",
                        r["count"], f"{float(r['value']):.2f}")
            )
            total_qty += int(r["count"])
            total_value += float(r["value"])
        self.sum_qty.set(str(total_qty))
        self.sum_val.set(f"{total_value:.2f}")
        _apply_zebra_tags(self.table)

    def _clear_all(self):
        for w in (self.f_type, self.f_school, self.f_color, self.f_size):
            w.set("")
        for k in self.multi.keys():
            self.multi[k] = []
        for b in self._multi_btns.values():
            b.configure(text="اختيار متعدد...")
        for w in (self.f_type, self.f_school, self.f_color, self.f_size):
            w.refresh_values()
        self._enforce_single_active_field()
        self._refresh()

    def _export_excel(self):
        try:
            db_rows = self.db.current_inventory(self._filters())
        except Exception as ex:
            messagebox.showerror("فشل التصدير", str(ex), parent=self)
            return

        path = filedialog.asksaveasfilename(
            title="تصدير المخزون إلى إكسل",
            defaultextension=".xlsx",
            filetypes=[("Excel Workbook", "*.xlsx"), ("Excel 97-2003 XML", "*.xls"), ("All files", "*.*")],
            initialfile=f"inventory_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx",
            parent=self,
        )
        if not path:
            return
        try:
            self.db.export_inventory_excel(path, db_rows)
            ToastNotification.show(self.winfo_toplevel(), f"تم حفظ المخزون إلى: {path}", toast_type="success")
        except Exception as ex:
            messagebox.showerror("فشل التصدير", str(ex), parent=self)

    def _remove_selected_dialog(self):
        if not self.db.is_manager_feature_enabled("allow_inventory_delete"):
            messagebox.showwarning("مقيد", _feature_restricted_message("حذف المخزون من نقطة البيع غير مسموح به حالياً."), parent=self)
            return
        sel = self.table.selection()
        if not sel:
            messagebox.showwarning("لم يتم التحديد", "اختر صفاً واحداً أو أكثر من المخزون أولاً.", parent=self)
            return

        rows = [self.table.item(i, "values") for i in sel]
        ids = [int(r[0]) for r in rows]

        if len(rows) == 1:
            r = rows[0]
            label = f"{r[1]} / {r[2]} / {r[3]} / {r[4]}"
            available = int(r[7])
            title = "حذف من المخزون (يتطلب كلمة مرور)"
        else:
            label = f"عدد الصفوف المحددة: {len(rows)}"
            available = sum(int(r[7]) for r in rows)
            title = "حذف متعدد من المخزون (يتطلب كلمة مرور)"

        dlg = tk.Toplevel(self)
        dlg.title(title)
        dlg.transient(self)
        dlg.grab_set()

        frm = ttk.Frame(dlg, padding=12)
        frm.pack(fill=tk.BOTH, expand=True)

        ttk.Label(frm, text=label, font=("Segoe UI", 10, "bold")).grid(row=0, column=0, columnspan=2, sticky="w", pady=(0, 6))
        ttk.Label(frm, text=f"إجمالي المتوفر: {available}").grid(row=1, column=0, columnspan=2, sticky="w")
        ttk.Label(frm, text="الكمية المطلوب حذفها (فارغ = الكل):").grid(row=2, column=0, sticky="e", padx=4, pady=8)
        qty_var = tk.StringVar()
        ttk.Entry(frm, textvariable=qty_var, width=12).grid(row=2, column=1, sticky="w", padx=4, pady=8)
        ttk.Label(frm, text="كلمة المرور:").grid(row=3, column=0, sticky="e", padx=4, pady=4)
        pw_var = tk.StringVar()
        ttk.Entry(frm, textvariable=pw_var, show="*").grid(row=3, column=1, sticky="w", padx=4, pady=4)

        btns = ttk.Frame(frm)
        btns.grid(row=4, column=0, columnspan=2, sticky="e", pady=(8, 0))

        def on_ok():
            if not self.db.verify_admin_password(pw_var.get()):
                messagebox.showerror("مرفوض", "كلمة المرور غير صحيحة.", parent=dlg)
                return
            qtext = qty_var.get().strip()
            qty = None
            if qtext:
                try:
                    qty = int(qtext)
                    if qty <= 0:
                        raise ValueError
                except ValueError:
                    messagebox.showerror("كمية غير صالحة", "أدخل عدداً موجباً أو اتركه فارغاً للكل.", parent=dlg)
                    return
            try:
                total_removed = 0
                for stock_id in ids:
                    total_removed += self.db.remove_from_stock(stock_id, qty, note="Admin remove (bulk)")
                dlg.destroy()
                ToastNotification.show(self.winfo_toplevel(), f"تم حذف {total_removed} وحدة من {len(ids)} صف(وف)", toast_type="success")
                self._refresh()
            except Exception as ex:
                messagebox.showerror("خطأ", str(ex), parent=dlg)

        ttk.Button(btns, text="إلغاء", command=dlg.destroy).pack(side=tk.RIGHT)
        ttk.Button(btns, text="حذف", command=on_ok).pack(side=tk.RIGHT, padx=6)

    def _edit_price_dialog(self):
        if not self.db.is_manager_feature_enabled("allow_inventory_price_edit"):
            messagebox.showwarning("مقيد", _feature_restricted_message("تعديل الأسعار من نافذة المخزون غير مسموح به حالياً."), parent=self)
            return
        row = 0
        sel = self.table.selection()
        if not sel:
            messagebox.showwarning("لم يتم التحديد", "اختر صفاً من المخزون أولاً.", parent=self)
            return

        rows = [self.table.item(i, "values") for i in sel]
        first = rows[0]

        current_price = float(first[6])
        multi = len(rows) > 1

        dlg = tk.Toplevel(self)
        dlg.title("تعديل السعر")
        dlg.transient(self)
        dlg.grab_set()

        frm = ttk.Frame(dlg, padding=12)
        frm.pack(fill=tk.BOTH, expand=True)

        if multi:
            ttk.Label(frm, text=f"عدد الصفوف المحددة: {len(rows)}", font=("Segoe UI", 10, "bold"))\
                .grid(row=row, column=0, columnspan=2, sticky="w", pady=(0, 6))
            row += 1
        else:
            ttk.Label(frm, text=f"الصنف: {first[1]} / {first[2]} / {first[3]} / {first[4]}")\
                .grid(row=row, column=0, columnspan=2, sticky="w", pady=(0, 6))
            row += 1
            ttk.Label(frm, text=f"السعر الحالي: {current_price:.2f}")\
                .grid(row=row, column=0, columnspan=2, sticky="w", pady=(0, 10))
            row += 1

        ttk.Label(frm, text="السعر الجديد:").grid(row=row, column=0, sticky="e", padx=4, pady=6)
        price_var = tk.StringVar(value=f"{current_price:.2f}")
        ttk.Entry(frm, textvariable=price_var, width=16).grid(row=row, column=1, sticky="w", padx=4, pady=6)
        row += 1

        scope_var = tk.StringVar(value="row")

        if not multi:
            scope_box = ttk.LabelFrame(frm, text="النطاق")
            scope_box.grid(row=row, column=0, columnspan=2, sticky="ew", pady=(6, 10))
            ttk.Radiobutton(scope_box, text="الصف المحدد فقط", variable=scope_var, value="row").pack(anchor="w", padx=8, pady=4)
            ttk.Radiobutton(scope_box, text="كل الصفوف بنفس (النوع/المدرسة/المقاس)",
                            variable=scope_var, value="same_specs").pack(anchor="w", padx=8, pady=4)
            row += 1

        def on_ok():
            try:
                new_price = float(price_var.get())
                if new_price < 0:
                    raise ValueError
            except Exception:
                messagebox.showerror("قيمة غير صالحة", "أدخل سعراً رقمياً غير سالب.", parent=dlg)
                return

            try:
                updated_total = 0
                if multi:
                    for vals in rows:
                        updated_total += self.db.update_prices({"id": int(vals[0])}, new_price, note="Price update (multi-selection)")
                else:
                    if scope_var.get() == "row":
                        updated_total = self.db.update_prices({"id": int(first[0])}, new_price, note="Price update (single row)")
                    else:
                        updated_total = self.db.update_prices({
                            "item_type": first[1], "school": first[2], "size": first[4],
                        }, new_price, note="Price update (same type/school/size)")

                dlg.destroy()
                if updated_total == 0:
                    ToastNotification.show(self.winfo_toplevel(), "لم يتم العثور على صفوف مطابقة", toast_type="warning")
                else:
                    ToastNotification.show(self.winfo_toplevel(), f"تم تحديث السعر في {updated_total} صف(وف)", toast_type="success")
                self._refresh()
            except Exception as ex:
                messagebox.showerror("فشل التحديث", str(ex), parent=dlg)

        btns = ttk.Frame(frm)
        btns.grid(row=row, column=0, columnspan=2, sticky="e", pady=(12, 0))
        ttk.Button(btns, text="إلغاء", command=dlg.destroy).pack(side=tk.RIGHT)
        ttk.Button(btns, text="حفظ", command=on_ok).pack(side=tk.RIGHT, padx=6)
        dlg.update_idletasks()

    def _edit_specs_dialog(self):
        if not self.db.is_manager_feature_enabled("allow_inventory_specs_edit"):
            messagebox.showwarning("مقيد", _feature_restricted_message("تعديل المواصفات من نقطة البيع غير مسموح به حالياً."), parent=self)
            return
        sel = self.table.selection()
        ids: List[int] = []
        if not sel:
            messagebox.showwarning("حدد النطاق",
                "اختر صفوفاً من الجدول لتطبيق التعديل.",
                parent=self)
            return
        for iid in sel:
            vals = self.table.item(iid, "values")
            ids.append(int(vals[0]))
        scope_text = f"عدد الصفوف المحددة: {len(ids)}"

        dlg = tk.Toplevel(self)
        dlg.title("تعديل المواصفات")
        dlg.transient(self)
        dlg.grab_set()

        frm = ttk.Frame(dlg, padding=12)
        frm.pack(fill=tk.BOTH, expand=True)

        ttk.Label(frm, text=scope_text, font=("", 10, "bold")).grid(row=0, column=0, columnspan=2, sticky="w", pady=(0,8))
        ttk.Label(frm, text="اترك الحقل فارغاً إذا كنت لا تريد تغييره.").grid(row=1, column=0, columnspan=2, sticky="w", pady=(0,10))

        it_var = tk.StringVar()
        sc_var = tk.StringVar()
        cl_var = tk.StringVar()
        sz_var = tk.StringVar()

        ttk.Label(frm, text="النوع (جديد):").grid(row=2, column=0, sticky="e", padx=6, pady=4)
        ttk.Entry(frm, textvariable=it_var, width=28).grid(row=2, column=1, sticky="w", padx=6, pady=4)
        ttk.Label(frm, text="المدرسة (جديد):").grid(row=3, column=0, sticky="e", padx=6, pady=4)
        ttk.Entry(frm, textvariable=sc_var, width=28).grid(row=3, column=1, sticky="w", padx=6, pady=4)
        ttk.Label(frm, text="اللون (جديد):").grid(row=4, column=0, sticky="e", padx=6, pady=4)
        ttk.Entry(frm, textvariable=cl_var, width=28).grid(row=4, column=1, sticky="w", padx=6, pady=4)
        ttk.Label(frm, text="المقاس (جديد):").grid(row=5, column=0, sticky="e", padx=6, pady=4)
        ttk.Entry(frm, textvariable=sz_var, width=28).grid(row=5, column=1, sticky="w", padx=6, pady=4)

        btns = ttk.Frame(frm); btns.grid(row=6, column=0, columnspan=2, sticky="e", pady=(10,0))
        ttk.Button(btns, text="إلغاء", command=dlg.destroy).pack(side=tk.RIGHT)

        def on_ok():
            changes = {}
            if it_var.get().strip(): changes["item_type"] = it_var.get().strip()
            if sc_var.get().strip(): changes["school"]    = sc_var.get().strip()
            if cl_var.get().strip(): changes["color"]     = cl_var.get().strip()
            if sz_var.get().strip(): changes["size"]      = sz_var.get().strip()

            if not changes:
                messagebox.showwarning("لا تغييرات", "لم تُدخل أي قيم جديدة.", parent=dlg)
                return
            try:
                updated = self.db.update_specs_by_ids(ids, **changes)

                dlg.destroy()
                if updated == 0:
                    ToastNotification.show(self.winfo_toplevel(), "لم يتم العثور على صفوف مطابقة", toast_type="warning")
                else:
                    ToastNotification.show(self.winfo_toplevel(), f"تم تعديل المواصفات في {updated} صف(وف)", toast_type="success")
                self._refresh()
            except Exception as ex:
                messagebox.showerror("فشل التحديث", str(ex), parent=dlg)

        ttk.Button(btns, text="حفظ", command=on_ok).pack(side=tk.RIGHT, padx=6)


class MultiSelectDialog(tk.Toplevel):
    def __init__(self, master, title: str, values: List[str], preselected: Optional[List[str]] = None):
        super().__init__(master)
        self.protocol("WM_DELETE_WINDOW", self._on_cancel)
        self.title(title)
        self._alive = True
        self._after_id = None
        self._result = None

        self.transient(master)
        self.grab_set()
        try:
            self.focus_force()
        except Exception:
            pass

        self.geometry("380x460")
        self.resizable(False, False)

        self._values = list(dict.fromkeys(values or []))
        self._ok = False

        frm = ttk.Frame(self, padding=10)
        frm.pack(fill=tk.BOTH, expand=True)

        frow = ttk.Frame(frm)
        frow.pack(fill=tk.X, pady=(0,6))
        ttk.Label(frow, text="بحث:").pack(side=tk.LEFT)
        self.q = tk.StringVar()
        ent = ttk.Entry(frow, textvariable=self.q)
        ent.pack(side=tk.LEFT, fill=tk.X, expand=True, padx=(6,0))
        ent.bind("<KeyRelease>", lambda e: self._refill())

        self.listbox = tk.Listbox(frm, selectmode="extended", activestyle="none")
        self.listbox.pack(fill=tk.BOTH, expand=True)
        ysb = ttk.Scrollbar(frm, orient="vertical", command=self.listbox.yview)
        self.listbox.configure(yscrollcommand=ysb.set)
        _bind_mousewheel(self.listbox)
        ysb.pack(side=tk.RIGHT, fill=tk.Y)

        btns = ttk.Frame(frm)
        btns.pack(fill=tk.X, pady=(8,0))
        ttk.Button(btns, text="مسح التحديد", command=self._clear_sel).pack(side=tk.LEFT)
        ttk.Button(btns, text="إلغاء", command=self._on_cancel).pack(side=tk.RIGHT)
        ttk.Button(btns, text="موافق", command=self._on_ok).pack(side=tk.RIGHT, padx=6)

        self._refill()

        if preselected:
            lower_map = {v.lower(): i for i, v in enumerate(self._shown)}
            sel_ids = [lower_map.get(v.lower()) for v in preselected if lower_map.get(v.lower()) is not None]
            for i in sel_ids:
                self.listbox.selection_set(i)

        self._after_id = self.after(10, ent.focus_set)

    def _refill(self):
        if not self._alive:
            return
        if not self.winfo_exists():
            return
        if not self.listbox.winfo_exists():
            return
        q = (self.q.get() or "").strip().lower()
        if q:
            self._shown = [v for v in self._values if q in str(v).lower()]
        else:
            self._shown = self._values[:]
        self.listbox.delete(0, tk.END)
        for v in self._shown:
            self.listbox.insert(tk.END, v)

    def _clear_sel(self):
        self.listbox.selection_clear(0, tk.END)

    def _cleanup(self):
        self._alive = False
        try:
            if self._after_id:
                self.after_cancel(self._after_id)
        except Exception:
            pass

    def _on_ok(self):
        try:
            sel_idx = list(self.listbox.curselection())
            self._result = [self._shown[i] for i in sel_idx]
        except Exception:
            self._result = []
        self._ok = True
        self._cleanup()
        self.destroy()

    def _on_cancel(self):
        self._result = None
        self._cleanup()
        self.destroy()

    def run(self) -> Optional[List[str]]:
        self.wait_window()
        return self._result


# ------------------- Bills History Window -------------------

class BillsHistoryWindow(tk.Toplevel):
    def __init__(self, master, db: SqliteDatabase):
        super().__init__(master)
        self.db = db
        self.title("سجل الفواتير")
        self.geometry("1120x560")
        self._build()

    def _build(self):
        top = ttk.Frame(self)
        top.pack(fill=tk.X, padx=8, pady=8)
        ttk.Label(top, text="سجل الفواتير", font=("Segoe UI", 12, "bold")).pack(side=tk.LEFT)
        ttk.Button(top, text="تحديث", command=self._refresh).pack(side=tk.RIGHT)
        ttk.Button(top, text="فتح", command=self._print_selected).pack(side=tk.RIGHT, padx=8)
        ttk.Button(top, text="تصدير المحدد إلى إكسل", command=self._export_selected).pack(side=tk.RIGHT)
        ttk.Button(top, text="VOID مع سبب", command=self._void_selected).pack(side=tk.RIGHT, padx=8)
        ttk.Label(top, text="نوع الفاتورة:").pack(side=tk.RIGHT, padx=(16, 4))
        self._type_filter = tk.StringVar(value="الكل")
        _types_cb = ttk.Combobox(
            top, textvariable=self._type_filter,
            values=("الكل", "بيع", "مرتجع", "استبدال", "إلى المصنع"),
            width=14, state="readonly",
        )
        _types_cb.pack(side=tk.RIGHT)
        _types_cb.bind("<<ComboboxSelected>>", lambda _e: self._refresh())

        bills_wrap = ttk.Frame(self)
        bills_wrap.pack(fill=tk.BOTH, expand=False, padx=8, pady=(0, 6))
        self.bills_table = ttk.Treeview(
            bills_wrap,
            columns=("id", "created_at", "bill_type", "customer", "total", "status"),
            show="headings",
            height=10,
        )
        for col, txt, w in [
            ("id", "المعرّف", 70),
            ("created_at", "التاريخ", 160),
            ("bill_type", "النوع", 90),
            ("customer", "العميل", 200),
            ("total", "الإجمالي", 100),
            ("status", "الحالة", 100),
        ]:
            self.bills_table.heading(col, text=txt)
            self.bills_table.column(col, width=w, anchor="center")
        bills_ysb = ttk.Scrollbar(bills_wrap, orient="vertical", command=self.bills_table.yview)
        bills_xsb = ttk.Scrollbar(bills_wrap, orient="horizontal", command=self.bills_table.xview)
        self.bills_table.configure(yscrollcommand=bills_ysb.set, xscrollcommand=bills_xsb.set)
        self.bills_table.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        bills_ysb.pack(side=tk.RIGHT, fill=tk.Y)
        bills_xsb.pack(side=tk.BOTTOM, fill=tk.X)
        self.bills_table.bind("<<TreeviewSelect>>", lambda e: self._load_items())
        _bind_mousewheel(self.bills_table)

        # Context menu on bills table
        _add_context_menu(self.bills_table, self)

        ttk.Label(self, text="بنود الفاتورة").pack(anchor="w", padx=8)

        items_wrap = ttk.Frame(self)
        items_wrap.pack(fill=tk.BOTH, expand=True, padx=8, pady=(0, 8))
        self.items_table = ttk.Treeview(
            items_wrap,
            columns=("type", "school", "color", "size", "origin", "price", "qty", "total"),
            show="headings",
            height=12,
        )
        for col, txt, w in [
            ("type","النوع",140), ("school","المدرسة",160), ("color","اللون",80), ("size","المقاس",70),
            ("origin","المصدر",90), ("price","السعر",80),
            ("qty","الكمية",60), ("total","إجمالي",100),
        ]:
            self.items_table.heading(col, text=txt)
            self.items_table.column(col, width=w, anchor="center")
        items_ysb = ttk.Scrollbar(items_wrap, orient="vertical", command=self.items_table.yview)
        items_xsb = ttk.Scrollbar(items_wrap, orient="horizontal", command=self.items_table.xview)
        self.items_table.configure(yscrollcommand=items_ysb.set, xscrollcommand=items_xsb.set)
        self.items_table.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        items_ysb.pack(side=tk.RIGHT, fill=tk.Y)
        items_xsb.pack(side=tk.BOTTOM, fill=tk.X)
        _bind_mousewheel(self.items_table)

        # Context menu on items table
        _add_context_menu(self.items_table, self)

        # Tooltip preview on bills table
        def _bill_tooltip():
            sel = self.bills_table.selection()
            if not sel:
                return ""
            bill_id = int(sel[0])
            try:
                items = self.db.list_bill_items(bill_id)
                if not items:
                    return "لا توجد بنود"
                lines = []
                for ln in items[:6]:
                    lines.append(f"{ln['item_type']} - {ln['school']} ({ln['color']}/{ln['size']}) x{ln['qty']}")
                if len(items) > 6:
                    lines.append(f"... و {len(items) - 6} أصناف أخرى")
                return "\n".join(lines)
            except Exception:
                return ""
        ToolTip(self.bills_table, text_func=_bill_tooltip, delay=600)

        self._refresh()

    @staticmethod
    def _bill_type_ar(bt: Optional[str]) -> str:
        u = str(bt or "SALE").upper()
        return {
            "SALE": "بيع",
            "RETURN": "مرتجع",
            "EXCHANGE": "استبدال",
            "WAREHOUSE_RETURN": "إلى المصنع",
        }.get(u, u)

    def _selected_bill_type_code(self) -> str:
        m = {
            "الكل": "",
            "بيع": "SALE",
            "مرتجع": "RETURN",
            "استبدال": "EXCHANGE",
            "إلى المصنع": "WAREHOUSE_RETURN",
        }
        return m.get(self._type_filter.get() or "الكل", "")

    def _refresh(self):
        self.bills_table.delete(*self.bills_table.get_children())
        want = self._selected_bill_type_code()
        for b in self.db.list_bills():
            bt = str(b.get("bill_type") or "SALE").upper()
            if want and bt != want:
                continue
            status_text = "ملغاة" if str(b.get("status") or "").upper() == "VOID" else "مؤكدة"
            self.bills_table.insert(
                "", tk.END, iid=str(b["id"]),
                values=(
                    b["id"],
                    b["created_at"],
                    self._bill_type_ar(b.get("bill_type")),
                    b.get("customer") or "",
                    f"{float(b['total']):.2f}",
                    status_text,
                ),
            )
        self.items_table.delete(*self.items_table.get_children())
        _apply_zebra_tags(self.bills_table)

    def _get_selected_bill_id(self) -> Optional[int]:
        sel = self.bills_table.selection()
        if not sel:
            return None
        return int(sel[0])

    def _load_items(self):
        bill_id = self._get_selected_bill_id()
        self.items_table.delete(*self.items_table.get_children())
        if bill_id is None:
            return
        items = self.db.list_bill_items(bill_id)
        for ln in items:
            origin_txt = "من المخزون" if ln.get("origin") == "STOCK" else ("من المصنع" if ln.get("origin") == "FACTORY" else "")
            self.items_table.insert(
                "", tk.END,
                values=(ln["item_type"], ln["school"], ln["color"], ln["size"], origin_txt,
                        f"{float(ln['unit_price']):.2f}",
                        ln["qty"], f"{float(ln['line_total']):.2f}")
            )
        _apply_zebra_tags(self.items_table)

    def _export_selected(self):
        bill_id = self._get_selected_bill_id()
        if bill_id is None:
            messagebox.showwarning("لم يتم التحديد", "اختر فاتورة أولاً.")
            return
        items = self.db.list_bill_items(bill_id)
        if not items:
            messagebox.showwarning("فارغ", "لا تحتوي هذه الفاتورة على بنود.")
            return
        path = filedialog.asksaveasfilename(
            title="تصدير الفاتورة إلى إكسل",
            defaultextension=".xlsx",
            filetypes=[("Excel Workbook", "*.xlsx"), ("Excel 97-2003 XML", "*.xls"), ("All files", "*.*")],
            initialfile=f"bill_{bill_id}.xlsx",
        )
        if not path:
            return
        headers = ["type", "school", "color", "size", "origin", "unit_price", "qty", "line_total"]
        def _origin_txt(o: Optional[str]) -> str:
            return "من المخزون" if o == "STOCK" else ("من المصنع" if o == "FACTORY" else "")
        rows = [
            [
                ln["item_type"], ln["school"], ln["color"], ln["size"],
                _origin_txt(ln.get("origin")),
                float(ln["unit_price"]), int(ln["qty"]), float(ln["line_total"]),
            ]
            for ln in items
        ]
        try:
            export_to_excel(path, headers, rows)
            ToastNotification.show(self.winfo_toplevel(), f"تم تصدير الفاتورة إلى: {path}", toast_type="success")
        except Exception as ex:
            messagebox.showerror("فشل التصدير", str(ex))

    def _print_selected(self):
        bill_id = self._get_selected_bill_id()
        if bill_id is None:
            messagebox.showwarning("لم يتم التحديد", "اختر فاتورة أولاً.")
            return
        try:
            bill = next(b for b in self.db.list_bills() if int(b["id"]) == bill_id)
            items = self.db.list_bill_items(bill_id)
        except StopIteration:
            messagebox.showerror("فشل الطباعة", "لم يتم العثور على الفاتورة.")
            return
        tmp_dir = tempfile.gettempdir()
        path = os.path.join(tmp_dir, f"bill_{bill_id}.html")
        save_bill_as_html(path, bill, items)
        _print_html_auto(path, copies=1, parent=self)

    def _void_selected(self):
        bill_id = self._get_selected_bill_id()
        if bill_id is None:
            messagebox.showwarning("لم يتم التحديد", "اختر فاتورة أولاً.", parent=self)
            return
        pw = simpledialog.askstring("كلمة مرور المدير", "أدخل كلمة مرور المدير:", show="*", parent=self)
        if not pw:
            return
        if not self.db.verify_admin_password(pw):
            messagebox.showerror("مرفوض", "كلمة المرور غير صحيحة.", parent=self)
            return
        reason = simpledialog.askstring("سبب الإلغاء", "أدخل سبب الـ VOID:", parent=self)
        if reason is None:
            return
        try:
            self.db.void_bill(bill_id, reason)
            ToastNotification.show(self.winfo_toplevel(), f"تم إلغاء الفاتورة #{bill_id} وتوثيق السبب", toast_type="success")
            self._refresh()
        except Exception as ex:
            messagebox.showerror("فشل الإلغاء", str(ex), parent=self)


# ------------------- Movements Window -------------------

def _summarize_movement_rows(rows: List[Dict[str, Any]]) -> Dict[str, Any]:
    """Aggregate counts/amounts for the movements log footer (filtered rows only)."""
    n = 0
    qty_in = 0
    qty_out = 0
    qty_adj_out = 0
    qty_return_in = 0
    qty_reserve = 0
    val_out = 0.0
    deliver_cash = 0.0
    for m in rows:
        n += 1
        d = str(m.get("direction") or "").strip().upper()
        q = int(m.get("qty") or 0)
        try:
            up = float(m.get("unit_price") or 0)
        except (TypeError, ValueError):
            up = 0.0
        if d == "IN":
            qty_in += q
        elif d in ("OUT", "OUT_FACTORY"):
            qty_out += q
            val_out += float(q) * up
        elif d == "ADJUST_OUT":
            qty_adj_out += q
            val_out += float(q) * up
        elif d == "RETURN_IN":
            qty_return_in += q
        elif d == "RESERVE":
            qty_reserve += q
        elif d == "DELIVER_PAY":
            deliver_cash += up
    income_moves = val_out + deliver_cash
    return {
        "n": n,
        "qty_in": qty_in,
        "qty_out": qty_out,
        "qty_adj_out": qty_adj_out,
        "qty_return_in": qty_return_in,
        "qty_reserve": qty_reserve,
        "val_out": val_out,
        "deliver_cash": deliver_cash,
        "income_moves": income_moves,
    }


class MovementsWindow(tk.Toplevel):
    def __init__(self, master, db: SqliteDatabase):
        super().__init__(master)
        self.db = db
        self.title("سجل الحركات")
        self.geometry("1180x620")
        self._build()

    def _build(self):
        top = ttk.LabelFrame(self, text="تصنيف")
        top.pack(fill=tk.X, padx=8, pady=8)

        self.customer = LabeledStaticCombo(top, "العميل", values=[""])
        self.customer.grid(row=0, column=0, sticky="ew", padx=6, pady=4)
        self.customer.cb.bind("<Button-1>", lambda e: self._reload_customers(), add="+")
        self.customer.cb.bind("<FocusIn>", lambda e: self._reload_customers(), add="+")

        self.ftype  = LabeledCombobox(top, "النوع",   self.db, "item_type");  self.ftype.grid(row=0, column=2, padx=6, pady=4, sticky="ew")
        self.fsch   = LabeledCombobox(top, "المدرسة", self.db, "school");     self.fsch.grid(row=0, column=3, padx=6, pady=4, sticky="ew")
        self.fclr   = LabeledCombobox(top, "اللون",   self.db, "color");      self.fclr.grid(row=0, column=4, padx=6, pady=4, sticky="ew")
        self.fsiz   = LabeledCombobox(top, "المقاس",  self.db, "size");       self.fsiz.grid(row=0, column=5, padx=6, pady=4, sticky="ew")

        self.df = DateField(top, "من (YYYY-MM-DD)"); self.df.grid(row=1, column=0, padx=6, pady=4, sticky="w")
        self.dt = DateField(top, "إلى");             self.dt.grid(row=1, column=1, padx=6, pady=4, sticky="w")

        ttk.Label(top, text="نص").grid(row=1, column=4, sticky="e", padx=4, pady=4)
        self.txt = ttk.Entry(top); self.txt.grid(row=1, column=5, sticky="ew", padx=6, pady=4)
        top.columnconfigure(5, weight=1)

        btns = ttk.Frame(top)
        btns.grid(row=0, column=6, rowspan=2, sticky="e", padx=6, pady=4)
        ttk.Button(btns, text="تحديث", command=self._refresh).pack(side=tk.LEFT)
        ttk.Button(btns, text="مسح", command=self._clear).pack(side=tk.LEFT, padx=6)
        ttk.Button(btns, text="تصدير إلى إكسل", command=self._export).pack(side=tk.LEFT)

        table_wrap = ttk.Frame(self)
        table_wrap.pack(fill=tk.BOTH, expand=True, padx=8, pady=(0,8))
        self.table = ttk.Treeview(
            table_wrap,
            columns=("id","ts","direction","type","school","color","size","qty","note","bill_id","stock_id"),
            show="headings",
            height=14,
        )
        for col, txt, w in [
            ("id","المعرّف",70), ("ts","الوقت",170), ("direction","الاتجاه",100), ("type","النوع",150),
            ("school","المدرسة",170), ("color","اللون",100), ("size","المقاس",80), ("qty","الكمية",70),
            ("note","ملاحظة",200), ("bill_id","الفاتورة",70), ("stock_id","المخزون",70),
        ]:
            self.table.heading(col, text=txt)
            self.table.column(col, width=w, anchor="center")
        ysb = ttk.Scrollbar(table_wrap, orient="vertical", command=self.table.yview)
        xsb = ttk.Scrollbar(table_wrap, orient="horizontal", command=self.table.xview)
        self.table.configure(yscrollcommand=ysb.set, xscrollcommand=xsb.set)
        self.table.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        ysb.pack(side=tk.RIGHT, fill=tk.Y)
        xsb.pack(side=tk.BOTTOM, fill=tk.X)
        _bind_mousewheel(self.table)

        # Context menu
        _add_context_menu(self.table, self)

        sum_fr = ttk.LabelFrame(self, text="ملخص النتائج المعروضة (حسب الفلاتر الحالية)")
        sum_fr.pack(fill=tk.X, padx=8, pady=(0, 8))
        self._summary_var = tk.StringVar(value="")
        ttk.Label(
            sum_fr, textvariable=self._summary_var, justify="right", wraplength=1120,
            font=("Segoe UI", 9),
        ).pack(anchor="e", padx=10, pady=8)

        self._refresh()
        self._reload_customers()

    def _reload_customers(self):
        vals = [""] + self.db.list_customers()
        self.customer.cb["values"] = vals

    def _filters(self) -> Dict[str, Any]:
        return {
            "customer":  self.customer.get() or None,
            "item_type": self.ftype.get(),
            "school":    self.fsch.get(),
            "color":     self.fclr.get(),
            "size":      self.fsiz.get(),
            "date_from": self.df.get() or None,
            "date_to":   self.dt.get() or None,
            "text":      (self.txt.get().strip() or None),
        }

    def _refresh(self):
        try:
            rows = self.db.list_movements(self._filters())
        except Exception as ex:
            messagebox.showerror("فشل البحث", str(ex), parent=self)
            self._summary_var.set("")
            return

        self.table.delete(*self.table.get_children())
        for m in rows:
            self.table.insert(
                "", tk.END,
                values=(m.get("id"), m.get("ts"), m.get("direction"),
                        m.get("item_type",""), m.get("school",""), m.get("color",""), m.get("size",""),
                        m.get("qty"), m.get("note",""),
                        m.get("bill_id") if m.get("bill_id") is not None else "",
                        m.get("stock_id") if m.get("stock_id") is not None else "")
            )
        _apply_zebra_tags(self.table)
        s = _summarize_movement_rows(rows)
        self._summary_var.set(
            f"عدد الحركات: {s['n']}  |  "
            f"إجمالي كمية الوارد: {s['qty_in']}  |  "
            f"إجمالي كمية المنصرف (بيع/مصنع): {s['qty_out']}  |  "
            f"تسويات سلبية (كمية): {s['qty_adj_out']}  |  "
            f"مرتجعات واردة (كمية): {s['qty_return_in']}  |  "
            f"حجوزات (كمية): {s['qty_reserve']}\n"
            f"قيمة المنصرف (كمية×سعر الحركة): {s['val_out']:.2f}  |  "
            f"تحصيل تسليم حجوزات: {s['deliver_cash']:.2f}  |  "
            f"إجمالي الدخل (منصرف + تحصيل تسليم): {s['income_moves']:.2f}"
        )

    def _clear(self):
        self.customer.set("")
        for w in (self.ftype, self.fsch, self.fclr, self.fsiz):
            w.set("")
        self.df.set("")
        self.dt.set("")
        self.txt.delete(0, tk.END)
        self._refresh()

    def _export(self):
        rows = self.db.list_movements(self._filters())
        if not rows:
            messagebox.showwarning("فارغ", "لا توجد صفوف للتصدير.", parent=self)
            return
        path = filedialog.asksaveasfilename(
            title="تصدير الحركات إلى إكسل",
            defaultextension=".xlsx",
            filetypes=[("Excel Workbook", "*.xlsx"), ("Excel 97-2003 XML", "*.xls"), ("All files", "*.*")],
            initialfile=f"movements_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx",
            parent=self,
        )
        if not path:
            return
        headers = ["id","ts","direction","item_type","school","color","size","unit_price","qty","note","bill_id","stock_id"]
        table = [[
            m.get("id"), m.get("ts"), m.get("direction"),
            m.get("item_type",""), m.get("school",""), m.get("color",""), m.get("size",""),
            m.get("unit_price",""),
            m.get("qty"), m.get("note",""), m.get("bill_id"), m.get("stock_id"),
        ] for m in rows]
        try:
            export_to_excel(path, headers, table)
            ToastNotification.show(self.winfo_toplevel(), f"تم حفظ الحركات إلى: {path}", toast_type="success")
        except Exception as ex:
            messagebox.showerror("فشل التصدير", str(ex), parent=self)


# ------------------- Admin Window -------------------

class AdminWindow(tk.Toplevel):
    """Admin configuration window."""

    def __init__(self, master, db: SqliteDatabase):
        super().__init__(master)
        self.db = db
        self.title("الإعدادات والإدارة")
        self.geometry("520x520")
        self.resizable(True, True)
        self.transient(master)
        self.grab_set()
        self._build()

    def _build(self):
        nb = ttk.Notebook(self)
        nb.pack(fill=tk.BOTH, expand=True, padx=6, pady=6)

        # ---- Tab 1: Password ----
        pw_tab = ttk.Frame(nb, padding=12)
        nb.add(pw_tab, text="كلمة المرور")

        ttk.Label(pw_tab, text="تغيير كلمة المرور", font=("Segoe UI", 11, "bold")).grid(
            row=0, column=0, columnspan=2, sticky="w", pady=(0, 12))

        ttk.Label(pw_tab, text="كلمة المرور الحالية:").grid(row=1, column=0, sticky="e", padx=6, pady=4)
        self._pw_cur = ttk.Entry(pw_tab, show="*")
        self._pw_cur.grid(row=1, column=1, sticky="ew", padx=6, pady=4)

        ttk.Label(pw_tab, text="كلمة المرور الجديدة:").grid(row=2, column=0, sticky="e", padx=6, pady=4)
        self._pw_new = ttk.Entry(pw_tab, show="*")
        self._pw_new.grid(row=2, column=1, sticky="ew", padx=6, pady=4)

        ttk.Label(pw_tab, text="تأكيد كلمة المرور:").grid(row=3, column=0, sticky="e", padx=6, pady=4)
        self._pw_confirm = ttk.Entry(pw_tab, show="*")
        self._pw_confirm.grid(row=3, column=1, sticky="ew", padx=6, pady=4)

        pw_tab.columnconfigure(1, weight=1)

        pw_btns = ttk.Frame(pw_tab)
        pw_btns.grid(row=4, column=0, columnspan=2, sticky="e", pady=(12, 0))
        ttk.Button(pw_btns, text="حفظ كلمة المرور", command=self._change_password).pack(side=tk.RIGHT)

        # ---- Tab 2: Feature Permissions ----
        perms_tab = ttk.Frame(nb, padding=12)
        nb.add(perms_tab, text="صلاحيات المدير")

        ttk.Label(perms_tab, text="الميزات المسموحة في نقطة البيع", font=("Segoe UI", 11, "bold")).grid(
            row=0, column=0, columnspan=2, sticky="w", pady=(0, 12))

        ttk.Label(
            perms_tab,
            text="أدخل كلمة مرور المدير ثم اختر ما تريد السماح به أو تقييده.",
            justify="right",
        ).grid(row=1, column=0, columnspan=2, sticky="w", pady=(0, 8))

        self._feature_vars: Dict[str, tk.BooleanVar] = {}
        feature_rows = [
            ("allow_inventory_window", "السماح بفتح نافذة المخزون"),
            ("allow_manual_incoming", "السماح بفاتورة الوارد اليدوية"),
            ("allow_bulk_price", "السماح بتعديل الأسعار الجماعي"),
            ("allow_excel_import", "السماح بالاستيراد من إكسل"),
            ("allow_manual_adjustment", "السماح بالتعديل اليدوي للكميات"),
            ("allow_reset_counts", "السماح بإعادة التعيين"),
            ("allow_inventory_delete", "السماح بحذف المخزون يدوياً"),
            ("allow_inventory_price_edit", "السماح بتعديل الأسعار من نافذة المخزون"),
            ("allow_inventory_specs_edit", "السماح بتعديل المواصفات من نافذة المخزون"),
            ("allow_size_profile_edit", "السماح بتعديل نطاقات المقاسات"),
        ]
        for row_idx, (key, label) in enumerate(feature_rows, start=2):
            var = tk.BooleanVar(value=self.db.is_manager_feature_enabled(key))
            self._feature_vars[key] = var
            ttk.Checkbutton(perms_tab, text=label, variable=var).grid(
                row=row_idx, column=0, columnspan=2, sticky="w", pady=2
            )

        ttk.Label(perms_tab, text="كلمة مرور المدير:").grid(
            row=2 + len(feature_rows), column=0, sticky="e", padx=6, pady=(10, 4)
        )
        self._perm_pw = ttk.Entry(perms_tab, show="*", width=20)
        self._perm_pw.grid(row=2 + len(feature_rows), column=1, sticky="w", padx=6, pady=(10, 4))
        perms_tab.columnconfigure(1, weight=1)

        perms_btns = ttk.Frame(perms_tab)
        perms_btns.grid(row=3 + len(feature_rows), column=0, columnspan=2, sticky="e", pady=(12, 0))
        ttk.Button(perms_btns, text="تحميل الحالي", command=self._reload_feature_permissions).pack(side=tk.RIGHT)
        ttk.Button(perms_btns, text="حفظ الصلاحيات", command=self._save_feature_permissions).pack(side=tk.RIGHT, padx=6)

        # ---- Tab 3: Import ----
        imp_tab = ttk.Frame(nb, padding=12)
        nb.add(imp_tab, text="استيراد")

        ttk.Label(imp_tab, text="استيراد من ملف إكسل", font=("Segoe UI", 11, "bold")).grid(
            row=0, column=0, columnspan=2, sticky="w", pady=(0, 12))

        ttk.Label(imp_tab, text="الملف:").grid(row=1, column=0, sticky="e", padx=6, pady=4)
        self._imp_path = tk.StringVar()
        ttk.Entry(imp_tab, textvariable=self._imp_path, width=40).grid(row=1, column=1, sticky="ew", padx=6, pady=4)
        ttk.Button(imp_tab, text="استعراض...", command=self._browse_import).grid(row=2, column=1, sticky="w", padx=6, pady=2)

        imp_tab.columnconfigure(1, weight=1)

        imp_btns = ttk.Frame(imp_tab)
        imp_btns.grid(row=3, column=0, columnspan=2, sticky="e", pady=(12, 0))
        ttk.Button(imp_btns, text="استيراد", command=self._do_import).pack(side=tk.RIGHT)

        # ---- Tab 4: Adjustment ----
        adj_tab = ttk.Frame(nb, padding=12)
        nb.add(adj_tab, text="تعديل يدوي")

        ttk.Label(adj_tab, text="تعديل الكميات يدوياً", font=("Segoe UI", 11, "bold")).grid(
            row=0, column=0, columnspan=2, sticky="w", pady=(0, 12))

        fields = [
            ("النوع:", "_adj_type"),
            ("المدرسة:", "_adj_school"),
            ("اللون:", "_adj_color"),
            ("المقاس:", "_adj_size"),
        ]
        for r, (lbl, attr) in enumerate(fields, 1):
            ttk.Label(adj_tab, text=lbl).grid(row=r, column=0, sticky="e", padx=6, pady=4)
            var = ttk.Entry(adj_tab, width=24)
            var.grid(row=r, column=1, sticky="ew", padx=6, pady=4)
            setattr(self, attr, var)

        ttk.Label(adj_tab, text="الكمية (موجب=إضافة، سالب=حذف):").grid(row=5, column=0, sticky="e", padx=6, pady=4)
        self._adj_qty = ttk.Entry(adj_tab, width=12)
        self._adj_qty.grid(row=5, column=1, sticky="w", padx=6, pady=4)

        ttk.Label(adj_tab, text="السعر:").grid(row=6, column=0, sticky="e", padx=6, pady=4)
        self._adj_price = ttk.Entry(adj_tab, width=12)
        self._adj_price.grid(row=6, column=1, sticky="w", padx=6, pady=4)

        ttk.Label(adj_tab, text="ملاحظة:").grid(row=7, column=0, sticky="e", padx=6, pady=4)
        self._adj_note = ttk.Entry(adj_tab, width=28)
        self._adj_note.grid(row=7, column=1, sticky="ew", padx=6, pady=4)

        ttk.Label(adj_tab, text="كلمة المرور:").grid(row=8, column=0, sticky="e", padx=6, pady=4)
        self._adj_pw = ttk.Entry(adj_tab, show="*", width=16)
        self._adj_pw.grid(row=8, column=1, sticky="w", padx=6, pady=4)

        adj_tab.columnconfigure(1, weight=1)

        adj_btns = ttk.Frame(adj_tab)
        adj_btns.grid(row=9, column=0, columnspan=2, sticky="e", pady=(12, 0))
        ttk.Button(adj_btns, text="تطبيق التعديل", command=self._do_adjustment).pack(side=tk.RIGHT)

        # ---- Tab 5: Reset Counts ----
        reset_tab = ttk.Frame(nb, padding=12)
        nb.add(reset_tab, text="إعادة تعيين")

        ttk.Label(reset_tab, text="إعادة تعيين العدادات", font=("Segoe UI", 11, "bold")).grid(
            row=0, column=0, columnspan=2, sticky="w", pady=(0, 12))

        ttk.Label(reset_tab, text=(
            "تحذير: هذا الإجراء سيحذف جميع سجلات الحركات والفواتير والحجوزات\n"
            "ولا يمكن التراجع عنه. يُستخدم عادةً في بداية موسم جديد."
        ), foreground="red", justify="right").grid(row=1, column=0, columnspan=2, sticky="w", pady=(0, 12))

        ttk.Label(reset_tab, text="كلمة المرور للتأكيد:").grid(row=2, column=0, sticky="e", padx=6, pady=4)
        self._reset_pw = ttk.Entry(reset_tab, show="*", width=20)
        self._reset_pw.grid(row=2, column=1, sticky="w", padx=6, pady=4)

        reset_tab.columnconfigure(1, weight=1)

        reset_btns = ttk.Frame(reset_tab)
        reset_btns.grid(row=3, column=0, columnspan=2, sticky="e", pady=(12, 0))
        ttk.Button(reset_btns, text="إعادة التعيين", command=self._reset_counts).pack(side=tk.RIGHT)

        # ---- Close button ----
        ttk.Button(self, text="إغلاق", command=self.destroy).pack(side=tk.BOTTOM, pady=8)

    # ---- Actions ----

    def _change_password(self):
        cur = self._pw_cur.get()
        new = self._pw_new.get().strip()
        confirm = self._pw_confirm.get().strip()
        if not self.db.verify_admin_password(cur):
            messagebox.showerror("مرفوض", "كلمة المرور الحالية غير صحيحة.", parent=self)
            return
        if not new:
            messagebox.showwarning("فارغة", "كلمة المرور الجديدة لا يمكن أن تكون فارغة.", parent=self)
            return
        if new != confirm:
            messagebox.showwarning("غير متطابقة", "كلمتا المرور غير متطابقتين.", parent=self)
            return
        self.db.set_admin_password(new)
        ToastNotification.show(self.winfo_toplevel(), "تم تغيير كلمة المرور بنجاح", toast_type="success")
        self._pw_cur.delete(0, tk.END)
        self._pw_new.delete(0, tk.END)
        self._pw_confirm.delete(0, tk.END)

    def _browse_import(self):
        path = filedialog.askopenfilename(
            title="اختر ملف إكسل",
            filetypes=[("Excel files", "*.xlsx *.xls"), ("All files", "*.*")],
            parent=self,
        )
        if path:
            self._imp_path.set(path)

    def _reload_feature_permissions(self):
        for key, var in self._feature_vars.items():
            var.set(self.db.is_manager_feature_enabled(key))

    def _save_feature_permissions(self):
        pw = self._perm_pw.get()
        if not self.db.verify_admin_password(pw):
            messagebox.showerror("مرفوض", "كلمة المرور غير صحيحة.", parent=self)
            return
        for key, var in self._feature_vars.items():
            self.db.set_manager_feature_enabled(key, bool(var.get()))
        ToastNotification.show(self.winfo_toplevel(), "تم حفظ صلاحيات المدير بنجاح", toast_type="success")
        self._perm_pw.delete(0, tk.END)

    def _do_import(self):
        if not self.db.is_manager_feature_enabled("allow_excel_import"):
            messagebox.showwarning("مقيد", _feature_restricted_message("الاستيراد إلى نقطة البيع غير مسموح به حالياً."), parent=self)
            return
        path = self._imp_path.get().strip()

        if not path:
            messagebox.showwarning("مسار ناقص", "اختر ملف إكسل أولاً.", parent=self)
            return

        try:
            count = self.db.import_from_excel(path)
            ToastNotification.show(self.winfo_toplevel(), f"تم استيراد {count} صف بنجاح", toast_type="success")
        except Exception as ex:
            messagebox.showerror("فشل الاستيراد", str(ex), parent=self)

    def _do_adjustment(self):
        if not self.db.is_manager_feature_enabled("allow_manual_adjustment"):
            messagebox.showwarning("مقيد", _feature_restricted_message("التعديل اليدوي على مخزون نقطة البيع غير مسموح به حالياً."), parent=self)
            return
        pw = self._adj_pw.get()
        if not self.db.verify_admin_password(pw):
            messagebox.showerror("مرفوض", "كلمة المرور غير صحيحة.", parent=self)
            return

        item_type = self._adj_type.get().strip()
        school    = self._adj_school.get().strip()
        color     = self._adj_color.get().strip()
        size      = self._adj_size.get().strip()
        qty_txt   = self._adj_qty.get().strip()
        price_txt = self._adj_price.get().strip()
        note      = self._adj_note.get().strip() or "Manual adjustment"

        if not (item_type and school and color and size):
            messagebox.showwarning("بيانات ناقصة", "أدخل النوع والمدرسة واللون والمقاس.", parent=self)
            return
        try:
            qty = int(qty_txt)
        except ValueError:
            messagebox.showerror("كمية غير صالحة", "أدخل كمية صحيحة.", parent=self)
            return
        try:
            price = float(price_txt) if price_txt else 0.0
            if price < 0:
                raise ValueError
        except ValueError:
            messagebox.showerror("سعر غير صالح", "أدخل سعراً صحيحاً.", parent=self)
            return

        try:
            self.db.manual_adjustment(
                item_type=item_type, school=school, color=color, size=size,
                qty=qty, unit_price=price, note=note
            )
            ToastNotification.show(self.winfo_toplevel(), f"تم تطبيق التعديل ({qty:+d}) على المخزون", toast_type="success")
        except Exception as ex:
            messagebox.showerror("فشل التعديل", str(ex), parent=self)

    def _reset_counts(self):
        if not self.db.is_manager_feature_enabled("allow_reset_counts"):
            messagebox.showwarning("مقيد", _feature_restricted_message("إعادة التعيين من نقطة البيع غير مسموح بها حالياً."), parent=self)
            return
        pw = self._reset_pw.get()
        if not self.db.verify_admin_password(pw):
            messagebox.showerror("مرفوض", "كلمة المرور غير صحيحة.", parent=self)
            return
        confirm = messagebox.askyesno(
            "تأكيد إعادة التعيين",
            "هل أنت متأكد من حذف جميع سجلات الحركات والفواتير والحجوزات؟\n"
            "هذا الإجراء لا يمكن التراجع عنه.",
            parent=self,
        )
        if not confirm:
            return
        try:
            self.db.reset_movement_counts()
            ToastNotification.show(self.winfo_toplevel(), "تم إعادة تعيين جميع العدادات بنجاح", toast_type="success")
            self._reset_pw.delete(0, tk.END)
        except Exception as ex:
            messagebox.showerror("فشل", str(ex), parent=self)


# ------------------- Shift Summary Dialog -------------------

class ShiftSummaryDialog(tk.Toplevel):
    """Show shift summary and confirm closure."""
    def __init__(self, master, db: SqliteDatabase, shift_id: int, on_closed=None):
        super().__init__(master)
        self.db = db
        self.shift_id = shift_id
        self._on_closed = on_closed
        self.title("ملخص الوردية")
        self.geometry("650x700")
        self.minsize(550, 500)
        self.resizable(True, True)
        self.transient(master)
        self.grab_set()
        self._build()

    def _build(self):
        summary = self.db.get_shift_summary(self.shift_id)

        # Scrollable container so content is never clipped
        outer = ttk.Frame(self)
        outer.pack(fill=tk.BOTH, expand=True)

        canvas = tk.Canvas(outer, highlightthickness=0)
        vsb = ttk.Scrollbar(outer, orient="vertical", command=canvas.yview)
        canvas.configure(yscrollcommand=vsb.set)
        canvas.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        vsb.pack(side=tk.RIGHT, fill=tk.Y)
        _bind_mousewheel(canvas)

        main = ttk.Frame(canvas, padding=10)
        main_win = canvas.create_window((0, 0), window=main, anchor="nw")

        def _on_cfg(_e=None):
            canvas.configure(scrollregion=canvas.bbox("all"))
            canvas.itemconfigure(main_win, width=canvas.winfo_width())
        main.bind("<Configure>", _on_cfg)
        canvas.bind("<Configure>", _on_cfg)

        ttk.Label(main, text="ملخص الوردية", font=("Segoe UI", 14, "bold")).pack(anchor="w", pady=(0, 8))

        info = ttk.LabelFrame(main, text="معلومات الوردية")
        info.pack(fill=tk.X, pady=4)
        started = summary["started_at"][:16].replace("T", " ")
        ttk.Label(info, text=f"بداية: {started}").pack(anchor="w", padx=8, pady=2)
        ttk.Label(info, text="نهاية: الآن").pack(anchor="w", padx=8, pady=2)

        inf = ttk.LabelFrame(main, text="الوارد (الأصناف المستلمة)")
        inf.pack(fill=tk.X, pady=4)
        ttk.Label(inf, text=f"عدد عمليات الإضافة: {summary['inflow_count']}").pack(anchor="w", padx=8, pady=2)
        ttk.Label(inf, text=f"إجمالي القطع الواردة: {summary['inflow_total_qty']}").pack(anchor="w", padx=8, pady=2)

        if summary["inflow_items"]:
            tree = ttk.Treeview(inf, columns=("type","school","color","size","qty"), show="headings", height=min(5, len(summary["inflow_items"])))
            for col, txt, w in [("type","النوع",100),("school","المدرسة",100),("color","اللون",70),("size","المقاس",50),("qty","الكمية",50)]:
                tree.heading(col, text=txt)
                tree.column(col, width=w, anchor="center")
            for it in summary["inflow_items"]:
                tree.insert("", tk.END, values=(it["item_type"], it["school"], it["color"], it["size"], it["qty"]))
            tree.pack(fill=tk.X, padx=8, pady=4)

        sal = ttk.LabelFrame(main, text="المبيعات")
        sal.pack(fill=tk.X, pady=4)
        ttk.Label(sal, text=f"عدد الفواتير: {summary['sales_count']}").pack(anchor="w", padx=8, pady=2)
        ttk.Label(sal, text=f"إجمالي المبيعات: {summary['sales_total']:.2f}").pack(anchor="w", padx=8, pady=2)

        res = ttk.LabelFrame(main, text="الحجوزات")
        res.pack(fill=tk.X, pady=4)
        ttk.Label(res, text=f"عدد الحجوزات: {summary['res_count']}").pack(anchor="w", padx=8, pady=2)
        ttk.Label(res, text=f"إجمالي الحجوزات: {summary['res_total']:.2f}").pack(anchor="w", padx=8, pady=2)
        ttk.Label(res, text=f"المدفوع من الحجوزات: {summary['res_paid']:.2f}").pack(anchor="w", padx=8, pady=2)

        deliver_count = summary.get("deliver_count", 0)
        deliver_total = summary.get("deliver_total", 0.0)
        if deliver_count > 0:
            dlv = ttk.LabelFrame(main, text="تسليم الحجوزات")
            dlv.pack(fill=tk.X, pady=4)
            ttk.Label(dlv, text=f"عدد عمليات التسليم: {deliver_count}").pack(anchor="w", padx=8, pady=2)
            ttk.Label(dlv, text=f"إجمالي المحصّل عند التسليم: {deliver_total:.2f}").pack(anchor="w", padx=8, pady=2)

        return_count = summary.get("return_count", 0)
        return_total = summary.get("return_total", 0.0)
        if return_count > 0:
            ret = ttk.LabelFrame(main, text="المرتجعات")
            ret.pack(fill=tk.X, pady=4)
            ttk.Label(ret, text=f"عدد فواتير المرتجع: {return_count}").pack(anchor="w", padx=8, pady=2)
            ttk.Label(ret, text=f"إجمالي المرتجعات: {return_total:.2f}").pack(anchor="w", padx=8, pady=2)

        exchange_count = summary.get("exchange_count", 0)
        exchange_total = summary.get("exchange_total", 0.0)
        if exchange_count > 0:
            exc = ttk.LabelFrame(main, text="الاستبدالات")
            exc.pack(fill=tk.X, pady=4)
            ttk.Label(exc, text=f"عدد فواتير الاستبدال: {exchange_count}").pack(anchor="w", padx=8, pady=2)
            if exchange_total >= 0:
                ttk.Label(exc, text=f"صافي الاستبدال (محصّل): {exchange_total:.2f}").pack(anchor="w", padx=8, pady=2)
            else:
                ttk.Label(exc, text=f"صافي الاستبدال (مسترد): {abs(exchange_total):.2f}").pack(anchor="w", padx=8, pady=2)

        grand = summary["sales_total"] + summary["res_paid"] + deliver_total - return_total + exchange_total
        ttk.Label(main, text=f"إجمالي النقدية المحصلة: {grand:.2f}", font=("Segoe UI", 12, "bold")).pack(anchor="w", pady=8)

        btns = ttk.Frame(main)
        btns.pack(fill=tk.X, pady=(4, 0))
        ttk.Button(btns, text="طباعة الملخص", command=lambda: self._print_summary(summary)).pack(side=tk.LEFT)
        ttk.Button(btns, text="إلغاء", command=self.destroy).pack(side=tk.RIGHT)
        ttk.Button(btns, text="إنهاء الوردية", command=lambda: self._confirm_close(summary)).pack(side=tk.RIGHT, padx=8)

    def _confirm_close(self, summary):
        if messagebox.askyesno("تأكيد", "هل أنت متأكد من إنهاء الوردية؟", parent=self):
            try:
                self.db.end_shift(self.shift_id, summary_json=json.dumps(summary, ensure_ascii=False, default=str))
                root = self.master
                if self._on_closed:
                    self._on_closed()
                self.destroy()
                ToastNotification.show(root, "تم إنهاء الوردية بنجاح", toast_type="success")
            except Exception as ex:
                messagebox.showerror("خطأ", str(ex), parent=self)

    def _print_summary(self, summary):
        started = summary["started_at"][:16].replace("T", " ")
        deliver_total = summary.get("deliver_total", 0.0)
        deliver_count = summary.get("deliver_count", 0)
        return_count = summary.get("return_count", 0)
        return_total = summary.get("return_total", 0.0)
        exchange_count = summary.get("exchange_count", 0)
        exchange_total = summary.get("exchange_total", 0.0)
        grand = summary["sales_total"] + summary["res_paid"] + deliver_total - return_total + exchange_total
        inflow_rows = ""
        for it in summary.get("inflow_items", []):
            inflow_rows += f"<tr><td>{it['item_type']}</td><td>{it['school']}</td><td>{it['color']}</td><td>{it['size']}</td><td>{it['qty']}</td></tr>\n"
        deliver_html = f'<div><b>تسليم حجوزات:</b> {deliver_count} عملية - محصّل {deliver_total:.2f}</div>\n<hr class="sep">' if deliver_count > 0 else ""
        return_html = f'<div><b>مرتجعات:</b> {return_count} فاتورة - {return_total:.2f}</div>\n<hr class="sep">' if return_count > 0 else ""
        exchange_html = f'<div><b>استبدالات:</b> {exchange_count} فاتورة - صافي {exchange_total:.2f}</div>\n<hr class="sep">' if exchange_count > 0 else ""

        html = f"""<!DOCTYPE html>
<html lang="ar" dir="rtl"><head><meta charset="utf-8"><title>ملخص الوردية</title>
<style>
@page {{ size: 80mm auto; margin: 2mm; }}
body {{ font-family: "Segoe UI", Tahoma, sans-serif; font-size: 11px; width: 76mm; direction: rtl; margin: 0; padding: 2mm; }}
h2 {{ font-size: 14px; text-align: center; margin: 4px 0; }}
.sep {{ border: none; border-top: 1px dashed #000; margin: 4px 0; }}
table {{ border-collapse: collapse; width: 100%; font-size: 10px; }}
th {{ border-bottom: 1px solid #000; padding: 2px; text-align: right; }}
td {{ padding: 2px; border-bottom: 1px dotted #ccc; }}
.total {{ font-size: 13px; font-weight: bold; text-align: center; margin: 6px 0; }}
</style></head><body>
<h2>ملخص الوردية #{summary['shift_id']}</h2>
<hr class="sep">
<div>بداية: {started}</div>
<div>نهاية: الآن</div>
<hr class="sep">
<div><b>الوارد:</b> {summary['inflow_count']} عملية - {summary['inflow_total_qty']} قطعة</div>
{"<table><thead><tr><th>النوع</th><th>المدرسة</th><th>اللون</th><th>المقاس</th><th>الكمية</th></tr></thead><tbody>" + inflow_rows + "</tbody></table>" if inflow_rows else ""}
<hr class="sep">
<div><b>المبيعات:</b> {summary['sales_count']} فاتورة - {summary['sales_total']:.2f}</div>
<hr class="sep">
<div><b>الحجوزات:</b> {summary['res_count']} - إجمالي {summary['res_total']:.2f} - مدفوع {summary['res_paid']:.2f}</div>
<hr class="sep">
{deliver_html}{return_html}{exchange_html}<div class="total">النقدية المحصلة: {grand:.2f}</div>
</body></html>"""

        tmp = os.path.join(tempfile.gettempdir(), f"shift_{summary['shift_id']}.html")
        with open(tmp, "w", encoding="utf-8") as f:
            f.write(html)
        _print_html_auto(tmp, copies=1, parent=self)


# ------------------- Bulk Price Dialog -------------------

class BulkPriceDialog(tk.Toplevel):
    """Dialog for bulk price update by percentage or fixed value."""
    def __init__(self, master, db: SqliteDatabase, on_done=None):
        super().__init__(master)
        self.db = db
        self._on_done = on_done
        self.title("تعديل الأسعار جماعياً")
        self.geometry("450x350")
        self.resizable(False, False)
        self.transient(master)
        self.grab_set()
        self._build()

    def _build(self):
        frm = ttk.Frame(self, padding=12)
        frm.pack(fill=tk.BOTH, expand=True)

        ttk.Label(frm, text="تعديل الأسعار جماعياً", font=("Segoe UI", 13, "bold")).pack(anchor="w", pady=(0, 8))

        # Mode selection
        mode_frm = ttk.LabelFrame(frm, text="طريقة التعديل")
        mode_frm.pack(fill=tk.X, pady=4)
        self._mode_var = tk.StringVar(value="percentage")
        ttk.Radiobutton(mode_frm, text="نسبة مئوية (يتم التقريب لأقرب 5)", variable=self._mode_var, value="percentage").pack(anchor="w", padx=8, pady=2)
        ttk.Radiobutton(mode_frm, text="قيمة ثابتة (تُضاف للسعر الحالي)", variable=self._mode_var, value="fixed").pack(anchor="w", padx=8, pady=2)

        val_frm = ttk.Frame(frm)
        val_frm.pack(fill=tk.X, pady=4)
        ttk.Label(val_frm, text="القيمة:").pack(side=tk.LEFT)
        self._value_var = tk.StringVar(value="10")
        ttk.Entry(val_frm, textvariable=self._value_var, width=12).pack(side=tk.LEFT, padx=8)
        ttk.Label(val_frm, text="(نسبة % أو مبلغ حسب الاختيار)").pack(side=tk.LEFT)

        # Optional filters
        filt = ttk.LabelFrame(frm, text="تطبيق على (اتركها فارغة لتعديل الكل)")
        filt.pack(fill=tk.X, pady=4)
        self._f_school = tk.StringVar()
        self._f_item = tk.StringVar()
        r0 = ttk.Frame(filt)
        r0.pack(fill=tk.X, padx=8, pady=2)
        ttk.Label(r0, text="المدرسة:").pack(side=tk.LEFT)
        ttk.Entry(r0, textvariable=self._f_school, width=20).pack(side=tk.LEFT, padx=4)
        ttk.Label(r0, text="النوع:").pack(side=tk.LEFT, padx=(12, 0))
        ttk.Entry(r0, textvariable=self._f_item, width=20).pack(side=tk.LEFT, padx=4)

        # Preview
        ttk.Label(frm, text="مثال: سعر 231 بنسبة 10% = 255 (231×1.1=254.1 → 255)", font=("Segoe UI", 9)).pack(anchor="w", pady=(8, 0))

        # Buttons
        btns = ttk.Frame(frm)
        btns.pack(fill=tk.X, pady=(12, 0))
        ttk.Button(btns, text="تطبيق", command=self._apply).pack(side=tk.LEFT)
        ttk.Button(btns, text="إلغاء", command=self.destroy).pack(side=tk.RIGHT)

    def _apply(self):
        try:
            value = float(self._value_var.get().strip())
        except ValueError:
            messagebox.showwarning("خطأ", "أدخل قيمة رقمية صحيحة.", parent=self)
            return

        mode = self._mode_var.get()
        constraints: Dict[str, Any] = {}
        sc = self._f_school.get().strip()
        it = self._f_item.get().strip()
        if sc:
            constraints["school"] = sc
        if it:
            constraints["item_type"] = it

        scope_text = "جميع الأصناف" if not constraints else f"الأصناف المحددة ({', '.join(constraints.values())})"
        if mode == "percentage":
            desc = f"زيادة {value}% على {scope_text}"
        else:
            desc = f"إضافة {value} على {scope_text}"

        if not messagebox.askyesno("تأكيد", f"هل تريد تطبيق: {desc}؟", parent=self):
            return

        try:
            count = self.db.bulk_update_prices(mode, value, constraints if constraints else None)
            ToastNotification.show(self.winfo_toplevel(), f"تم تعديل أسعار {count} صنف بنجاح", toast_type="success")
            if self._on_done:
                self._on_done()
            self.destroy()
        except Exception as ex:
            messagebox.showerror("خطأ", str(ex), parent=self)


# ------------------- Main Application -------------------

class WarehouseApp:
    """Main application controller."""

    def __init__(self, root: tk.Tk, db: SqliteDatabase):
        self.root = root
        self.db = db
        self._current_shift_id: Optional[int] = None
        self.root._app_controller = self
        root.title("\u0625\u062F\u0627\u0631\u0629 \u0627\u0644\u0645\u062E\u0627\u0632\u0646 \u0648\u0627\u0644\u0645\u0628\u064A\u0639\u0627\u062A")
        root.geometry("1280x760")
        root.minsize(900, 600)
        self._build()
        self.root.protocol("WM_DELETE_WINDOW", self._on_close)
        self._check_shift_state()
        self._bind_shortcuts()
        self._update_clock()
        try:
            import sync_periodic

            sync_periodic.attach_periodic_sync(self.root, self.db.path)
        except Exception:
            pass

    def _build(self):
        # ---- Branded Header Bar ----
        T = getattr(self.root, "_theme", THEME_LIGHT)
        self._header = tk.Frame(self.root, bg=T["HEADER_BG"], height=48)
        self._header.pack(fill=tk.X, side=tk.TOP)
        self._header.pack_propagate(False)

        tk.Label(self._header, text="\u0625\u062F\u0627\u0631\u0629 \u0627\u0644\u0645\u062E\u0627\u0632\u0646 \u0648\u0627\u0644\u0645\u0628\u064A\u0639\u0627\u062A",
                 bg=T["HEADER_BG"], fg=T["HEADER_FG"],
                 font=("Segoe UI", 14, "bold")).pack(side=tk.LEFT, padx=16)

        # Shift status in header
        self._shift_status_var = tk.StringVar(value="")
        tk.Label(self._header, textvariable=self._shift_status_var,
                 bg=T["HEADER_BG"], fg="#93c5fd",
                 font=("Segoe UI", 9, "bold")).pack(side=tk.LEFT, padx=20)

        # Clock on right
        self._clock_var = tk.StringVar(value="")
        tk.Label(self._header, textvariable=self._clock_var,
                 bg=T["HEADER_BG"], fg=T["HEADER_FG"],
                 font=("Segoe UI", 11)).pack(side=tk.RIGHT, padx=16)

        # ---- Toolbar ----
        self._toolbar = ttk.Frame(self.root)
        self._toolbar.pack(fill=tk.X, side=tk.TOP, padx=6, pady=(4, 0))

        # Old fully-disabled toolbar buttons were commented during the first lockdown pass.
        ttk.Button(self._toolbar, text="\u0627\u0644\u0645\u062E\u0632\u0648\u0646", command=self._open_inventory).pack(side=tk.LEFT, padx=2)
        ttk.Button(self._toolbar, text="\u0633\u062C\u0644 \u0627\u0644\u0641\u0648\u0627\u062A\u064A\u0631", command=self._open_bills_history).pack(side=tk.LEFT, padx=2)
        ttk.Button(self._toolbar, text="\u0633\u062C\u0644 \u0627\u0644\u062D\u0631\u0643\u0627\u062A", command=self._open_movements).pack(side=tk.LEFT, padx=2)
        ttk.Button(self._toolbar, text="\u062A\u0639\u062F\u064A\u0644 \u0627\u0644\u0623\u0633\u0639\u0627\u0631", command=self._open_bulk_price).pack(side=tk.LEFT, padx=2)

        # Right side toolbar
        ttk.Button(self._toolbar, text="\u0627\u0644\u0625\u0639\u062F\u0627\u062F\u0627\u062A", command=self._open_admin).pack(side=tk.RIGHT, padx=2)
        ttk.Button(self._toolbar, text="\u0627\u0644\u0645\u0632\u0627\u0645\u0646\u0629", command=self._open_sync_dialog).pack(side=tk.RIGHT, padx=2)

        # Dark mode toggle
        self._dark_mode_var = tk.BooleanVar(value=False)
        self._dark_btn = ttk.Button(self._toolbar, text="\u263E \u0627\u0644\u0648\u0636\u0639 \u0627\u0644\u0644\u064A\u0644\u064A",
                                    command=self._toggle_dark_mode, style="Secondary.TButton")
        self._dark_btn.pack(side=tk.RIGHT, padx=4)

        # Shift button
        self._shift_btn = ttk.Button(self._toolbar, text="\u0628\u062F\u0621 \u0648\u0631\u062F\u064A\u0629", command=self._toggle_shift)
        self._shift_btn.pack(side=tk.LEFT, padx=(16, 2))

        # ---- Main notebook ----
        self._nb = ttk.Notebook(self.root)
        self._nb.pack(fill=tk.BOTH, expand=True, padx=6, pady=(4, 0))

        # Tab 0: Dashboard (Home)
        dash_tab = ttk.Frame(self._nb)
        self._nb.add(dash_tab, text="\u2302 \u0627\u0644\u0631\u0626\u064A\u0633\u064A\u0629")
        self._dashboard = DashboardFrame(dash_tab, self.db)
        self._dashboard.pack(fill=tk.BOTH, expand=True)

        # Tab 1: Income
        income_tab = ttk.Frame(self._nb)
        self._nb.add(income_tab, text="\u0627\u0644\u0648\u0627\u0631\u062F")
        self._income_frame = IncomeFrame(income_tab, self.db)
        self._income_frame.pack(fill=tk.BOTH, expand=True)

        # Tab 2: POS
        pos_tab = ttk.Frame(self._nb)
        self._nb.add(pos_tab, text="\u0646\u0642\u0637\u0629 \u0627\u0644\u0628\u064A\u0639")
        self._pos_frame = POSFrame(pos_tab, self.db)
        self._pos_frame.pack(fill=tk.BOTH, expand=True)

        # Tab 3: Reservations
        res_tab = ttk.Frame(self._nb)
        self._nb.add(res_tab, text="\u0627\u0644\u062D\u062C\u0648\u0632\u0627\u062A")
        self._res_frame = ReservationsFrame(res_tab, self.db)
        self._res_frame.pack(fill=tk.BOTH, expand=True)

        # Tab 4: Statistics
        stats_tab = ttk.Frame(self._nb)
        self._nb.add(stats_tab, text="\u0627\u0644\u0625\u062D\u0635\u0627\u0626\u064A\u0627\u062A")
        self._stats_frame = StatisticsFrame(stats_tab, self.db)
        self._stats_frame.pack(fill=tk.BOTH, expand=True)

        # Tab 5: Shifts Summary
        shifts_tab = ttk.Frame(self._nb)
        self._nb.add(shifts_tab, text="ملخص الورديات")
        self._shifts_frame = ShiftsSummaryFrame(shifts_tab, self.db)
        self._shifts_frame.pack(fill=tk.BOTH, expand=True)

        self._nb.bind("<<NotebookTabChanged>>", self._on_tab_changed)

        # ---- Status Bar ----
        self._statusbar = tk.Frame(self.root, bg=T.get("EDGE", "#cbd5e1"), height=26)
        self._statusbar.pack(fill=tk.X, side=tk.BOTTOM)
        self._statusbar.pack_propagate(False)
        self._status_var = tk.StringVar(value="\u062C\u0627\u0647\u0632")
        tk.Label(self._statusbar, textvariable=self._status_var,
                 bg=T.get("EDGE", "#cbd5e1"), fg=T.get("TEXT", "#0f172a"),
                 font=("Segoe UI", 8), anchor="w").pack(side=tk.LEFT, padx=8)
        self._status_items_var = tk.StringVar(value="")
        tk.Label(self._statusbar, textvariable=self._status_items_var,
                 bg=T.get("EDGE", "#cbd5e1"), fg=T.get("TEXT", "#0f172a"),
                 font=("Segoe UI", 8), anchor="e").pack(side=tk.RIGHT, padx=8)

        try:
            self._dashboard.refresh()
        except Exception:
            pass

    def refresh_dashboard(self):
        try:
            self._dashboard.refresh()
        except Exception:
            pass

    def _update_status(self, msg):
        self._status_var.set(msg)
        self.root.after(5000, lambda: self._status_var.set("\u062C\u0627\u0647\u0632"))

    def _update_clock(self):
        try:
            now = datetime.now()
            self._clock_var.set(now.strftime("%Y-%m-%d  %H:%M:%S"))
        except Exception:
            pass
        self.root.after(1000, self._update_clock)

    def _bind_shortcuts(self):
        self.root.bind("<F5>", lambda e: self._shortcut_refresh())
        self.root.bind("<F8>", lambda e: self._shortcut_finalize())
        self.root.bind("<Escape>", lambda e: self._shortcut_escape())
        self.root.bind("<Control-p>", lambda e: self._shortcut_print())
        self.root.bind("<Delete>", lambda e: self._shortcut_delete())

    def _shortcut_refresh(self):
        try:
            tab = self._nb.tab(self._nb.select(), "text")
            if "\u0627\u0644\u0625\u062D\u0635\u0627\u0626\u064A\u0627\u062A" in tab:
                self._stats_frame._refresh_all()
            elif "\u0627\u0644\u062D\u062C\u0648\u0632\u0627\u062A" in tab:
                self._res_frame._refresh()
            elif "\u0627\u0644\u0631\u0626\u064A\u0633\u064A\u0629" in tab:
                self._dashboard.refresh()
            self._update_status("F5: \u062A\u0645 \u0627\u0644\u062A\u062D\u062F\u064A\u062B")
        except Exception:
            pass

    def _shortcut_finalize(self):
        try:
            tab = self._nb.tab(self._nb.select(), "text")
            if "\u0646\u0642\u0637\u0629 \u0627\u0644\u0628\u064A\u0639" in tab:
                self._pos_frame._finalize()
            elif "\u0627\u0644\u0648\u0627\u0631\u062F" in tab and self._income_frame is not None:
                self._income_frame._finalize_income()
        except Exception:
            pass

    def _shortcut_escape(self):
        try:
            tab = self._nb.tab(self._nb.select(), "text")
            if "\u0646\u0642\u0637\u0629 \u0627\u0644\u0628\u064A\u0639" in tab:
                self._pos_frame._clear_bill()
        except Exception:
            pass

    def _shortcut_print(self):
        try:
            tab = self._nb.tab(self._nb.select(), "text")
            if "\u0627\u0644\u0625\u062D\u0635\u0627\u0626\u064A\u0627\u062A" in tab:
                self._stats_frame._print_movements()
        except Exception:
            pass

    def _shortcut_delete(self):
        try:
            tab = self._nb.tab(self._nb.select(), "text")
            if "\u0646\u0642\u0637\u0629 \u0627\u0644\u0628\u064A\u0639" in tab:
                self._pos_frame._remove_line()
        except Exception:
            pass

    def _toggle_dark_mode(self):
        is_dark = not getattr(self.root, "_theme_dark", False)
        _apply_colorful_theme_to_root(self.root, dark=is_dark)
        T = self.root._theme
        # Update header
        for w in (self._header,):
            w.configure(bg=T["HEADER_BG"])
        for w in self._header.winfo_children():
            try:
                w.configure(bg=T["HEADER_BG"])
            except Exception:
                pass
        # Update status bar
        self._statusbar.configure(bg=T["EDGE"])
        for w in self._statusbar.winfo_children():
            try:
                w.configure(bg=T["EDGE"], fg=T["TEXT"])
            except Exception:
                pass
        self._dark_btn.configure(
            text="\u2600 \u0627\u0644\u0648\u0636\u0639 \u0627\u0644\u0639\u0627\u062F\u064A" if is_dark else "\u263E \u0627\u0644\u0648\u0636\u0639 \u0627\u0644\u0644\u064A\u0644\u064A")
        self._update_status("\u062A\u0645 \u062A\u063A\u064A\u064A\u0631 \u0627\u0644\u0633\u0645\u0629" + (" \u0627\u0644\u0644\u064A\u0644\u064A\u0629" if is_dark else " \u0627\u0644\u0639\u0627\u062F\u064A\u0629"))

    def _on_tab_changed(self, event=None):
        try:
            current = self._nb.select()
            tab_text = self._nb.tab(current, "text")
            if "\u0627\u0644\u0625\u062D\u0635\u0627\u0626\u064A\u0627\u062A" in tab_text:
                self._stats_frame._refresh_all()
            elif "\u0627\u0644\u0631\u0626\u064A\u0633\u064A\u0629" in tab_text:
                self._dashboard.refresh()
            elif "\u0627\u0644\u062D\u062C\u0648\u0632\u0627\u062A" in tab_text:
                self._res_frame._refresh()
            elif "ملخص الورديات" in tab_text:
                self._shifts_frame._refresh_all()
        except Exception:
            pass

    def _open_inventory(self):
        if not self.db.is_manager_feature_enabled("allow_inventory_window"):
            messagebox.showwarning("مقيد", _feature_restricted_message("نافذة المخزون مقيدة حالياً في نقطة البيع."), parent=self.root)
            return
        try:
            InventoryWindow(self.root, self.db)
        except Exception as ex:
            messagebox.showerror("خطأ", str(ex))

    def _open_bills_history(self):
        try:
            BillsHistoryWindow(self.root, self.db)
        except Exception as ex:
            messagebox.showerror("خطأ", str(ex))

    def _open_movements(self):
        try:
            MovementsWindow(self.root, self.db)
        except Exception as ex:
            messagebox.showerror("خطأ", str(ex))

    def _open_admin(self):
        try:
            AdminWindow(self.root, self.db)
        except Exception as ex:
            messagebox.showerror("خطأ", str(ex))

    def _open_sync_dialog(self):
        try:
            import sync_ui
            sync_ui.open_sync_dialog(self.root, self.db.conn)
        except Exception as ex:
            messagebox.showerror("المزامنة", f"تعذّر فتح نافذة المزامنة:\n{ex}")

    def _open_bulk_price(self):
        if not self.db.is_manager_feature_enabled("allow_bulk_price"):
            messagebox.showwarning("مقيد", _feature_restricted_message("تعديل الأسعار من نقطة البيع غير مسموح به حالياً."), parent=self.root)
            return
        pw = simpledialog.askstring("كلمة مرور المدير", "أدخل كلمة مرور المدير:", show="*", parent=self.root)
        if not pw:
            return
        if not self.db.verify_admin_password(pw):
            messagebox.showerror("مرفوض", "كلمة المرور غير صحيحة.", parent=self.root)
            return
        BulkPriceDialog(self.root, self.db)

    def _check_shift_state(self):
        shift = self.db.get_open_shift()
        if shift:
            self._current_shift_id = shift["id"]
            self.db.active_shift_id = shift["id"]
            self._update_shift_ui(shift)
            self._set_app_enabled(True)
            return
        total = self.db.conn.execute("SELECT COUNT(*) AS c FROM shifts").fetchone()["c"]
        if total == 0:
            sid = self.db.start_shift()
            self._current_shift_id = sid
            self.db.active_shift_id = sid
            shift = self.db.get_open_shift()
            self._update_shift_ui(shift)
            self._set_app_enabled(True)
            messagebox.showinfo("نظام الورديات", "تم تفعيل نظام الورديات.\nتم فتح وردية جديدة تلقائياً.")
        else:
            self._current_shift_id = None
            self.db.active_shift_id = None
            self._shift_status_var.set("لا توجد وردية مفتوحة")
            self._shift_btn.configure(text="بدء وردية")
            self._set_app_enabled(False)
            self._prompt_start_shift()

    def _update_shift_ui(self, shift):
        started = shift["started_at"][:16].replace("T", " ")
        self._shift_status_var.set(f"الوردية مفتوحة منذ {started}")
        self._shift_btn.configure(text="إنهاء الوردية")

    def _toggle_shift(self):
        if self._current_shift_id:
            self._prompt_end_shift()
        else:
            self._prompt_start_shift()

    def _prompt_start_shift(self):
        if messagebox.askyesno("بدء وردية جديدة", "هل تريد بدء وردية جديدة؟\nيجب فتح وردية قبل البدء بالعمل."):
            try:
                sid = self.db.start_shift()
                self._current_shift_id = sid
                self.db.active_shift_id = sid
                shift = self.db.get_open_shift()
                self._update_shift_ui(shift)
                self._set_app_enabled(True)
            except Exception as ex:
                messagebox.showerror("خطأ", str(ex))

    def _prompt_end_shift(self):
        ShiftSummaryDialog(self.root, self.db, self._current_shift_id, on_closed=self._on_shift_closed)

    def _on_shift_closed(self):
        self._current_shift_id = None
        self.db.active_shift_id = None
        self._shift_status_var.set("لا توجد وردية مفتوحة")
        self._shift_btn.configure(text="بدء وردية")
        self._set_app_enabled(False)

    def _set_app_enabled(self, enabled: bool):
        state = "normal" if enabled else "disabled"
        for tab_id in self._nb.tabs():
            try:
                self._nb.tab(tab_id, state=state)
            except Exception:
                pass
        for child in self._toolbar.winfo_children():
            if child is not self._shift_btn:
                try:
                    child.configure(state=state)
                except Exception:
                    pass

    def _on_close(self):
        if self._current_shift_id:
            result = messagebox.askyesnocancel(
                "إغلاق البرنامج",
                "يوجد وردية مفتوحة.\nهل تريد إنهاء الوردية قبل الإغلاق؟\n\nنعم = إنهاء الوردية ثم إغلاق\nلا = إغلاق بدون إنهاء\nإلغاء = العودة")
            if result is True:
                ShiftSummaryDialog(self.root, self.db, self._current_shift_id,
                                   on_closed=lambda: self.root.destroy())
            elif result is False:
                self.root.destroy()
        else:
            self.root.destroy()


# ------------------- Entry Point -------------------

def main():
    db = SqliteDatabase(DB_PATH)
    root = tk.Tk()
    # Apply the colorful theme
    _apply_colorful_theme_to_root(root)
    _app = WarehouseApp(root, db)
    root.mainloop()


_DARK_MODE = False  # Global toggle

THEME_LIGHT = {
    "BG": "#f8fafc", "SURFACE": "#e2e8f0", "CARD": "#ffffff",
    "ACCENT": "#2563eb", "ACCENT2": "#1d4ed8", "SUCCESS": "#16a34a", "SUCCESS2": "#15803d",
    "DANGER": "#dc2626", "DANGER2": "#b91c1c", "WARNING": "#f59e0b",
    "TEXT": "#0f172a", "TEXT2": "#475569", "SUBTLE": "#94a3b8",
    "ROW": "#ffffff", "ROW_ALT": "#f1f5f9", "SELBG": "#bfdbfe",
    "EDGE": "#cbd5e1", "HEADER_BG": "#1e3a5f", "HEADER_FG": "#ffffff",
}

THEME_DARK = {
    "BG": "#0f172a", "SURFACE": "#1e293b", "CARD": "#334155",
    "ACCENT": "#3b82f6", "ACCENT2": "#2563eb", "SUCCESS": "#22c55e", "SUCCESS2": "#16a34a",
    "DANGER": "#ef4444", "DANGER2": "#dc2626", "WARNING": "#fbbf24",
    "TEXT": "#f1f5f9", "TEXT2": "#94a3b8", "SUBTLE": "#64748b",
    "ROW": "#1e293b", "ROW_ALT": "#334155", "SELBG": "#1e40af",
    "EDGE": "#475569", "HEADER_BG": "#020617", "HEADER_FG": "#e2e8f0",
}


def _apply_colorful_theme_to_root(root, dark=False):
    """Apply a modern colour palette with semantic accent colours."""
    global _DARK_MODE
    _DARK_MODE = dark
    T = THEME_DARK if dark else THEME_LIGHT
    s = ttk.Style(root)
    try:
        s.theme_use("clam")
    except Exception:
        pass

    BG, SURFACE, ACCENT, ACCENT2 = T["BG"], T["SURFACE"], T["ACCENT"], T["ACCENT2"]
    TEXT, SUBTLE, ROW, SELBG, EDGE = T["TEXT"], T["SUBTLE"], T["ROW"], T["SELBG"], T["EDGE"]
    ROW_ALT = T["ROW_ALT"]
    SUCCESS, SUCCESS2 = T["SUCCESS"], T["SUCCESS2"]
    DANGER, DANGER2 = T["DANGER"], T["DANGER2"]

    try:
        root.configure(bg=BG)
    except Exception:
        pass

    # Base widgets
    s.configure("TFrame", background=SURFACE)
    s.configure("TLabelframe", background=SURFACE, bordercolor=EDGE, relief="groove")
    s.configure("TLabelframe.Label", background=SURFACE, foreground=TEXT,
                font=("Segoe UI", 10, "bold"))
    s.configure("TLabel", background=SURFACE, foreground=TEXT)
    s.configure("TPanedwindow", background=SURFACE)
    s.configure("TSeparator", background=EDGE)

    # Primary button — blue
    s.configure("TButton", background=ACCENT, foreground="#ffffff",
                bordercolor=ACCENT, focusthickness=1, focuscolor=ACCENT2,
                padding=(12, 6), font=("Segoe UI", 9, "bold"))
    s.map("TButton",
          background=[("active", ACCENT2), ("pressed", ACCENT2), ("disabled", SUBTLE)],
          foreground=[("disabled", "#94a3b8")],
          bordercolor=[("focus", ACCENT2)])

    # Success button — green
    s.configure("Success.TButton", background=SUCCESS, foreground="#ffffff",
                bordercolor=SUCCESS, focuscolor=SUCCESS2, padding=(12, 6),
                font=("Segoe UI", 10, "bold"))
    s.map("Success.TButton",
          background=[("active", SUCCESS2), ("pressed", SUCCESS2)],
          bordercolor=[("focus", SUCCESS2)])

    # Danger button — red
    s.configure("Danger.TButton", background=DANGER, foreground="#ffffff",
                bordercolor=DANGER, focuscolor=DANGER2, padding=(12, 6),
                font=("Segoe UI", 9, "bold"))
    s.map("Danger.TButton",
          background=[("active", DANGER2), ("pressed", DANGER2)],
          bordercolor=[("focus", DANGER2)])

    # Secondary / subtle button
    s.configure("Secondary.TButton", background=SURFACE, foreground=TEXT,
                bordercolor=EDGE, padding=(10, 5), font=("Segoe UI", 9))
    s.map("Secondary.TButton",
          background=[("active", EDGE), ("pressed", EDGE)],
          foreground=[("active", "#ffffff")])

    # Inputs
    for sty in ("TEntry", "TCombobox", "TSpinbox"):
        s.configure(sty, fieldbackground=ROW, background=ROW, foreground=TEXT,
                    bordercolor=EDGE, lightcolor=EDGE, darkcolor=EDGE,
                    padding=4)
        s.map(sty, fieldbackground=[("focus", "#ffffff" if not dark else "#475569")],
              bordercolor=[("focus", ACCENT)])

    s.configure("TCheckbutton", background=SURFACE, foreground=TEXT)
    s.map("TCheckbutton", background=[("active", SURFACE)])
    s.configure("TRadiobutton", background=SURFACE, foreground=TEXT)
    s.map("TRadiobutton", background=[("active", SURFACE)])

    # Notebook tabs
    s.configure("TNotebook", background=BG, bordercolor=EDGE)
    s.configure("TNotebook.Tab", background=SURFACE, foreground=SUBTLE,
                padding=[14, 6], font=("Segoe UI", 9, "bold"))
    s.map("TNotebook.Tab",
          background=[("selected", BG)],
          foreground=[("selected", ACCENT)])

    # Treeview — with alternating row support
    s.configure("Treeview",
                background=ROW, fieldbackground=ROW, foreground=TEXT,
                bordercolor=EDGE, rowheight=26,
                font=("Segoe UI", 9))
    s.configure("Treeview.Heading",
                background=SURFACE, foreground=TEXT, bordercolor=EDGE,
                font=("Segoe UI", 9, "bold"), padding=4)
    s.map("Treeview",
          background=[("selected", SELBG)],
          foreground=[("selected", TEXT)])

    # Scrollbar
    s.configure("TScrollbar", background=SURFACE, troughcolor=BG,
                bordercolor=EDGE, arrowcolor=TEXT)

    # Store theme in root for access
    root._theme = T
    root._theme_dark = dark


if __name__ == "__main__":
    main()
